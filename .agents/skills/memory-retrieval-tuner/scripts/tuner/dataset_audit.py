from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook

from .benchmark_support import load_dataset, parse_gold_requirements
from .fact_evaluator import parse_required_context, uses_context_gold
from .io_utils import atomic_write_json, atomic_write_text, load_json, load_jsonl, sha256_file
from .models import Dataset, Requirement, Turn


class DatasetAuditFailed(RuntimeError):
    pass


COLUMN_ALIASES = {
    "id": ("编号", "query_id", "turn_id", "id"),
    "question": ("当前问题", "question", "query"),
    "answer": ("最终回答", "answer", "response"),
    "gold": ("关联前序对话", "gold", "gold_ids", "dependency_turn_ids"),
    "required_context": ("所需前文信息", "required_context", "gold_context"),
    "dependency_type": ("依赖类型", "dependency_type"),
}

EXPLICIT_HISTORY_PATTERNS = (
    re.compile(r"回到.{0,8}(?:轮前|前面|之前).{0,40}[【\[]"),
    re.compile(r"(?:第|Q)\s*\d+\s*(?:轮|次)?", re.IGNORECASE),
    re.compile(r"(?:go back|turn)\s+\d+", re.IGNORECASE),
    re.compile(r"(?:上|前)\s*\d+\s*轮"),
)


def _header_index(headers: Sequence[Any], sheet_name: str) -> dict[str, int]:
    normalized = {str(value or "").strip().lower(): index for index, value in enumerate(headers)}
    result: dict[str, int] = {}
    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias.lower() in normalized:
                result[canonical] = normalized[alias.lower()]
                break
    missing = [name for name in ("id", "question", "answer", "gold") if name not in result]
    if missing:
        raise ValueError(f"Sheet {sheet_name} missing required benchmark columns: {missing}")
    return result


def _session_code(sheet_name: str, query_id: str) -> str:
    match = re.match(r"(S\d{3})", query_id, re.IGNORECASE)
    return match.group(1).upper() if match else sheet_name.split("_", 1)[0]


def load_benchmark_dataset(path: Path, sessions: Sequence[str] | None = None) -> Dataset:
    """Load the repository workbook schema while preserving AND/OR Gold groups."""
    path = path.resolve()
    requested = {value.upper() for value in sessions or []}
    workbook = load_workbook(path, read_only=True, data_only=True)
    loaded: dict[str, tuple[Turn, ...]] = {}
    try:
        for sheet_name in workbook.sheetnames:
            sheet_code = sheet_name.split("_", 1)[0].upper()
            if requested and sheet_name.upper() not in requested and sheet_code not in requested:
                continue
            rows = workbook[sheet_name].iter_rows(values_only=True)
            try:
                headers = next(rows)
            except StopIteration:
                continue
            try:
                index = _header_index(headers, sheet_name)
            except ValueError:
                explicitly_requested = bool(requested and (sheet_name.upper() in requested or sheet_code in requested))
                if explicitly_requested:
                    raise
                continue
            turns: list[Turn] = []
            for row in rows:
                values = list(row)

                def cell(key: str) -> str:
                    position = index.get(key)
                    return (
                        str(values[position] or "").strip() if position is not None and position < len(values) else ""
                    )

                query_id = cell("id").upper()
                question = cell("question")
                answer = cell("answer")
                if not query_id and not question and not answer:
                    continue
                gold_raw = cell("gold")
                parsed = parse_gold_requirements(gold_raw)
                requirements = tuple(Requirement(item.members, item.raw_text) for item in parsed)
                turns.append(
                    Turn(
                        session_id=sheet_name,
                        session_code=_session_code(sheet_name, query_id),
                        turn_index=len(turns),
                        query_id=query_id,
                        question=question,
                        answer=answer,
                        gold_raw=gold_raw,
                        requirements=requirements,
                        required_context=cell("required_context"),
                        dependency_type=cell("dependency_type"),
                    )
                )
            if turns:
                if not any(re.fullmatch(r"S\d{3}-Q\d{3}", turn.query_id, re.IGNORECASE) for turn in turns):
                    explicitly_requested = bool(
                        requested and (sheet_name.upper() in requested or sheet_code in requested)
                    )
                    if explicitly_requested:
                        raise ValueError(f"Sheet {sheet_name} contains no valid Session Query IDs")
                    continue
                loaded[sheet_name] = tuple(turns)
    finally:
        workbook.close()
    if not loaded:
        raise ValueError(f"No benchmark Sessions found in {path}")
    return Dataset(path=str(path), sha256=sha256_file(path), sessions=loaded)


def _normalize_template(question: str) -> str:
    value = re.sub(r"\d+(?:\.\d+)?", "#", question.lower())
    value = re.sub(r"【[^】]{4,}】|\[[^\]]{4,}\]", "<quoted-history>", value)
    value = re.sub(r"\s+", "", value)
    return value


def _source_run_check(
    source_run: Path | None,
    expected_query_ids: set[str],
    *,
    expected_checkpoint_query_ids: set[str],
    expected_dataset_sha256: str,
    expected_session_turn_counts: Mapping[str, int],
) -> dict[str, Any]:
    if source_run is None:
        return {"required": False, "status": "AUTO_DISCOVERY_OR_PRODUCTION_GENERATION"}
    result_file = source_run / "recall_turn_results.jsonl" if source_run.is_dir() else source_run
    if result_file.exists():
        actual: list[str] = []
        failed = 0
        for row in load_jsonl(result_file):
            query_id = str(row.get("turn_id") or row.get("query_id") or "").upper()
            if query_id:
                actual.append(query_id)
            failed += int(bool(row.get("error")))
        missing = sorted(expected_query_ids - set(actual))
        duplicates = sorted(item for item, count in Counter(actual).items() if count > 1)
        status = "COMPLETE" if not missing and not duplicates and failed == 0 else "INCOMPLETE"
        return {
            "required": True,
            "kind": "full_production_trace",
            "status": status,
            "path": str(result_file),
            "expected_queries": len(expected_query_ids),
            "actual_unique_queries": len(set(actual)),
            "failed_turns": failed,
            "missing_queries": missing,
            "duplicate_results": duplicates,
        }

    if not source_run.is_dir():
        return {"required": True, "status": "INCOMPLETE", "reason": f"missing {result_file}"}
    actual_checkpoints: list[str] = []
    valid_sessions: set[str] = set()
    invalid_manifests: list[str] = []
    for manifest_path in source_run.glob("**/production_midterm_manifest.json"):
        try:
            manifest = load_json(manifest_path)
            session_id = str(manifest.get("session_id") or "")
            checkpoint_path = Path(str(manifest.get("checkpoints_path") or ""))
            valid = (
                manifest.get("status") == "COMPLETE"
                and manifest.get("dataset_sha256") == expected_dataset_sha256
                and session_id in expected_session_turn_counts
                and int(manifest.get("turn_count") or 0) == expected_session_turn_counts[session_id]
                and int(manifest.get("failed_turns") or 0) == 0
                and checkpoint_path.exists()
                and manifest.get("checkpoints_sha256") == sha256_file(checkpoint_path)
            )
            if not valid:
                continue
            valid_sessions.add(session_id)
            actual_checkpoints.extend(
                str(row.get("query_id") or "").upper() for row in load_jsonl(checkpoint_path) if row.get("query_id")
            )
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            invalid_manifests.append(f"{manifest_path}: {type(exc).__name__}")
    missing_sessions = sorted(set(expected_session_turn_counts) - valid_sessions)
    missing_queries = sorted(expected_checkpoint_query_ids - set(actual_checkpoints))
    duplicates = sorted(item for item, count in Counter(actual_checkpoints).items() if count > 1)
    complete = not missing_sessions and not missing_queries and not duplicates
    return {
        "required": True,
        "kind": "production_midterm_checkpoints",
        "status": "COMPLETE" if complete else "INCOMPLETE",
        "path": str(source_run),
        "expected_sessions": len(expected_session_turn_counts),
        "actual_sessions": len(valid_sessions),
        "expected_checkpoint_queries": len(expected_checkpoint_query_ids),
        "actual_unique_checkpoint_queries": len(set(actual_checkpoints)),
        "missing_sessions": missing_sessions,
        "missing_queries": missing_queries,
        "duplicate_results": duplicates,
        "invalid_manifests": invalid_manifests,
    }


def audit_dataset(
    dataset_path: Path,
    *,
    output_dir: Path,
    shortterm_qa_turns: int,
    warning_config: Mapping[str, float],
    sessions: Sequence[str] | None = None,
    source_run: Path | None = None,
) -> tuple[Dataset, dict[str, Any]]:
    hard_errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    repository_schema_error: str | None = None
    try:
        # Reuse the production benchmark schema validator before the richer audit.
        load_dataset(dataset_path, include_sheets=None)
    except Exception as exc:
        repository_schema_error = str(exc)
    try:
        dataset = load_benchmark_dataset(dataset_path, sessions)
    except Exception as exc:
        audit = {
            "status": "DATASET_AUDIT_FAILED",
            "dataset": str(dataset_path.resolve()),
            "hard_errors": [*hard_errors, {"code": "DATASET_PARSE_ERROR", "message": str(exc)}],
            "warnings": [],
        }
        _write_audit(output_dir, audit)
        raise DatasetAuditFailed(str(exc)) from exc

    all_turns = [turn for turns in dataset.sessions.values() for turn in turns]
    query_id_counts = Counter(turn.query_id for turn in all_turns)
    duplicate_ids = sorted(query_id for query_id, count in query_id_counts.items() if count > 1)
    if duplicate_ids:
        hard_errors.append({"code": "DUPLICATE_QUERY_ID", "query_ids": duplicate_ids})

    positions = {
        turn.query_id: (turn.session_id, turn.turn_index, turn.session_code)
        for turn in all_turns
        if query_id_counts[turn.query_id] == 1
    }
    requirement_count = 0
    or_count = 0
    shortterm_count = 0
    outside_shortterm_count = 0
    distance_counts: Counter[int] = Counter()
    future: list[dict[str, str]] = []
    cross_session: list[dict[str, str]] = []
    missing: list[dict[str, str]] = []

    for turn in all_turns:
        for requirement in turn.requirements:
            requirement_count += 1
            or_count += int(requirement.is_or)
            short_visible = {
                candidate.query_id
                for candidate in dataset.sessions[turn.session_id][
                    max(0, turn.turn_index - shortterm_qa_turns) : turn.turn_index
                ]
            }
            if any(member in short_visible for member in requirement.members):
                shortterm_count += 1
            else:
                outside_shortterm_count += 1
            for member in requirement.members:
                target = positions.get(member)
                if target is None:
                    missing.append({"query_id": turn.query_id, "target": member})
                    continue
                target_session, target_index, target_code = target
                if target_session != turn.session_id or target_code != turn.session_code:
                    cross_session.append({"query_id": turn.query_id, "target": member})
                elif target_index >= turn.turn_index:
                    future.append({"query_id": turn.query_id, "target": member})
                else:
                    distance_counts[turn.turn_index - target_index] += 1

    if missing:
        hard_errors.append({"code": "MISSING_GOLD_TARGET", "count": len(missing), "examples": missing[:20]})
    if future:
        hard_errors.append({"code": "FUTURE_DEPENDENCY", "count": len(future), "examples": future[:20]})
    if cross_session:
        warnings.append(
            {
                "code": "CROSS_SESSION_DEPENDENCY_EXCLUDED_FROM_SESSION_RECALL",
                "count": len(cross_session),
                "examples": cross_session[:20],
            }
        )
    if repository_schema_error:
        if cross_session and not missing and not future:
            warnings.append(
                {
                    "code": "REPOSITORY_SCHEMA_VALIDATION_CROSS_SESSION_ONLY",
                    "message": repository_schema_error,
                }
            )
        else:
            hard_errors.append({"code": "REPOSITORY_SCHEMA_VALIDATION", "message": repository_schema_error})

    source_check = _source_run_check(
        source_run,
        {turn.query_id for turn in all_turns},
        expected_checkpoint_query_ids={
            turn.query_id for turn in all_turns if turn.requirements or uses_context_gold(turn)
        },
        expected_dataset_sha256=dataset.sha256,
        expected_session_turn_counts={session_id: len(turns) for session_id, turns in dataset.sessions.items()},
    )
    if source_check["status"] == "INCOMPLETE":
        hard_errors.append({"code": "INCOMPLETE_SOURCE_RUN", **source_check})

    fixed_fact_count = sum(
        len(parse_required_context(turn.required_context)) for turn in all_turns if uses_context_gold(turn)
    )
    fixed_context_query_count = sum(uses_context_gold(turn) for turn in all_turns)
    gold_query_count = sum(bool(turn.requirements or uses_context_gold(turn)) for turn in all_turns)
    leakage = [
        turn.query_id
        for turn in all_turns
        if any(pattern.search(turn.question) for pattern in EXPLICIT_HISTORY_PATTERNS)
    ]
    leakage_ratio = len(leakage) / max(gold_query_count, 1)
    if leakage_ratio >= float(warning_config.get("explicit_history_target_leakage_ratio", 0.2)):
        warnings.append(
            {
                "code": "EXPLICIT_HISTORY_POSITION_LEAKAGE",
                "ratio": leakage_ratio,
                "count": len(leakage),
                "examples": leakage[:20],
            }
        )

    total_distances = sum(distance_counts.values())
    dominant_distance, dominant_count = distance_counts.most_common(1)[0] if distance_counts else (None, 0)
    concentration = dominant_count / max(total_distances, 1)
    if concentration >= float(warning_config.get("single_distance_concentration_ratio", 0.6)):
        warnings.append(
            {
                "code": "DEPENDENCY_DISTANCE_CONCENTRATION",
                "distance": dominant_distance,
                "ratio": concentration,
                "count": dominant_count,
            }
        )

    template_counts = Counter(_normalize_template(turn.question) for turn in all_turns if turn.question)
    repeated_template_count = sum(count for count in template_counts.values() if count > 1)
    repeated_template_ratio = repeated_template_count / max(len(all_turns), 1)
    if repeated_template_ratio >= float(warning_config.get("repeated_question_template_ratio", 0.3)):
        warnings.append(
            {"code": "REPEATED_QUERY_TEMPLATE", "ratio": repeated_template_ratio, "count": repeated_template_count}
        )

    audit = {
        "status": "DATASET_AUDIT_FAILED" if hard_errors else "DATASET_QUALITY_WARNING" if warnings else "PASS",
        "dataset": dataset.path,
        "dataset_sha256": dataset.sha256,
        "session_count": len(dataset.sessions),
        "query_count": len(all_turns),
        "gold_bearing_query_count": gold_query_count,
        "independent_query_count": len(all_turns) - gold_query_count,
        "gold_requirement_count": requirement_count + fixed_fact_count,
        "required_context_fact_count": fixed_fact_count,
        "required_context_query_count": fixed_context_query_count,
        "and_requirement_count": requirement_count - or_count,
        "or_requirement_count": or_count,
        "shortterm_requirement_count": shortterm_count,
        "outside_shortterm_requirement_count": outside_shortterm_count,
        "shortterm_qa_turns": shortterm_qa_turns,
        "dependency_distance_distribution": {str(key): value for key, value in sorted(distance_counts.items())},
        "dominant_dependency_distance": dominant_distance,
        "dominant_dependency_distance_ratio": concentration,
        "explicit_history_leakage_ratio": leakage_ratio,
        "future_dependency_count": len(future),
        "cross_session_dependency_count": len(cross_session),
        # Ordinary workbook rows carry only within-session dependency Gold;
        # temporal promotion labels are a separate future contract.
        "cross_session_gold_available": any(
            any(marker in str(turn.dependency_type or "").lower() for marker in ("cross", "temporal", "promotion"))
            for turn in all_turns
        ),
        "missing_gold_target_count": len(missing),
        "duplicate_query_ids": duplicate_ids,
        "source_run": source_check,
        "hard_errors": hard_errors,
        "warnings": warnings,
    }
    _write_audit(output_dir, audit)
    if hard_errors:
        raise DatasetAuditFailed("DATASET_AUDIT_FAILED")
    return dataset, audit


def _write_audit(output_dir: Path, audit: Mapping[str, Any]) -> None:
    atomic_write_json(output_dir / "dataset_audit.json", audit)
    lines = [
        "# Dataset Audit",
        "",
        f"Status: **{audit['status']}**",
        "",
        f"- Dataset: `{audit.get('dataset', '')}`",
        f"- Sessions: {audit.get('session_count', 0)}",
        f"- Queries: {audit.get('query_count', 0)}",
        f"- Gold requirements: {audit.get('gold_requirement_count', 0)}",
        f"- ShortTerm / outside ShortTerm: {audit.get('shortterm_requirement_count', 0)} / "
        f"{audit.get('outside_shortterm_requirement_count', 0)}",
        f"- AND / OR requirements: {audit.get('and_requirement_count', 0)} / {audit.get('or_requirement_count', 0)}",
        "",
        "## Hard errors",
        "",
    ]
    errors = list(audit.get("hard_errors") or [])
    lines.extend(f"- `{item.get('code')}`: {item}" for item in errors)
    if not errors:
        lines.append("- None")
    lines.extend(("", "## Quality warnings", ""))
    warnings = list(audit.get("warnings") or [])
    lines.extend(f"- `{item.get('code')}`: {item}" for item in warnings)
    if not warnings:
        lines.append("- None")
    atomic_write_text(output_dir / "dataset_audit.md", "\n".join(lines) + "\n")
