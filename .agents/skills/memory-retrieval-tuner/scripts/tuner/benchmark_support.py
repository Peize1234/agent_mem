from __future__ import annotations

import asyncio
import json
import os
import re
import shutil
import subprocess
import time
from collections import defaultdict, deque
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

TURN_ID_PATTERN = re.compile(r"S\d{3}-Q\d{3}", re.IGNORECASE)
MIGRATION_TERMINAL_STATUSES = {"succeeded", "succeeded_degraded", "completed_with_loss"}
EMPTY_GOLD_VALUES = {"", "无", "none", "null", "nan"}


@dataclass(frozen=True)
class BenchmarkTurn:
    session_id: str
    sheet_name: str
    turn_index: int
    turn_id: str
    question: str
    answer: str
    needs_history: bool
    dependency_turn_ids: tuple[str, ...]
    required_context: str
    dependency_type: str
    max_lookback: int


@dataclass(frozen=True)
class BenchmarkSession:
    session_id: str
    sheet_name: str
    turns: tuple[BenchmarkTurn, ...]

    @property
    def turn_position_by_id(self) -> dict[str, int]:
        return {turn.turn_id: turn.turn_index for turn in self.turns}


@dataclass(frozen=True)
class GoldRequirement:
    members: tuple[str, ...]
    raw_text: str

    @property
    def is_or(self) -> bool:
        return len(self.members) > 1

    def hit_by(self, retrieved_turn_ids: Iterable[str]) -> bool:
        retrieved = {str(value).upper() for value in retrieved_turn_ids}
        return any(member in retrieved for member in self.members)


def load_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"JSON root must be an object: {path}")
    return value


def expand_env_placeholders(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: expand_env_placeholders(item) for key, item in value.items()}
    if isinstance(value, list):
        return [expand_env_placeholders(item) for item in value]
    if isinstance(value, str):
        match = re.fullmatch(r"\$\{([A-Za-z_][A-Za-z0-9_]*)\}", value)
        if match:
            return os.environ.get(match.group(1), value)
        return os.path.expandvars(value)
    return value


_SENSITIVE_KEYS = {
    "api_key",
    "password",
    "secret",
    "secret_key",
    "private_key",
    "access_key",
    "credentials",
    "credential",
    "token",
    "access_token",
    "refresh_token",
    "auth_token",
    "session_token",
    "client_secret",
}
_SENSITIVE_SUFFIXES = ("_password", "_secret", "_token", "_credential", "_credentials")


def redact_secrets(value: Any) -> Any:
    if isinstance(value, Mapping):
        result: dict[str, Any] = {}
        for key, item in value.items():
            lowered = str(key).lower()
            if lowered in _SENSITIVE_KEYS or any(lowered.endswith(suffix) for suffix in _SENSITIVE_SUFFIXES):
                result[str(key)] = "***REDACTED***" if item not in (None, "") else item
            else:
                result[str(key)] = redact_secrets(item)
        return result
    if isinstance(value, list):
        return [redact_secrets(item) for item in value]
    if isinstance(value, tuple):
        return tuple(redact_secrets(item) for item in value)
    return value


def prepare_runtime_config(
    base_config: Mapping[str, Any],
    *,
    runtime_dir: Path,
    collection_name: str,
    reset_storage: bool,
    agentic_retrieval_enabled: bool,
    profile_enabled: bool,
    profile_update_on_add: bool,
) -> dict[str, Any]:
    runtime_dir = runtime_dir.resolve()
    if reset_storage and runtime_dir.exists():
        shutil.rmtree(runtime_dir)
    runtime_dir.mkdir(parents=True, exist_ok=True)
    config = expand_env_placeholders(deepcopy(dict(base_config)))
    config["history_db_path"] = str(runtime_dir / "history.db")
    config.setdefault("vector_store", {}).setdefault("config", {})
    config["vector_store"]["config"]["path"] = str(runtime_dir / "qdrant")
    config["vector_store"]["config"]["collection_name"] = collection_name
    config.setdefault("agentic_retrieval", {})["enabled"] = bool(agentic_retrieval_enabled)
    config.setdefault("profile", {})["enabled"] = bool(profile_enabled)
    config["profile"]["update_on_add"] = bool(profile_update_on_add)
    config.setdefault("background", {})["enabled"] = True
    return config


def _to_bool(value: Any) -> bool:
    return str(value or "").strip().lower() in {"是", "true", "1", "yes", "y"}


def _to_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(value)) if value not in (None, "") else default
    except (TypeError, ValueError):
        return default


def parse_dependency_turn_ids(value: Any) -> tuple[str, ...]:
    return tuple(dict.fromkeys(match.upper() for match in TURN_ID_PATTERN.findall(str(value or ""))))


def validate_session_turns(sheet_name: str, turns: Sequence[BenchmarkTurn]) -> None:
    positions: dict[str, int] = {}
    for turn in turns:
        if turn.turn_id in positions:
            raise ValueError(f"Sheet {sheet_name} contains duplicate id: {turn.turn_id}")
        positions[turn.turn_id] = turn.turn_index
    for turn in turns:
        for dependency in turn.dependency_turn_ids:
            if dependency not in positions:
                raise ValueError(f"{turn.turn_id} depends on missing turn: {dependency}")
            if positions[dependency] >= turn.turn_index:
                raise ValueError(f"{turn.turn_id} depends on current/future turn: {dependency}")
        if turn.needs_history and not turn.dependency_turn_ids:
            raise ValueError(f"{turn.turn_id} requires history but has no dependency")


def load_dataset(
    path: Path,
    *,
    include_sheets: Sequence[str] | None = None,
    max_sessions: int | None = None,
    max_turns_per_session: int | None = None,
) -> list[BenchmarkSession]:
    """Load benchmark Session sheets while ignoring audit/readme/helper sheets."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    requested = set(include_sheets or [])
    sessions: list[BenchmarkSession] = []
    required = {"编号", "当前问题", "最终回答", "是否需要前文", "关联前序对话"}
    try:
        for sheet_name in workbook.sheetnames:
            if requested and sheet_name not in requested:
                continue
            rows = workbook[sheet_name].iter_rows(values_only=True)
            try:
                headers = [str(value or "").strip() for value in next(rows)]
            except StopIteration:
                continue
            index = {name: position for position, name in enumerate(headers) if name}
            missing = sorted(required - set(index))
            if missing:
                if requested:
                    raise ValueError(f"Sheet {sheet_name} missing columns: {', '.join(missing)}")
                continue
            turns: list[BenchmarkTurn] = []
            for row_number, row in enumerate(rows, start=2):
                if max_turns_per_session is not None and len(turns) >= max_turns_per_session:
                    break
                values = list(row)

                def get(name: str, default: Any = None) -> Any:
                    position = index.get(name)
                    return values[position] if position is not None and position < len(values) else default

                turn_id = str(get("编号") or "").strip().upper()
                question = str(get("当前问题") or "").strip()
                answer = str(get("最终回答") or "").strip()
                if not turn_id and not question and not answer:
                    continue
                if not turn_id or not question or not answer:
                    raise ValueError(f"Sheet {sheet_name} row {row_number} has an incomplete benchmark turn")
                turns.append(
                    BenchmarkTurn(
                        session_id=sheet_name,
                        sheet_name=sheet_name,
                        turn_index=len(turns),
                        turn_id=turn_id,
                        question=question,
                        answer=answer,
                        needs_history=_to_bool(get("是否需要前文")),
                        dependency_turn_ids=parse_dependency_turn_ids(get("关联前序对话")),
                        required_context=str(get("所需前文信息") or "").strip(),
                        dependency_type=str(get("依赖类型") or "").strip(),
                        max_lookback=_to_int(get("最大回溯轮数")),
                    )
                )
            if not turns:
                continue
            if not any(TURN_ID_PATTERN.fullmatch(turn.turn_id) for turn in turns):
                if requested:
                    raise ValueError(f"Sheet {sheet_name} contains no valid Session ids")
                continue
            validate_session_turns(sheet_name, turns)
            sessions.append(BenchmarkSession(sheet_name, sheet_name, tuple(turns)))
            if max_sessions is not None and len(sessions) >= max_sessions:
                break
    finally:
        workbook.close()
    if requested:
        missing = sorted(requested - {session.sheet_name for session in sessions})
        if missing:
            raise ValueError(f"Requested Session sheets not found: {', '.join(missing)}")
    if not sessions:
        raise ValueError(f"Dataset has no benchmark Sessions: {path}")
    return sessions


def _split_top_level(value: str) -> list[str]:
    groups: list[str] = []
    start = 0
    stack: list[str] = []
    opening = {"（": "）", "(": ")"}
    for index, character in enumerate(value):
        if character in opening:
            stack.append(opening[character])
        elif character in {"）", ")"}:
            if not stack or stack.pop() != character:
                raise ValueError(f"Gold expression has mismatched parentheses: {value}")
        elif character in {"；", ";"} and not stack:
            groups.append(value[start:index].strip())
            start = index + 1
    if stack:
        raise ValueError(f"Gold expression has unclosed parentheses: {value}")
    groups.append(value[start:].strip())
    if any(not group for group in groups):
        raise ValueError(f"Gold expression has an empty requirement: {value}")
    return groups


def _strip_outer_parentheses(value: str) -> tuple[str, bool]:
    pairs = {"（": "）", "(": ")"}
    if not value or value[0] not in pairs:
        return value, False
    if value[-1] != pairs[value[0]]:
        raise ValueError(f"Gold OR group has mismatched parentheses: {value}")
    depth = 0
    for index, character in enumerate(value):
        if character in pairs:
            depth += 1
        elif character in {"）", ")"}:
            depth -= 1
            if depth == 0 and index != len(value) - 1:
                return value, False
    return value[1:-1].strip(), True


def parse_gold_requirements(value: Any) -> tuple[GoldRequirement, ...]:
    text = str(value or "").strip()
    if text.lower() in EMPTY_GOLD_VALUES:
        return ()
    requirements: list[GoldRequirement] = []
    for raw_group in _split_top_level(text):
        inner, parenthesized = _strip_outer_parentheses(raw_group)
        member_texts = re.split(r"[；;]", inner) if parenthesized else [inner]
        members: list[str] = []
        for member_text in member_texts:
            matches = [match.upper() for match in TURN_ID_PATTERN.findall(member_text)]
            if len(matches) != 1:
                raise ValueError(f"Gold member must contain exactly one turn id: {member_text!r}")
            if matches[0] not in members:
                members.append(matches[0])
        if parenthesized and len(members) < 2:
            raise ValueError(f"Gold OR group must contain at least two members: {raw_group}")
        requirements.append(GoldRequirement(tuple(members), raw_group))
    return tuple(requirements)


class LineageTracker:
    def __init__(self, short_term_qa_capacity: int):
        if short_term_qa_capacity < 0:
            raise ValueError("short_term_qa_capacity cannot be negative")
        self.short_term_qa_capacity = short_term_qa_capacity
        self._active_turns: dict[str, deque[str]] = defaultdict(deque)
        self.job_to_turn_ids: dict[str, tuple[str, ...]] = {}

    def register_add(self, session_id: str, turn_id: str, migration_job_id: str | None) -> tuple[str, ...]:
        active = self._active_turns[session_id]
        active.append(turn_id)
        overflow: list[str] = []
        while len(active) > self.short_term_qa_capacity:
            overflow.append(active.popleft())
        if overflow and not migration_job_id:
            raise RuntimeError(f"Session {session_id} evicted {overflow} without a migration job")
        if migration_job_id:
            self.job_to_turn_ids[migration_job_id] = tuple(overflow)
        return tuple(overflow)

    def turn_ids_for_job(self, job_id: str | None) -> tuple[str, ...]:
        return self.job_to_turn_ids.get(job_id or "", ())


def ordered_unique(values: Iterable[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


async def wait_for_migration_jobs(
    memory: Any,
    job_ids: Sequence[str],
    *,
    timeout_seconds: float,
    poll_interval_seconds: float,
) -> list[dict[str, Any]]:
    async def wait_one(job_id: str) -> dict[str, Any]:
        deadline = time.monotonic() + timeout_seconds
        while True:
            job = await asyncio.to_thread(memory.db.get_background_job, job_id, "migration")
            if job is None:
                raise RuntimeError(f"Migration job not found: {job_id}")
            if str(job.get("status")) in MIGRATION_TERMINAL_STATUSES:
                return job
            if time.monotonic() >= deadline:
                raise TimeoutError(f"Migration job timed out: {job_id} ({job.get('status')})")
            await asyncio.sleep(poll_interval_seconds)

    return await asyncio.gather(*(wait_one(job_id) for job_id in ordered_unique(job_ids)))


def safe_git_commit(repo_root: Path) -> str | None:
    try:
        completed = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=repo_root,
            capture_output=True,
            text=True,
            timeout=2,
            check=False,
        )
        return completed.stdout.strip() or None
    except (OSError, subprocess.SubprocessError):
        return None
