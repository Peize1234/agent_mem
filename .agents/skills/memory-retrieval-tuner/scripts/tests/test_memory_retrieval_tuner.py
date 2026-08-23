from __future__ import annotations

import hashlib
import json
import sys
import threading
import time
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

SCRIPTS = Path(__file__).resolve().parents[1]
REPO_ROOT = next(path for path in Path(__file__).resolve().parents if (path / "pyproject.toml").exists())
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(REPO_ROOT))

import tuner.orchestrator as orchestrator  # noqa: E402
import tuner.production_midterm_adapter as production_adapter  # noqa: E402
from tuner.agentic_retrieval_artifacts import (  # noqa: E402
    PRODUCTION_AGENTIC_EXECUTION_CONTRACT,
    PRODUCTION_AGENTIC_TRACE_SCHEMA,
    agentic_parent_retrieval_identity_sha256,
    build_agentic_parent_retrieval_identity,
    load_production_agentic_trace,
)
from tuner.artifact_registry import ArtifactRegistry  # noqa: E402
from tuner.benchmark_support import load_dataset, parse_gold_requirements  # noqa: E402
from tuner.build_report import write_outputs  # noqa: E402
from tuner.candidate_selector import classify_overfit, select_best  # noqa: E402
from tuner.dataset_audit import DatasetAuditFailed, audit_dataset  # noqa: E402
from tuner.derived_artifacts import DerivedArtifactBuilder  # noqa: E402
from tuner.diagnostic_midterm_retriever import DiagnosticMidTermRetriever  # noqa: E402
from tuner.encoding_contract import EncodingContract  # noqa: E402
from tuner.evaluate_candidate import (  # noqa: E402
    _eligible_requirements,
    _evaluate_session,
    candidate_hash,
    evaluate_candidate,
)
from tuner.experiment_branches import (  # noqa: E402
    ALL_REGIMES,
    AgenticRetrievalBranch,
    BranchContext,
    BranchOutcome,
    BranchRegistry,
    BranchSpec,
    EmbeddingBranch,
    FineGrainedLongtermExtractionPromptBranch,
    QueryRepresentationBranch,
    RerankingBranch,
    RetrievalControlBranch,
    SourcePromptBranch,
    _derived_dimensions,
)
from tuner.generated_source_artifacts import prepare_generated_source_candidate  # noqa: E402
from tuner.io_utils import sha256_file  # noqa: E402
from tuner.model_discovery import (  # noqa: E402
    ModelCandidate,
    ModelDiscovery,
    ResourceEnvelope,
    _annotate_relative_benchmark_scores,
    _candidate_score_value,
    _metadata_evidence,
)
from tuner.models import Candidate, CandidateResult, Dataset, Requirement, Turn  # noqa: E402
from tuner.orchestrator import (  # noqa: E402
    TunerConfig,
    _artifact_cache_root,
    _evaluate_loso_many,
    _resolve_memory_config,
    _resolve_run_memory_config,
    _resolve_shortterm_window,
)
from tuner.parameter_schema import parameter_class  # noqa: E402
from tuner.production_midterm_adapter import (  # noqa: E402
    ADAPTER_SCHEMA,
    ProductionMidtermAdapter,
    generate_production_sources,
    isolated_runtime_layout,
    production_candidate_from_manifests,
    production_prompt_hashes,
)
from tuner.prompt_artifacts import (  # noqa: E402
    PRODUCTION_QUERY_PROMPT_HASH,
    PRODUCTION_SHORTTERM_HISTORY_POLICY,
    QueryPromptArtifactGenerator,
    QueryPromptVariant,
    controlled_query_prompt_variants,
)
from tuner.source_prompt_variants import (  # noqa: E402
    controlled_fine_grained_longterm_prompt_variants,
    controlled_page_prompt_variants,
)
from tuner.split_sessions import create_or_load_split  # noqa: E402
from tuner.staged_search import candidate_config_hash, run_staged_search  # noqa: E402

from mem0.configs.base import AgenticRetrievalConfig, MemoryConfig, MidTermMemoryConfig  # noqa: E402
from mem0.configs.midterm_prompts import MIDTERM_PAGE_SUMMARY_PROMPT  # noqa: E402
from mem0.configs.production import load_production_memory_config  # noqa: E402
from mem0.configs.query_prompts import QUERY_REFERENCE_RESOLUTION_PROMPT  # noqa: E402
from mem0.embeddings.encoding_contract import EncodingContractEmbedding  # noqa: E402
from mem0.memory import main as memory_main  # noqa: E402
from mem0.memory.main import Memory  # noqa: E402
from mem0.memory.midterm_retriever import MidTermRetriever  # noqa: E402
from mem0.memory.midterm_updater import PRODUCTION_PAGE_CONTEXT_CONTRACT  # noqa: E402
from mem0.memory.query_resolver import QueryResolver as ProductionQueryResolver  # noqa: E402
from mem0.memory.query_resolver import build_query_resolution_messages  # noqa: E402
from mem0.memory.storage import SQLiteManager  # noqa: E402


def make_turn(session: str, index: int, gold: tuple[Requirement, ...] = ()) -> Turn:
    code = session.split("_", 1)[0]
    return Turn(
        session_id=session,
        session_code=code,
        turn_index=index,
        query_id=f"{code}-Q{index + 1:03d}",
        question=f"问题 {index + 1} 指标 {index % 2}",
        answer=f"答案 {index + 1} 财务 指标 {index % 2}",
        gold_raw="",
        requirements=gold,
    )


def make_dataset(tmp_path: Path, session_count: int = 2) -> Dataset:
    sessions = {}
    for session_index in range(session_count):
        session = f"S{session_index + 1:03d}_测试"
        code = session.split("_", 1)[0]
        turns = [make_turn(session, index) for index in range(4)]
        turns.append(make_turn(session, 4, (Requirement((f"{code}-Q001",), f"{code}-Q001"),)))
        sessions[session] = tuple(turns)
    return Dataset(path=str(tmp_path / "synthetic.xlsx"), sha256="a" * 64, sessions=sessions)


def result(name: str, recall: float, *, macro: float | None = None, stddev: float = 0.0) -> CandidateResult:
    return CandidateResult(
        name=name,
        candidate_hash=name,
        stage="test",
        config={"name": name},
        metrics={
            "recall_at_k": recall,
            "macro_session_recall_at_k": recall if macro is None else macro,
            "session_stddev": stddev,
            "mrr": recall,
            "recall_at_2k": recall,
            "recall_at_4k": recall,
        },
        requirement_rows=[],
        session_rows=[],
        runtime_seconds=0.1,
        work_seconds=0.1,
        cache_hits=0,
        cache_misses=1,
    )


def write_workbook(path: Path, sessions: dict[str, list[tuple[str, str]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sessions.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["编号", "当前问题", "最终回答", "是否需要前文", "关联前序对话", "所需前文信息", "依赖类型"])
        for query_id, gold in rows:
            sheet.append([query_id, f"问题 {query_id}", f"答案 {query_id}", bool(gold), gold, "", ""])
    workbook.save(path)


def test_gold_and_or_semantics() -> None:
    requirements = parse_gold_requirements("S001-Q001；（S001-Q002；S001-Q003）")
    assert len(requirements) == 2
    assert not requirements[0].is_or
    assert requirements[1].is_or
    assert requirements[1].hit_by(["S001-Q003"])


@pytest.mark.parametrize(("k", "expected"), [(5, False), (10, True)])
def test_evaluator_parameterizes_k(tmp_path: Path, k: int, expected: bool) -> None:
    dataset = make_dataset(tmp_path, 1)
    session = next(iter(dataset.sessions))
    target = dataset.sessions[session][-1]
    ranking = [
        {"source_turn_id": f"S001-Q{index:03d}", "page_id": f"S001-Q{index:03d}"} for index in (2, 3, 4, 8, 9, 10, 1)
    ]
    evaluated = _evaluate_session(
        dataset, session, {target.query_id: ranking}, k=k, target="midterm", shortterm_window=3
    )
    assert evaluated["requirements"][0]["hit_at_k"] is expected


@pytest.mark.parametrize(
    ("count", "method", "validation_count"),
    [
        (8, "deterministic_stratified_holdout", 2),
        (6, "deterministic_holdout", 2),
        (4, "leave_one_session_out", 0),
        (2, "exploratory_only", 0),
    ],
)
def test_session_split_policy(tmp_path: Path, count: int, method: str, validation_count: int) -> None:
    dataset = make_dataset(tmp_path, count)
    output = tmp_path / f"split-{count}"
    split = create_or_load_split(dataset, output_dir=output, seed=42, shortterm_window=3)
    assert split["method"] == method
    assert len(split["validation_sessions"]) == validation_count
    assert not set(split["tune_sessions"]) & set(split["validation_sessions"])
    if method == "leave_one_session_out":
        assert len(split["folds"]) == count
        assert sorted(fold["validation_sessions"][0] for fold in split["folds"]) == sorted(dataset.sessions)
        assert all(len(fold["tune_sessions"]) == count - 1 for fold in split["folds"])
    assert create_or_load_split(dataset, output_dir=output, seed=42, shortterm_window=3) == split


def test_shortterm_window_is_derived_and_mismatch_is_detected() -> None:
    actual, validation = _resolve_shortterm_window(
        {"dataset": {"shortterm_qa_turns": 2}},
        {"midterm": {"short_term_capacity": 6}},
    )
    assert actual == 3
    assert validation == {
        "status": "OVERRIDDEN_BY_PRODUCTION_CONFIG",
        "dataset_config_qa_turns": 2,
        "production_capacity_messages": 6,
        "actual_qa_turns": 3,
        "source": "memory_config.midterm.short_term_capacity / 2",
    }
    with pytest.raises(ValueError, match="positive even message count"):
        _resolve_shortterm_window(
            {"dataset": {"shortterm_qa_turns": 3}},
            {"midterm": {"short_term_capacity": 5}},
        )


def test_default_memory_config_is_resolved_from_current_production_defaults(tmp_path: Path) -> None:
    resolved, source = _resolve_memory_config(None)
    production = load_production_memory_config(resolve_environment=False).model_dump(mode="json", warnings=False)

    assert source == "repository_production_config"
    assert TunerConfig(dataset=tmp_path / "dataset.xlsx").memory_config is None
    assert not (SCRIPTS.parent / "memory_config.json").exists()
    assert resolved == production
    assert resolved["llm"]["provider"] == "deepseek"
    assert resolved["llm"]["config"]["model"] == "deepseek-v4-flash"
    assert resolved["embedder"]["provider"] == "huggingface"
    assert resolved["embedder"]["config"]["model"] == "BAAI/bge-small-zh-v1.5"
    assert resolved["embedder"]["config"]["embedding_dims"] == 512
    assert resolved["agentic_retrieval"]["max_tool_result_chars"] == 30000
    assert resolved["agentic_retrieval"]["max_tool_result_chars"] == AgenticRetrievalConfig().max_tool_result_chars


def test_partial_explicit_memory_config_overrides_only_declared_values(tmp_path: Path) -> None:
    config_path = tmp_path / "partial-memory-config.json"
    config_path.write_text(
        json.dumps(
            {
                "midterm": {"top_k_pages": 9},
                "benchmark_runtime": {"llm_observability": True},
            }
        ),
        encoding="utf-8",
    )

    resolved, source = _resolve_memory_config(config_path)

    assert source == "repository_production_config_with_explicit_override"
    assert resolved["midterm"]["top_k_pages"] == 9
    assert resolved["llm"]["provider"] == "deepseek"
    assert resolved["llm"]["config"]["model"] == "deepseek-v4-flash"
    assert resolved["embedder"]["provider"] == "huggingface"
    assert resolved["embedder"]["config"]["model"] == "BAAI/bge-small-zh-v1.5"
    assert resolved["embedder"]["config"]["embedding_dims"] == 512
    assert resolved["agentic_retrieval"]["max_tool_result_chars"] == 30000
    assert resolved["midterm"]["top_k_sessions"] == MidTermMemoryConfig().top_k_sessions
    assert resolved["agentic_retrieval"]["max_iterations"] == AgenticRetrievalConfig().max_iterations
    assert resolved["agentic_retrieval"]["max_tool_calls"] == AgenticRetrievalConfig().max_tool_calls
    assert resolved["benchmark_runtime"] == {"llm_observability": True}


def test_legacy_non_thinking_flags_become_production_request_options(tmp_path: Path) -> None:
    config_path = tmp_path / "non-thinking.json"
    config_path.write_text(
        json.dumps(
            {
                "benchmark_runtime": {
                    "deepseek_midterm_non_thinking": True,
                    "deepseek_longterm_non_thinking": True,
                }
            }
        ),
        encoding="utf-8",
    )

    resolved, _ = _resolve_memory_config(config_path)
    disabled = {"extra_body": {"thinking": {"type": "disabled"}}}
    assert resolved["midterm"]["page_summary_request_options"] == disabled
    assert resolved["midterm"]["session_merge_request_options"] == disabled
    assert resolved["fine_grained_longterm"]["extraction_request_options"] == disabled


def test_resume_uses_frozen_config_after_repository_production_changes(tmp_path: Path, monkeypatch) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    config_a = load_production_memory_config(
        {"midterm": {"top_k_pages": 7}},
        resolve_environment=False,
    )
    config_b = load_production_memory_config(
        {"midterm": {"top_k_pages": 11}},
        resolve_environment=False,
    )
    monkeypatch.setattr(orchestrator, "load_production_memory_config", lambda *args, **kwargs: config_a)

    frozen, _source, frozen_path, frozen_sha256 = _resolve_run_memory_config(
        TunerConfig(dataset=tmp_path / "dataset.xlsx"),
        run_dir,
        {},
    )
    monkeypatch.setattr(orchestrator, "load_production_memory_config", lambda *args, **kwargs: config_b)

    resumed, source, resumed_path, resumed_sha256 = _resolve_run_memory_config(
        TunerConfig(dataset=tmp_path / "dataset.xlsx", resume=run_dir),
        run_dir,
        {"resolved_memory_config_sha256": frozen_sha256},
    )

    assert resumed == frozen
    assert resumed["midterm"]["top_k_pages"] == 7
    assert source == "frozen_run_config"
    assert resumed_path == frozen_path
    assert resumed_sha256 == frozen_sha256 == sha256_file(frozen_path)


def test_resume_accepts_matching_explicit_config_and_rejects_conflict(tmp_path: Path) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    matching_path = tmp_path / "matching.json"
    matching_path.write_text(json.dumps({"midterm": {"top_k_pages": 8}}), encoding="utf-8")
    conflict_path = tmp_path / "conflict.json"
    conflict_path.write_text(json.dumps({"midterm": {"top_k_pages": 9}}), encoding="utf-8")

    frozen, _source, _path, frozen_sha256 = _resolve_run_memory_config(
        TunerConfig(dataset=tmp_path / "dataset.xlsx", memory_config=matching_path),
        run_dir,
        {},
    )
    resumed, source, _path, _sha256 = _resolve_run_memory_config(
        TunerConfig(dataset=tmp_path / "dataset.xlsx", resume=run_dir, memory_config=matching_path),
        run_dir,
        {"resolved_memory_config_sha256": frozen_sha256},
    )

    assert resumed == frozen
    assert source == "frozen_run_config_verified_explicit"
    with pytest.raises(ValueError, match="Resume memory_config conflicts with the frozen configuration"):
        _resolve_run_memory_config(
            TunerConfig(dataset=tmp_path / "dataset.xlsx", resume=run_dir, memory_config=conflict_path),
            run_dir,
            {"resolved_memory_config_sha256": frozen_sha256},
        )


def test_resume_legacy_run_without_frozen_config_fails_closed(tmp_path: Path) -> None:
    run_dir = tmp_path / "legacy-run"
    run_dir.mkdir()

    with pytest.raises(ValueError, match="legacy tuning run.*resolved_memory_config.json is missing"):
        _resolve_run_memory_config(
            TunerConfig(dataset=tmp_path / "dataset.xlsx", resume=run_dir),
            run_dir,
            {},
        )


def test_resume_derives_cache_from_run_directory(tmp_path: Path) -> None:
    run_dir = tmp_path / "custom-output" / "run-id"
    resumed = TunerConfig(
        dataset=tmp_path / "dataset.xlsx",
        output_root=tmp_path / "unrelated-default",
        resume=run_dir,
    )
    fresh = TunerConfig(dataset=tmp_path / "dataset.xlsx", output_root=tmp_path / "fresh-output")
    assert _artifact_cache_root(resumed, run_dir) == tmp_path / "custom-output" / ".cache"
    assert (
        _artifact_cache_root(fresh, tmp_path / "fresh-output" / "run")
        == (tmp_path / "fresh-output" / ".cache").resolve()
    )


def test_loso_executes_every_frozen_fold(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 3)
    split = create_or_load_split(dataset, output_dir=tmp_path / "split", seed=42, shortterm_window=3)
    ranking_path = tmp_path / "rankings.jsonl"
    ranking_path.write_text(
        "\n".join(
            json.dumps(
                {
                    "query_id": turns[-1].query_id,
                    "source_turn_id": turns[0].query_id,
                    "rank": 1,
                    "score": 1.0,
                }
            )
            for turns in dataset.sessions.values()
        )
        + "\n",
        encoding="utf-8",
    )
    candidate = Candidate(
        name="baseline",
        stage="baseline",
        config={"backend": "frozen_ranking", "ranking_path": str(ranking_path)},
    )
    registry = ArtifactRegistry(tmp_path / "cache", tmp_path / "results")
    trace_path = tmp_path / "search_trace.jsonl"
    trace_path.touch()
    common = {
        "dataset": dataset,
        "candidates": [candidate],
        "folds": split["folds"],
        "config": TunerConfig(dataset=Path(dataset.path), k=5),
        "shortterm_window": 3,
        "ranking_depth": 20,
        "registry": registry,
        "run_dir": tmp_path / "run",
        "execution": {"max_parallel_candidates": 2, "max_parallel_sessions": 2, "adaptive_reductions": []},
        "trace_path": trace_path,
    }
    tune = _evaluate_loso_many(
        partition="tune_sessions",
        scope="tune_loso",
        **common,
    )[0]
    validation = _evaluate_loso_many(
        partition="validation_sessions",
        scope="validation_loso",
        **common,
    )[0]

    assert len(tune.session_rows) == 6
    assert len(validation.session_rows) == 3
    assert sorted(row["session_id"] for row in validation.session_rows) == sorted(dataset.sessions)
    scopes = [json.loads(line)["scope"] for line in trace_path.read_text(encoding="utf-8").splitlines()]
    assert {scope for scope in scopes if scope.startswith("validation_loso")} == {
        "validation_loso_fold_1",
        "validation_loso_fold_2",
        "validation_loso_fold_3",
    }


def test_loso_prunes_invalid_candidate_without_aborting(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 3)
    split = create_or_load_split(dataset, output_dir=tmp_path / "split-invalid", seed=42, shortterm_window=3)
    candidate = Candidate(name="invalid", stage="test", config={"backend": "not-supported"})
    trace = tmp_path / "invalid-trace.jsonl"
    trace.touch()
    combined = _evaluate_loso_many(
        dataset,
        [candidate],
        split["folds"],
        partition="tune_sessions",
        scope="invalid_loso",
        config=TunerConfig(dataset=Path(dataset.path), k=5),
        shortterm_window=3,
        ranking_depth=20,
        registry=ArtifactRegistry(tmp_path / "invalid-cache", tmp_path / "results"),
        run_dir=tmp_path / "invalid-run",
        execution={"max_parallel_candidates": 1, "max_parallel_sessions": 2, "adaptive_reductions": []},
        trace_path=trace,
    )
    assert len(combined) == 1
    assert combined[0].status == "INVALID"
    assert combined[0].metrics["recall_at_k"] == 0.0


def test_audit_future_dependency_and_cross_session(tmp_path: Path) -> None:
    future = tmp_path / "future.xlsx"
    write_workbook(future, {"S001_x": [("S001-Q001", "S001-Q002"), ("S001-Q002", "")]})
    with pytest.raises(DatasetAuditFailed):
        audit_dataset(
            future,
            output_dir=tmp_path / "future-output",
            shortterm_qa_turns=3,
            warning_config={},
        )
    audit = json.loads((tmp_path / "future-output/dataset_audit.json").read_text(encoding="utf-8"))
    assert audit["status"] == "DATASET_AUDIT_FAILED"
    assert any(error["code"] in {"FUTURE_DEPENDENCY", "REPOSITORY_SCHEMA_VALIDATION"} for error in audit["hard_errors"])

    cross = tmp_path / "cross.xlsx"
    write_workbook(
        cross,
        {
            "S001_x": [("S001-Q001", ""), ("S001-Q002", "S002-Q001")],
            "S002_x": [("S002-Q001", "")],
        },
    )
    audit_dataset(
        cross,
        output_dir=tmp_path / "cross-output",
        shortterm_qa_turns=3,
        warning_config={},
    )
    audit = json.loads((tmp_path / "cross-output/dataset_audit.json").read_text(encoding="utf-8"))
    assert audit["cross_session_dependency_count"] == 1
    assert audit["cross_session_gold_available"] is False
    assert not any(error["code"] == "CROSS_SESSION_DEPENDENCY" for error in audit["hard_errors"])


def test_auxiliary_sheet_is_not_a_session(tmp_path: Path) -> None:
    dataset_path = tmp_path / "with-helper.xlsx"
    write_workbook(dataset_path, {"S001_x": [("S001-Q001", ""), ("S001-Q002", "S001-Q001")]})
    workbook = load_workbook(dataset_path)
    helper = workbook.create_sheet("Dataset_Quality_Audit")
    helper.append(["编号", "当前问题", "最终回答", "是否需要前文", "关联前序对话"])
    helper.append(["CHECK-1", "检查项", "这不是 Session", False, ""])
    workbook.save(dataset_path)
    workbook.close()

    repository_sessions = load_dataset(dataset_path)
    assert [session.session_id for session in repository_sessions] == ["S001_x"]
    dataset, audit = audit_dataset(
        dataset_path,
        output_dir=tmp_path / "audit-helper",
        shortterm_qa_turns=3,
        warning_config={},
    )
    assert list(dataset.sessions) == ["S001_x"]
    assert audit["session_count"] == 1


def test_audit_accepts_complete_production_midterm_checkpoint_run(tmp_path: Path) -> None:
    dataset_path = tmp_path / "dataset.xlsx"
    write_workbook(
        dataset_path,
        {
            "S001_x": [
                ("S001-Q001", ""),
                ("S001-Q002", ""),
                ("S001-Q003", ""),
                ("S001-Q004", ""),
                ("S001-Q005", "S001-Q001"),
            ]
        },
    )
    source = tmp_path / "source/S001"
    source.mkdir(parents=True)
    checkpoints = source / "production_midterm_checkpoints.jsonl"
    checkpoints.write_text(json.dumps({"query_id": "S001-Q005"}) + "\n", encoding="utf-8")
    (source / "production_midterm_manifest.json").write_text(
        json.dumps(
            {
                "status": "COMPLETE",
                "dataset_sha256": sha256_file(dataset_path),
                "session_id": "S001_x",
                "turn_count": 5,
                "failed_turns": 0,
                "checkpoints_path": str(checkpoints),
                "checkpoints_sha256": sha256_file(checkpoints),
            }
        ),
        encoding="utf-8",
    )
    _, audit = audit_dataset(
        dataset_path,
        output_dir=tmp_path / "audit",
        shortterm_qa_turns=3,
        warning_config={
            "explicit_history_target_leakage_ratio": 1.1,
            "single_distance_concentration_ratio": 1.1,
            "repeated_question_template_ratio": 1.1,
        },
        source_run=source.parent,
    )
    assert audit["source_run"]["status"] == "COMPLETE"
    assert audit["source_run"]["kind"] == "production_midterm_checkpoints"


def test_true_midterm_requirement_eligibility_for_or_gold(tmp_path: Path) -> None:
    session = "S001_test"
    turns = [make_turn(session, index) for index in range(5)]
    target = make_turn(
        session,
        5,
        (
            Requirement(("S001-Q001",), "S001-Q001"),
            Requirement(("S001-Q001", "S001-Q004"), "(S001-Q001；S001-Q004)"),
        ),
    )
    turns.append(target)
    eligible = _eligible_requirements(target, turns, "midterm", shortterm_window=3)
    assert eligible == [Requirement(("S001-Q001",), "S001-Q001")]


def test_production_adapter_manifest_and_runtime_isolation(tmp_path: Path) -> None:
    manifests = []
    for index, session in enumerate(("S001_test", "S002_test"), start=1):
        checkpoints = tmp_path / session / "checkpoints.jsonl"
        checkpoints.parent.mkdir(parents=True)
        checkpoints.write_text("", encoding="utf-8")
        manifest = checkpoints.parent / "production_midterm_manifest.json"
        manifest.write_text(
            json.dumps(
                {
                    "schema": ADAPTER_SCHEMA,
                    "status": "COMPLETE",
                    "backend": "production_midterm",
                    "session_id": session,
                    "production_config": {
                        "top_k_sessions": 5,
                        "top_k_pages": 5,
                        "max_total_pages": 5,
                    },
                    "effective_memory_config": {
                        "vector_store": {"config": {"bm25_language": "zh"}},
                        "agentic_retrieval": {
                            "max_iterations": 2,
                            "max_tool_calls": 1,
                            "max_queries": 3,
                            "max_total_results": 6,
                            "max_tool_result_chars": 23456,
                        },
                    },
                    "checkpoints_path": str(checkpoints),
                    "failed_turns": 0,
                    "llm_calls": index,
                    "embedding_calls": index * 10,
                }
            ),
            encoding="utf-8",
        )
        manifests.append(manifest)
    config, provenance = production_candidate_from_manifests(manifests)
    assert config["backend"] == "production_midterm"
    assert config["retrieval_method"] == "dense"
    assert config["bm25_language"] == "zh"
    assert config["max_total_results"] == 6
    assert config["benchmark_constraints"]["agentic_result_cap"] == 5
    assert config["agentic_fixed_max_tool_result_chars"] == 23456
    assert provenance["source"] == "real AsyncMemory Add/MidTerm pipeline"
    assert provenance["llm_calls"] == 3
    assert provenance["embedding_calls"] == 30
    assert provenance["production_agentic_max_total_results"] == 6
    assert provenance["production_agentic_max_tool_result_chars"] == 23456
    assert provenance["tuner_agentic_context_cap"] == 5
    ProductionMidtermAdapter.supported(config)
    with pytest.raises(ValueError, match="regenerated production source artifacts"):
        ProductionMidtermAdapter._validated_production_config({**config, "page_representation": "summary"})
    with pytest.raises(ValueError, match="regenerated production source artifacts"):
        ProductionMidtermAdapter._validated_production_config(
            {
                **config,
                "production_overrides": {
                    "fine_grained_longterm": {
                        "extraction_request_options": {"extra_body": {"thinking": {"type": "disabled"}}}
                    }
                },
            }
        )
    with pytest.raises(ValueError, match="Production QueryResolver artifact"):
        ProductionMidtermAdapter._validated_production_config(
            {**config, "production_overrides": {"query_rewrite_prompt": "custom query prompt"}}
        )

    first = isolated_runtime_layout(tmp_path, candidate_hash="candidate-a", session_id="S001")
    second = isolated_runtime_layout(tmp_path, candidate_hash="candidate-b", session_id="S001")
    third = isolated_runtime_layout(tmp_path, candidate_hash="candidate-a", session_id="S002")
    assert len({first.root, second.root, third.root}) == 3
    assert all(layout.sqlite_path.exists() for layout in (first, second, third))


def test_production_candidate_inherits_deployed_midterm_behavior(tmp_path: Path) -> None:
    checkpoints = tmp_path / "checkpoints.jsonl"
    checkpoints.write_text("", encoding="utf-8")
    manifest = tmp_path / "production_midterm_manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "schema": ADAPTER_SCHEMA,
                "status": "COMPLETE",
                "production_config": {
                    "retrieval_method": "dense_bm25_fusion",
                    "page_representation": "summary_keywords",
                    "dense_weight": 0.35,
                },
                "effective_memory_config": {},
                "checkpoints_path": str(checkpoints),
            }
        ),
        encoding="utf-8",
    )

    config, _ = production_candidate_from_manifests([manifest])

    assert config["retrieval_method"] == "dense_bm25_fusion"
    assert config["page_representation"] == "summary_keywords"
    assert config["dense_weight"] == pytest.approx(0.35)


def test_production_source_workers_respect_llm_concurrency_limit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    dataset_path = tmp_path / "dataset.xlsx"
    dataset_path.write_bytes(b"dataset")
    memory_config = tmp_path / "memory_config.json"
    memory_config.write_text("{}", encoding="utf-8")
    lock = threading.Lock()
    active = 0
    max_active = 0

    def fake_run(command: list[str], **_: object) -> object:
        nonlocal active, max_active
        with lock:
            active += 1
            max_active = max(max_active, active)
        try:
            time.sleep(0.05)
            spec = json.loads(Path(command[-1]).read_text(encoding="utf-8"))
            output_dir = Path(spec["output_dir"])
            checkpoints = output_dir / "production_midterm_checkpoints.jsonl"
            checkpoints.write_text("", encoding="utf-8")
            (output_dir / "production_midterm_manifest.json").write_text(
                json.dumps(
                    {
                        "status": "COMPLETE",
                        "turn_count": 1,
                        "failed_turns": 0,
                        "checkpoints_path": str(checkpoints),
                    }
                ),
                encoding="utf-8",
            )
        finally:
            with lock:
                active -= 1
        return production_adapter.subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(production_adapter.subprocess, "run", fake_run)
    session_ids = [f"S{index:03d}_test" for index in range(1, 5)]
    paths = generate_production_sources(
        dataset_path=dataset_path,
        dataset_sha256="dataset-sha",
        session_ids=session_ids,
        session_turn_counts={session_id: 1 for session_id in session_ids},
        memory_config_path=memory_config,
        run_dir=tmp_path / "run",
        ranking_depth=20,
        llm_mode="real",
        max_parallel_sessions=4,
        max_parallel_llm_calls=2,
    )
    assert len(paths) == 4
    assert 1 < max_active <= 2


def test_production_adapter_calls_real_midterm_retriever(tmp_path: Path) -> None:
    session_point_id = "11111111-1111-1111-1111-111111111111"
    page_point_id = "22222222-2222-2222-2222-222222222222"
    scope = {"user_id": "recall::S001_test", "run_id": "S001_test"}
    checkpoint = {
        "query_id": "S001-Q005",
        "query": "真实生产检索查询",
        "filters": scope,
        "query_vector": [1.0, 0.0],
        "sessions": [
            {
                "id": session_point_id,
                "vector": [1.0, 0.0],
                "payload": {
                    **scope,
                    "summary": "生产 Session 摘要",
                    "summary_keywords": ["生产"],
                    "data": "生产 Session 摘要\nKeywords: 生产",
                    "output_state": "committed",
                    "page_ids": [page_point_id],
                },
            }
        ],
        "pages": [
            {
                "id": page_point_id,
                "vector": [1.0, 0.0],
                "payload": {
                    **scope,
                    "session_id": session_point_id,
                    "source_job_id": "job-1",
                    "summary": "真实生产检索查询的 Page prompt 摘要",
                    "keywords": ["生产", "摘要"],
                    "user_input": "源问题",
                    "data": "真实生产检索查询的 Page prompt 摘要\nKeywords: 生产, 摘要\nUser: 源问题",
                    "text_lemmatized": "真实生产检索查询 生产 page 摘要",
                    "output_state": "committed",
                },
            }
        ],
        "source_turn_ids_by_job": {"job-1": ["S001-Q001"]},
    }
    config = {
        "backend": "production_midterm",
        "retrieval_method": "dense",
        "query_representation": "original",
        "page_representation": "production",
        "top_k_sessions": 1,
        "top_k_pages": 1,
        "max_total_pages": 1,
    }
    adapter = ProductionMidtermAdapter(
        run_dir=tmp_path,
        candidate_hash="candidate",
        session_id="S001_test",
        ranking_depth=20,
    )
    result = adapter.rank(checkpoint, config)
    assert len(result) == 1
    assert result[0]["page_id"] == page_point_id
    assert result[0]["source_turn_id"] == "S001-Q001"
    assert result[0]["source"] == "mid_term_page"
    assert result[0]["score"] == pytest.approx(1.0)
    assert result[0]["rank"] == 1
    assert result[0]["candidate_pool_count"] == 1
    assert result[0]["final_rank"] == 1

    hybrid = ProductionMidtermAdapter(
        run_dir=tmp_path,
        candidate_hash="candidate-hybrid",
        session_id="S001_test",
        ranking_depth=20,
    )
    hybrid_ranking = hybrid.rank(
        checkpoint,
        {
            **config,
            "retrieval_method": "dense_bm25_fusion",
            "dense_weight": 0.7,
            "bm25_language": "zh",
        },
    )
    assert hybrid_ranking[0]["source_turn_id"] == "S001-Q001"
    assert hybrid_ranking[0]["source"] == "mid_term_page"


def test_diagnostic_retriever_matches_production_and_keeps_trace_isolated() -> None:
    scope = {"user_id": "u1", "run_id": "r1"}
    session = SimpleNamespace(
        id="s1",
        score=0.9,
        payload={
            **scope,
            "summary": "session",
            "page_ids": ["p1", "p2", "p3"],
            "last_visit_turn_index": 1,
            "N_visit": 1,
            "L_interaction": 0,
        },
    )
    pages = [
        SimpleNamespace(
            id=f"p{index}",
            score=score,
            payload={
                **scope,
                "session_id": "s1",
                "summary": f"page {index}",
                "turn_index": index,
                "source_job_id": f"job-{index}",
            },
        )
        for index, score in ((1, 0.9), (2, 0.8), (3, 0.7))
    ]

    class Store:
        def __init__(self) -> None:
            self.global_depths: list[int] = []

        def search_sessions(self, query: str, filters: dict[str, Any], top_k: int):
            return [session][:top_k]

        def search_pages(self, query: str, filters: dict[str, Any], top_k: int):
            if "session_id" not in filters:
                self.global_depths.append(top_k)
            return pages[:top_k]

        def current_turn_index(self, filters: dict[str, Any]) -> int:
            return 4

        def get_session(self, session_id: str):
            return session

    config = MidTermMemoryConfig(top_k_sessions=1, top_k_pages=3, max_total_pages=2)
    production_store = Store()
    diagnostic_store = Store()
    production = MidTermRetriever(production_store, config)
    diagnostic = DiagnosticMidTermRetriever(diagnostic_store, config)

    production_results = production.search("query", scope)
    diagnostic_results = diagnostic.search("query", scope)

    assert production_results == diagnostic_results
    assert not hasattr(production, "last_search_diagnostics")
    assert len(diagnostic.last_search_diagnostics["deduplicated_candidate_pool"]) == 3
    assert len(diagnostic.last_search_diagnostics["final_visible_pages"]) == 2
    # The default remains the production's historical 4 * max_total_pages.
    assert production_store.global_depths == diagnostic_store.global_depths == [8]
    assert "search" not in DiagnosticMidTermRetriever.__dict__


def test_production_parameterization_defaults_match_historical_constants() -> None:
    config = MemoryConfig()
    assert config.midterm.midterm_candidate_pool_multiplier == 4
    assert config.longterm_candidate_pool_multiplier == 4
    assert config.entity_similarity_threshold == 0.5


def test_production_default_longterm_overfetch_matches_historical_formula() -> None:
    requested_depths: list[int] = []
    memory = Memory.__new__(Memory)
    memory.config = MemoryConfig()
    memory.embedding_model = SimpleNamespace(embed=lambda query, mode: [1.0, 0.0])
    memory.vector_store = SimpleNamespace(
        search=lambda *, query, vectors, top_k, filters: requested_depths.append(top_k) or [],
        keyword_search=lambda *, query, top_k, filters: requested_depths.append(top_k) or [],
    )
    memory._run_entity_extraction = lambda extractor, query: []

    assert memory._search_vector_store("query", {}, limit=20) == []
    assert requested_depths == [80, 80]


def test_production_default_entity_threshold_matches_historical_half() -> None:
    memory = Memory.__new__(Memory)
    memory.config = MemoryConfig()
    memory.embedding_model = SimpleNamespace(embed_batch=lambda texts, mode: [[1.0] for _ in texts])
    memory._entity_store = SimpleNamespace(
        search=lambda **kwargs: [
            SimpleNamespace(score=0.49, payload={"linked_memory_ids": ["below"]}),
            SimpleNamespace(score=0.50, payload={"linked_memory_ids": ["at-threshold"]}),
        ]
    )

    boosts = memory._compute_entity_boosts([("company", "Acme")], {"user_id": "u1"})

    assert "below" not in boosts
    assert boosts["at-threshold"] == pytest.approx(0.25)


def test_prompt_override_is_a_production_config_and_does_not_mutate_defaults() -> None:
    production_prompt = MIDTERM_PAGE_SUMMARY_PROMPT
    config = MidTermMemoryConfig(page_summary_prompt="tuner candidate prompt")

    assert config.page_summary_prompt == "tuner candidate prompt"
    assert MidTermMemoryConfig().page_summary_prompt == production_prompt


def _checkpoint_with_history(
    monkeypatch: pytest.MonkeyPatch,
    history: list[dict[str, Any]],
) -> tuple[dict[str, Any], Any]:
    class ResolverLLM:
        def __init__(self) -> None:
            self.requests: list[list[dict[str, str]]] = []

        async def generate_response_async(self, *, messages: list[dict[str, str]], **kwargs: Any) -> str:
            del kwargs
            self.requests.append(messages)
            return json.dumps({"resolved_query": "Acme 的第二季度利润是多少？"}, ensure_ascii=False)

    class DiagnosticRetriever:
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            del args, kwargs
            self.last_search_diagnostics = {"contract": "production"}

        def search(self, query: str, filters: dict[str, str], *, record_visits: bool) -> list[dict[str, Any]]:
            assert query in {"它的第二季度利润是多少？", "Acme 的第二季度利润是多少？"}
            assert filters == {"user_id": "u1", "run_id": "s1"}
            assert record_visits is False
            return []

    llm = ResolverLLM()
    midterm_memory = SimpleNamespace(
        pages_store=object(),
        sessions_store=object(),
        current_turn_index=lambda filters: 1,
    )

    class CheckpointMemory:
        def __init__(self) -> None:
            self.llm = llm
            self.embedding_model = SimpleNamespace(embed=lambda query, mode: [float(len(query)), 1.0])
            self.midterm_memory = midterm_memory
            self.config = SimpleNamespace(midterm=SimpleNamespace(), longterm_top_k=30)

        async def _retrieve_base_context(self, query: str, **kwargs: Any) -> dict[str, Any]:
            assert kwargs == {"user_id": "u1", "session_id": "s1"}
            return {
                "query": query,
                "session_id": "s1",
                "short_term_messages": history,
                "retrieved_memories": [],
            }

        def search(self, *args: Any, **kwargs: Any) -> list[dict[str, Any]]:
            del args, kwargs
            return []

    monkeypatch.setattr(production_adapter, "DiagnosticMidTermRetriever", DiagnosticRetriever)
    monkeypatch.setattr(production_adapter, "_scroll_points", lambda store: [])
    monkeypatch.setattr(production_adapter, "_longterm_candidate_pool", lambda *args, **kwargs: {})
    monkeypatch.setattr(production_adapter, "_session_heat_states", lambda *args, **kwargs: [])

    checkpoint = production_adapter._checkpoint(
        CheckpointMemory(),
        query_id="Q2",
        query="它的第二季度利润是多少？",
        session_id="s1",
        user_id="u1",
        lineage=production_adapter.LineageTracker(3),
        ranking_depth=20,
    )
    return checkpoint, llm


def test_tuner_checkpoint_passes_short_term_messages_to_production_query_resolver(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    history = [
        {"role": "user", "content": "Acme 的第一季度利润是多少？", "turn_index": 1},
        {"role": "assistant", "content": "第一季度利润为 10。", "turn_index": 1},
    ]

    checkpoint, llm = _checkpoint_with_history(monkeypatch, history)

    assert checkpoint["retrieval_query"] == "Acme 的第二季度利润是多少？"
    assert len(llm.requests) == 1
    assert llm.requests[0] == build_query_resolution_messages("它的第二季度利润是多少？", history)
    assert production_adapter.QueryResolver is ProductionQueryResolver


def test_tuner_checkpoint_without_visible_history_keeps_original_query(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    checkpoint, llm = _checkpoint_with_history(monkeypatch, [])

    assert checkpoint["retrieval_query"] == "它的第二季度利润是多少？"
    assert llm.requests == []


@pytest.mark.asyncio
async def test_production_source_replay_requests_real_per_qa_inference(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    dataset_path = tmp_path / "dataset.xlsx"
    dataset_path.write_bytes(b"isolated-test-dataset")
    memory_config_path = tmp_path / "memory.json"
    memory_config_path.write_text("{}", encoding="utf-8")
    turn = SimpleNamespace(turn_id="S001-Q001", question="Q1", answer="A1", turn_index=1, requirements=())
    session = SimpleNamespace(session_id="S001", turns=(turn,))
    add_calls: list[dict[str, Any]] = []

    class FakeLLM:
        def generate_response(self, *args: Any, **kwargs: Any) -> str:
            del args, kwargs
            return '{"memory": []}'

    class FakeEmbedding:
        def embed(self, *args: Any, **kwargs: Any) -> list[float]:
            del args, kwargs
            return [1.0]

    class FakeMemory:
        def __init__(self) -> None:
            self.llm = FakeLLM()
            self.embedding_model = FakeEmbedding()
            self._midterm_updater = None
            self._midterm_memory = None
            self.flushes = 0
            self.closed = False

        def _short_term_capacity(self) -> int:
            return 6

        async def flush_background_tasks(self, *, timeout: float) -> bool:
            assert timeout == 30.0
            self.flushes += 1
            return True

        async def add(self, messages: list[dict[str, Any]], **kwargs: Any) -> dict[str, Any]:
            add_calls.append({"messages": messages, **kwargs})
            return {"results": [], "background": {"migration_job_id": None, "profile_job_id": None}}

        def close(self) -> bool:
            self.closed = True
            return True

    fake_memory = FakeMemory()
    monkeypatch.setattr(production_adapter, "load_dataset", lambda *args, **kwargs: [session])
    monkeypatch.setattr(production_adapter, "load_json", lambda path: {})
    monkeypatch.setattr(
        production_adapter,
        "prepare_runtime_config",
        lambda *args, **kwargs: {
            "background": {},
            "midterm": {"short_term_capacity": 6},
            "embedder": {},
            "vector_store": {},
        },
    )
    monkeypatch.setattr(production_adapter, "create_production_memory", lambda *args, **kwargs: fake_memory)
    monkeypatch.setattr(
        production_adapter,
        "_checkpoint",
        lambda *args, **kwargs: {
            "query_id": "S001-Q001",
            "query": "Q1",
            "retrieval_query": "Q1",
            "current_turn_index": 0,
            "baseline_ranking": [],
            "layered_results": [],
            "heat_states": [],
        },
    )

    await production_adapter.build_production_source(
        {
            "dataset_path": str(dataset_path),
            "session_id": "S001",
            "output_dir": str(tmp_path / "output"),
            "runtime_dir": str(tmp_path / "runtime"),
            "memory_config_path": str(memory_config_path),
            "collection_name": "test-collection",
            "ranking_depth": 20,
            "job_timeout_seconds": 30,
            "llm_mode": "mock",
        }
    )

    assert len(add_calls) == 1
    assert add_calls[0]["infer"] is True
    assert [message["content"] for message in add_calls[0]["messages"]] == ["Q1", "A1"]
    assert fake_memory.flushes == 2
    assert fake_memory.closed is True


def test_production_query_baseline_is_the_p0_prompt_hash(tmp_path: Path) -> None:
    session_id = "S001_query_baseline"
    dataset = Dataset(
        path=str(tmp_path / "dataset.xlsx"),
        sha256="a" * 64,
        sessions={session_id: (make_turn(session_id, 0),)},
    )
    _, manifest = _production_branch_inputs(tmp_path, dataset)
    baseline_config, _ = production_candidate_from_manifests([manifest])
    expected = hashlib.sha256(QUERY_REFERENCE_RESOLUTION_PROMPT.encode()).hexdigest()

    assert PRODUCTION_QUERY_PROMPT_HASH == expected
    assert baseline_config["query_prompt_hash"] == expected
    assert baseline_config["query_prompt_text"] == QUERY_REFERENCE_RESOLUTION_PROMPT


def test_all_page_prompt_candidates_use_one_production_context_contract() -> None:
    variants = controlled_page_prompt_variants({"regime": "balanced_or_plateau"})

    assert {variant["context_contract"] for variant in variants.values()} == {PRODUCTION_PAGE_CONTEXT_CONTRACT}
    assert all("context_mode" not in variant for variant in variants.values())


def test_other_session_weight_is_production_fixed_not_a_search_parameter() -> None:
    import yaml

    space = yaml.safe_load((Path(__file__).resolve().parents[2] / "search_space.yaml").read_text())
    classes = space["parameters"]["classes"]
    all_searchable = {
        value for name, values in classes.items() if name != "production_fixed_not_searched" for value in values
    }

    assert parameter_class("longterm_other_session_weight") == "production-fixed"
    assert "longterm_other_session_weight" not in all_searchable
    assert "max_tool_result_chars" not in all_searchable
    assert AgenticRetrievalConfig().max_tool_result_chars == 30000
    assert classes["production_fixed_not_searched"] == ["longterm_other_session_weight"]


def test_production_adapter_exports_full_candidate_threshold_and_cap_trace(tmp_path: Path) -> None:
    session_id = "11111111-1111-1111-1111-111111111111"
    scope = {"user_id": "recall::S001_trace", "run_id": "S001_trace"}
    pages = []
    lineage = {}
    for index in range(6):
        page_id = f"22222222-2222-2222-2222-{index:012d}"
        job_id = f"job-{index}"
        pages.append(
            {
                "id": page_id,
                "vector": [1.0, 0.0],
                "payload": {
                    **scope,
                    "session_id": session_id,
                    "source_job_id": job_id,
                    "summary": f"trace page {index}",
                    "raw_dialogue": f"trace page {index}",
                    "output_state": "committed",
                },
            }
        )
        lineage[job_id] = [f"S001-Q00{index + 1}"]
    checkpoint = {
        "query_id": "S001-Q005",
        "query": "trace query",
        "filters": scope,
        "query_vector": [1.0, 0.0],
        "sessions": [
            {
                "id": session_id,
                "vector": [1.0, 0.0],
                "payload": {
                    **scope,
                    "session_id": session_id,
                    "summary": "trace session",
                    "page_ids": [row["id"] for row in pages],
                    "output_state": "committed",
                },
            }
        ],
        "pages": pages,
        "source_turn_ids_by_job": lineage,
    }
    adapter = ProductionMidtermAdapter(
        run_dir=tmp_path,
        candidate_hash="trace-candidate",
        session_id="S001_trace",
        ranking_depth=3,
    )
    rows = adapter.rank(
        checkpoint,
        {
            "backend": "production_midterm",
            "retrieval_method": "dense",
            "query_representation": "original",
            "page_representation": "production",
            "top_k_sessions": 1,
            "top_k_pages": 6,
            "max_total_pages": 2,
            "midterm_rag_threshold": 0.0,
        },
    )
    page_rows = [row for row in rows if row.get("source") == "mid_term_page"]
    assert len(page_rows) == 6
    assert len([row for row in page_rows if row.get("final_visible")]) <= 2
    assert all(row.get("rank_before_threshold") for row in page_rows)
    assert all("threshold_filtered" in row for row in page_rows)
    assert any(row.get("ranking_loss") for row in page_rows if row.get("rank_before_threshold", 0) > 3)


def test_candidate_selection_near_tie_and_overfit() -> None:
    baseline_tune = result("baseline", 0.50)
    baseline_validation = result("baseline", 0.50)
    complex_tune = result("complex", 0.60)
    complex_validation = result("complex", 0.49)
    assert classify_overfit(
        complex_tune,
        complex_validation,
        baseline_tune,
        baseline_validation,
        validation_regression_pp=-0.5,
    )
    stable_tune = result("stable", 0.505, macro=0.51)
    stable_validation = result("stable", 0.501, macro=0.51)
    stable_validation.complexity = 0
    best, overfit, _ = select_best(
        {"baseline": baseline_tune, "complex": complex_tune, "stable": stable_tune},
        {"baseline": baseline_validation, "complex": complex_validation, "stable": stable_validation},
        baseline_name="baseline",
        tie_tolerance_pp=0.25,
        overfit_regression_pp=-0.5,
    )
    assert best.name == "stable"
    assert overfit == ["complex"]


def test_report_marks_missing_full_memory_trace_as_na(tmp_path: Path) -> None:
    baseline = result("baseline", 0.5)
    baseline.metrics["midterm_recall_at_k"] = 0.5
    reason = "No complete production trace is available"
    write_outputs(
        run_dir=tmp_path,
        dataset_audit={
            "dataset": "dataset.xlsx",
            "dataset_sha256": "sha",
            "session_count": 1,
            "query_count": 1,
            "gold_requirement_count": 1,
            "status": "OK",
            "warnings": [],
        },
        split={"method": "exploratory_only", "confidence": "exploratory", "tune_sessions": ["S001"]},
        k=5,
        baseline_tune=baseline,
        baseline_validation=baseline,
        tune_results=[baseline],
        validation_by_name={"baseline": baseline},
        best=baseline,
        selection={},
        overfit=[],
        skipped_branches=[],
        diagnostics={},
        stop_reason="test",
        run_metadata={
            "shortterm_qa_turns": 3,
            "shortterm_window_validation": {"status": "MATCH"},
            "full_memory_regression": {
                "status": "SKIPPED",
                "backend": None,
                "reason": reason,
                "metrics": None,
            },
            "midterm_baseline_backend": "production_midterm",
            "full_memory_regression_baseline_backend": None,
        },
    )
    report = (tmp_path / "final_report.md").read_text(encoding="utf-8")
    assert "Regression ShortTerm coverage: N/A" in report
    assert "Regression LongTerm R@5: N/A" in report
    assert reason in report


def test_ranking_cache_reuse_across_k(tmp_path: Path) -> None:
    registry = ArtifactRegistry(tmp_path / "cache", tmp_path / "results")
    identity = {"dataset": "x", "candidate": "y"}
    registry.store_ranking(identity, [{"query_id": "q", "page_id": "p"}], depth=40)
    assert registry.get_ranking(identity, required_depth=20) is not None
    assert registry.get_ranking(identity, required_depth=40) is not None
    assert registry.get_ranking(identity, required_depth=41) is None


def test_production_trace_discovery_checks_completeness(tmp_path: Path) -> None:
    dataset_path = tmp_path / "dataset.xlsx"
    dataset_path.write_bytes(b"dataset")
    from tuner.io_utils import sha256_file

    dataset_sha256 = sha256_file(dataset_path)
    result_dir = tmp_path / "results/run/S001"
    result_dir.mkdir(parents=True)
    trace_path = result_dir / "recall_turn_results.jsonl"
    trace_path.write_text(
        "\n".join(
            json.dumps(
                {
                    "session_id": "S001_test",
                    "turn_id": f"S001-Q{index:03d}",
                    "error": None,
                    "mid_retrieved_turn_ids": [],
                }
            )
            for index in (1, 2)
        )
        + "\n",
        encoding="utf-8",
    )
    summary = {
        "dataset": str(dataset_path),
        "sessions": [{"result_dir": str(result_dir), "failed_turns": 0}],
    }
    (tmp_path / "results/run/isolated_run_summary.json").write_text(json.dumps(summary), encoding="utf-8")
    registry = ArtifactRegistry(tmp_path / "cache", tmp_path / "results")
    candidate = registry.discover_production_trace(
        dataset_path=dataset_path,
        dataset_sha256=dataset_sha256,
        session_query_counts={"S001_test": 2},
    )
    assert candidate is not None
    assert candidate.config["backend"] == "production_trace"
    assert candidate.provenance["failed_turns"] == 0
    assert (
        registry.discover_production_trace(
            dataset_path=dataset_path,
            dataset_sha256=dataset_sha256,
            session_query_counts={"S001_test": 3},
        )
        is None
    )


def test_trace_does_not_replace_midterm_checkpoints_and_regression_is_separate(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    dataset_path = tmp_path / "dataset.xlsx"
    sessions = {}
    for index in range(1, 9):
        code = f"S{index:03d}"
        sessions[f"{code}_test"] = [
            (f"{code}-Q001", ""),
            (f"{code}-Q002", ""),
            (f"{code}-Q003", ""),
            (f"{code}-Q004", ""),
            (f"{code}-Q005", f"{code}-Q001"),
        ]
    write_workbook(dataset_path, sessions)
    memory_config = tmp_path / "memory_config.json"
    memory_config.write_text(
        json.dumps(
            {
                "midterm": {
                    "enabled": True,
                    "short_term_capacity": 6,
                    "top_k_sessions": 5,
                    "top_k_pages": 5,
                    "max_total_pages": 5,
                }
            }
        ),
        encoding="utf-8",
    )
    trace_candidate = Candidate(
        name="baseline",
        stage="baseline",
        config={"backend": "production_trace", "trace_paths": ["trace.jsonl"]},
        provenance={"trace_sha256": {"trace.jsonl": "trace-sha"}},
    )
    generated = False
    observed: list[tuple[str, str, str, dict[str, object]]] = []
    manifest_path = tmp_path / "manifest.json"

    monkeypatch.setattr(ArtifactRegistry, "discover_production_midterm", lambda self, **kwargs: None)
    monkeypatch.setattr(
        ArtifactRegistry,
        "discover_production_trace",
        lambda self, **kwargs: trace_candidate,
    )
    monkeypatch.setattr(ArtifactRegistry, "discover_frozen_rankings", lambda self, *args, **kwargs: [])

    def fake_generate(**_: object) -> list[Path]:
        nonlocal generated
        generated = True
        manifest_path.write_text("{}", encoding="utf-8")
        return [manifest_path]

    monkeypatch.setattr(orchestrator, "generate_production_sources", fake_generate)
    monkeypatch.setattr(
        orchestrator,
        "production_candidate_from_manifests",
        lambda paths: (
            {
                "backend": "production_midterm",
                "retrieval_method": "dense",
                "top_k_sessions": 5,
                "top_k_pages": 5,
                "max_total_pages": 5,
                "agentic_fixed_max_tool_result_chars": 30000,
                "manifest_paths": [str(manifest_path.resolve())],
                "manifest_sha256": {str(manifest_path.resolve()): sha256_file(manifest_path)},
            },
            {"llm_calls": 2, "embedding_calls": 11, "source": "test production source"},
        ),
    )
    monkeypatch.setattr(orchestrator, "_gpu_count", lambda: 0)
    monkeypatch.setattr(orchestrator, "_memory_gib", lambda: 16.0)

    def fake_evaluate_many(
        dataset: Dataset,
        candidates: list[Candidate] | tuple[Candidate, ...],
        session_ids: list[str] | tuple[str, ...],
        *,
        scope: str,
        **_: object,
    ) -> list[CandidateResult]:
        values = []
        for candidate in candidates:
            backend = str(candidate.config.get("backend"))
            observed.append((scope, candidate.name, backend, dict(candidate.config)))
            if backend == "production_trace":
                recall = 0.99
            elif candidate.config.get("top_k_pages") == 10:
                recall = 0.70
            else:
                recall = 0.50
            metrics = {
                "evaluated_query_count": len(session_ids),
                "eligible_requirement_count": len(session_ids),
                "recall_at_k": recall,
                "recall_at_2k": recall,
                "recall_at_4k": recall,
                "macro_session_recall_at_k": recall,
                "session_stddev": 0.0,
                "worst_session_recall_at_k": recall,
                "mrr": recall,
                "midterm_recall_at_k": recall,
                "shortterm_coverage": 0.25,
                "target_layer_union": 0.60,
                "all_memory_union": None,
                "query_completion": None,
            }
            if backend == "production_trace":
                metrics.update(
                    {
                        "longterm_recall_at_k": 0.40,
                        "all_memory_recall_at_k": 0.80,
                        "all_memory_union": 0.80,
                        "query_completion": 0.80,
                    }
                )
            values.append(
                CandidateResult(
                    name=candidate.name,
                    candidate_hash=candidate.name,
                    stage=candidate.stage,
                    config=candidate.config,
                    metrics=metrics,
                    requirement_rows=[],
                    session_rows=[],
                    runtime_seconds=0.01,
                    work_seconds=0.01,
                    cache_hits=0,
                    cache_misses=len(session_ids),
                    complexity=candidate.complexity,
                )
            )
        return values

    monkeypatch.setattr(orchestrator, "_evaluate_many", fake_evaluate_many)
    run_dir = orchestrator.run_tuning(
        TunerConfig(
            dataset=dataset_path,
            budget="quick",
            output_root=tmp_path / "results",
            memory_config=memory_config,
            max_parallel_sessions=4,
            max_parallel_llm_calls=2,
        ),
        skill_root=SCRIPTS.parent,
    )

    metadata = json.loads((run_dir / "run_metadata.json").read_text(encoding="utf-8"))
    best = json.loads((run_dir / "best_config.json").read_text(encoding="utf-8"))
    report = (run_dir / "final_report.md").read_text(encoding="utf-8")
    assert generated
    assert metadata["midterm_baseline_backend"] == "production_midterm"
    assert metadata["full_memory_regression_baseline_backend"] == "production_trace"
    assert metadata["source_generation_llm_calls"] == 2
    assert metadata["source_generation_embedding_calls"] == 11
    assert metadata["llm_calls"] == 2
    assert metadata["embedding_calls"] == 11
    assert metadata["execution"]["source_worker_parallelism"] == 2
    assert metadata["shortterm_qa_turns"] == 3
    assert metadata["resolved_memory_config_sha256"] == sha256_file(Path(metadata["resolved_memory_config_path"]))
    assert best["candidate"].startswith("RetrievalControl:top_k_pages=10")
    assert best["config"]["top_k_pages"] == 10
    assert best["validation_metrics"]["recall_at_k"] == pytest.approx(0.70)
    cheap_configs = [value for scope, _, _, value in observed if scope == "stage_1_tune"]
    assert all(value.get("top_k_sessions") == 5 for value in cheap_configs)
    assert any(value.get("top_k_pages") != 5 for value in cheap_configs)
    assert any(value.get("max_total_pages") != 5 for value in cheap_configs)
    assert all(scope == "full_memory_regression" for scope, _, backend, _ in observed if backend == "production_trace")
    assert metadata["full_memory_regression"]["metrics"]["longterm_recall_at_k"] == 0.40
    assert "Regression LongTerm R@5: 0.4000" in report
    assert "Regression All-memory R@5 / query completion: 0.8000 / 0.8000" in report


def test_frozen_query_artifact_discovery_validates_dataset_and_original_query(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 1)
    results = tmp_path / "results/query_run"
    results.mkdir(parents=True)
    (results / "run_metadata.json").write_text(
        json.dumps({"dataset_sha256": dataset.sha256}),
        encoding="utf-8",
    )
    target = dataset.sessions["S001_测试"][-1]
    artifact = results / "resolved_queries.jsonl"
    artifact.write_text(
        json.dumps(
            {
                "query_id": target.query_id,
                "original_query": target.question,
                "resolved_query": f"已消解：{target.question}",
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    base = Candidate(
        name="baseline",
        stage="baseline",
        config={"backend": "production_midterm", "query_representation": "original"},
    )
    registry = ArtifactRegistry(tmp_path / "cache", tmp_path / "results")
    discovered = registry.discover_frozen_queries(
        dataset.sha256,
        query_text_by_id={turn.query_id: turn.question for turns in dataset.sessions.values() for turn in turns},
        base_candidate=base,
    )
    assert len(discovered) == 1
    assert discovered[0].config["query_representation"] == "bounded_reference_resolution"
    assert discovered[0].provenance["matched_query_count"] == 1

    artifact.write_text(
        json.dumps(
            {
                "query_id": target.query_id,
                "original_query": "不属于当前数据集的问题",
                "resolved_query": target.question,
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    assert not registry.discover_frozen_queries(
        dataset.sha256,
        query_text_by_id={target.query_id: target.question},
        base_candidate=base,
    )


def test_parallel_session_candidate_isolation_and_resume(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 2)
    registry = ArtifactRegistry(tmp_path / "cache", tmp_path / "results")
    run_dir = tmp_path / "run"
    ranking_path = tmp_path / "frozen_rankings.jsonl"
    ranking_path.write_text(
        "\n".join(
            json.dumps(
                {
                    "query_id": turns[-1].query_id,
                    "source_turn_id": turns[0].query_id,
                    "rank": 1,
                    "score": 1.0,
                }
            )
            for turns in dataset.sessions.values()
        )
        + "\n",
        encoding="utf-8",
    )
    baseline = Candidate(
        name="baseline",
        stage="baseline",
        config={
            "backend": "frozen_ranking",
            "ranking_path": str(ranking_path),
            "max_total_pages": 20,
        },
    )
    alternative = Candidate(
        name="alternative",
        stage="cheap",
        config={**baseline.config, "max_total_pages": 5},
    )
    evaluate_candidate(
        dataset,
        baseline,
        list(dataset.sessions),
        scope="tune",
        k=5,
        target="midterm",
        shortterm_window=3,
        ranking_depth=20,
        max_parallel_sessions=2,
        registry=registry,
        run_dir=run_dir,
    )
    second = evaluate_candidate(
        dataset,
        baseline,
        list(dataset.sessions),
        scope="tune",
        k=5,
        target="midterm",
        shortterm_window=3,
        ranking_depth=20,
        max_parallel_sessions=2,
        registry=registry,
        run_dir=run_dir,
    )
    assert second.cache_hits == 2
    baseline_hash = candidate_hash(dataset.sha256, baseline)
    alternative_hash = candidate_hash(dataset.sha256, alternative)
    assert baseline_hash != alternative_hash
    worker_scope = "tune__k5__midterm__short3"
    assert len(list((run_dir / "workers" / baseline_hash / worker_scope).glob("*.json"))) == 2

    interrupted = registry.worker_path(run_dir, alternative_hash, worker_scope, next(iter(dataset.sessions)))
    interrupted.parent.mkdir(parents=True, exist_ok=True)
    interrupted.write_text('{"status": "RUNNING"}', encoding="utf-8")
    evaluated = evaluate_candidate(
        dataset,
        alternative,
        list(dataset.sessions),
        scope="tune",
        k=5,
        target="midterm",
        shortterm_window=3,
        ranking_depth=20,
        max_parallel_sessions=2,
        registry=registry,
        run_dir=run_dir,
    )
    assert evaluated.metrics["eligible_requirement_count"] == 2
    assert json.loads(interrupted.read_text(encoding="utf-8"))["status"] == "COMPLETE"


def test_k_change_reuses_deep_raw_ranking_cache(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 1)
    registry = ArtifactRegistry(tmp_path / "cache", tmp_path / "results")
    turns = next(iter(dataset.sessions.values()))
    ranking_path = tmp_path / "frozen_rankings.jsonl"
    ranking_path.write_text(
        json.dumps(
            {
                "query_id": turns[-1].query_id,
                "source_turn_id": turns[0].query_id,
                "rank": 1,
                "score": 1.0,
            }
        )
        + "\n",
        encoding="utf-8",
    )
    base_config = {
        "backend": "frozen_ranking",
        "ranking_path": str(ranking_path),
        "max_total_pages": 20,
    }
    baseline = Candidate(name="baseline", stage="baseline", config=base_config)
    same_ranking = Candidate(name="same-ranking", stage="cheap", config=base_config)
    common = {
        "dataset": dataset,
        "sessions": list(dataset.sessions),
        "scope": "tune",
        "k": 5,
        "target": "midterm",
        "shortterm_window": 3,
        "ranking_depth": 20,
        "max_parallel_sessions": 1,
        "registry": registry,
        "run_dir": tmp_path / "run",
    }
    evaluate_candidate(candidate=baseline, **common)
    reused = evaluate_candidate(candidate=same_ranking, **{**common, "k": 10})
    assert reused.cache_hits == 1
    assert reused.cache_misses == 0


def test_branch_registry_exposes_required_experiment_contracts() -> None:
    registry = BranchRegistry()
    descriptions = {item["name"]: item for item in registry.describe()}
    assert {
        "RetrievalControl",
        "AgenticRetrieval",
        "QueryRepresentation",
        "PageRepresentation",
        "HybridRetrieval",
        "Reranking",
        "Embedding",
        "FieldAwareMultiVector",
        "QueryRewritePrompt",
        "MidtermPageSummaryPrompt",
        "MidtermSessionMergePrompt",
        "FineGrainedLongtermExtractionPrompt",
    } <= set(descriptions)
    for item in descriptions.values():
        assert item["diagnostic_regimes"]
        assert item["cost_level"] in {"cheap", "medium", "high", "expensive"}
        assert item["required_artifacts"]
        assert item["execution_adapter"]
        assert item["provenance_contract"]
        assert isinstance(item["resource_requirements"], dict)


def _write_production_agentic_trace(
    path: Path,
    dataset: Dataset,
    variants: list[tuple[int, int]],
    *,
    parent_config: dict[str, Any] | None = None,
    parent_identity: dict[str, Any] | None = None,
    supplement_label: str = "Agentic",
    max_tool_result_chars: int = 30000,
) -> None:
    parent = parent_identity or build_agentic_parent_retrieval_identity(parent_config or {})
    parent_sha256 = agentic_parent_retrieval_identity_sha256(parent)
    rows = []
    for max_queries, max_total_results in variants:
        for session_id, turns in dataset.sessions.items():
            for turn in turns:
                rows.append(
                    {
                        "schema": PRODUCTION_AGENTIC_TRACE_SCHEMA,
                        "dataset_sha256": dataset.sha256,
                        "session_id": session_id,
                        "query_id": turn.query_id,
                        "production_agentic_execution": True,
                        "execution_contract": PRODUCTION_AGENTIC_EXECUTION_CONTRACT,
                        "parent_retrieval_identity": parent,
                        "parent_retrieval_identity_sha256": parent_sha256,
                        "max_iterations": 2,
                        "max_tool_calls": 1,
                        "max_tool_result_chars": max_tool_result_chars,
                        "max_queries": max_queries,
                        "max_total_results": max_total_results,
                        "error": None,
                        "agentic_result": {
                            "status": "supplemented",
                            "supplement": (
                                f"华辰公司授信额度100万元；{supplement_label} "
                                f"{max_queries}/{max_total_results} {turn.query_id}"
                            ),
                            "iterations": 2,
                            "tool_call_count": 1,
                            "tool_trace": [
                                {
                                    "iteration": 1,
                                    "name": "search_memory",
                                    "arguments": {"queries": [turn.question][:max_queries]},
                                    "result_summary": {"ok": True, "item_count": max_total_results},
                                }
                            ],
                        },
                    }
                )
    path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def _agentic_query_ids(dataset: Dataset) -> dict[str, list[str]]:
    return {session_id: [turn.query_id for turn in turns] for session_id, turns in dataset.sessions.items()}


def _agentic_trace_config(paths: list[Path], *, max_tool_result_chars: int = 30000) -> dict[str, Any]:
    return {
        "agentic_fixed_max_tool_result_chars": max_tool_result_chars,
        "production_agentic_trace_paths": [str(path) for path in paths],
        "production_agentic_trace_sha256": {str(path.resolve()): sha256_file(path) for path in paths},
    }


def _synthetic_agentic_source(label: str = "source-a") -> dict[str, Any]:
    return {
        "source_generation_spec": {
            "source_identity": {
                "kind": "synthetic-production-source",
                "source_content_sha256": label,
            }
        }
    }


def _write_semantic_agentic_manifest(
    path: Path,
    *,
    runtime_root: Path,
    checkpoints_sha256: str = "checkpoints-a",
    trace_sha256: str = "trace-a",
    embedding_model_id: str = "embedding-a",
    embedding_model_revision: str = "revision-a",
    prompt_hash: str = "prompt-a",
    source_label: str = "source-a",
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    manifest = {
        "schema": ADAPTER_SCHEMA,
        "status": "COMPLETE",
        "backend": "production_midterm",
        "dataset_path": str(runtime_root / "dataset.xlsx"),
        "dataset_sha256": "dataset-a",
        "session_id": "S001",
        "turn_count": 5,
        "checkpoint_count": 2,
        "failed_turns": 0,
        "ranking_depth": 20,
        "shortterm_qa_turns": 3,
        "checkpoints_path": str(runtime_root / "checkpoints.jsonl"),
        "checkpoints_sha256": checkpoints_sha256,
        "trace_path": str(runtime_root / "trace.jsonl"),
        "trace_sha256": trace_sha256,
        "memory_config_path": str(runtime_root / "memory.json"),
        "production_config": {
            "short_term_capacity": 6,
            "top_k_sessions": 5,
            "top_k_pages": 5,
            "max_total_pages": 5,
        },
        "effective_memory_config": {
            "llm": {"provider": "mock", "config": {"model": "source-llm"}},
            "embedder": {
                "provider": "huggingface",
                "config": {"model": embedding_model_id, "embedding_dims": 512},
            },
            "vector_store": {
                "provider": "qdrant",
                "config": {
                    "path": str(runtime_root / "qdrant"),
                    "collection_name": f"temporary-{runtime_root.name}",
                    "embedding_model_dims": 512,
                    "bm25_language": "zh",
                },
            },
            "history_db_path": str(runtime_root / "history.db"),
            "midterm": {
                "short_term_capacity": 6,
                "session_similarity_threshold": 0.6,
                "embedding_similarity_weight": 0.7,
                "keyword_overlap_weight": 0.3,
                "top_k_sessions": 5,
                "top_k_pages": 5,
                "max_total_pages": 5,
            },
        },
        "config_overrides": {"embedder": {"config": {"model": embedding_model_id}}},
        "prompt_hashes": {"page_summary": prompt_hash, "session_merge": "merge-a"},
        "source_identity": {"source_content_sha256": source_label, "runtime_dir": str(runtime_root)},
        "source_variant": "production",
        "page_context_contract": PRODUCTION_PAGE_CONTEXT_CONTRACT,
        "production_memory_contract": "production-memory-v1",
        "embedding_model": {
            "provider": "huggingface",
            "config": {"model": embedding_model_id, "embedding_dims": 512},
        },
        "embedding_model_id": embedding_model_id,
        "embedding_model_revision": embedding_model_revision,
        "embedding_encoding_contract": {"normalize_embeddings": True, "pooling": "mean"},
        "stateful_replay": True,
        "llm_mode": "real",
    }
    path.write_text(json.dumps(manifest), encoding="utf-8")
    return path


def test_agentic_trace_accepts_exact_parent_retrieval_identity(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 1)
    parent_config = {
        "top_k_sessions": 3,
        "top_k_pages": 8,
        "max_total_pages": 4,
        "midterm_candidate_pool_multiplier": 4,
        "midterm_rag_threshold": 0.1,
        "retrieval_method": "dense",
        "query_representation": "original",
        "manifest_sha256": {"/mutable/runtime/manifest.json": "manifest-a"},
        **_synthetic_agentic_source(),
        "longterm_top_k": 20,
        "longterm_rag_threshold": 0.1,
    }
    expected = build_agentic_parent_retrieval_identity(parent_config)
    trace_path = tmp_path / "production_agentic_trace_exact.jsonl"
    _write_production_agentic_trace(trace_path, dataset, [(3, 5)], parent_config=parent_config)

    trace = load_production_agentic_trace(
        _agentic_trace_config([trace_path]),
        dataset_sha256=dataset.sha256,
        query_ids_by_session=_agentic_query_ids(dataset),
        expected_parent_retrieval_identity=expected,
    )

    assert trace.variants == ((3, 5),)
    assert trace.max_tool_result_chars == 30000
    assert trace.serializable()["fixed"]["max_tool_result_chars"] == 30000
    assert trace.parent_retrieval_identity == expected
    assert trace.parent_retrieval_identity_sha256 == agentic_parent_retrieval_identity_sha256(expected)


def test_agentic_parent_identity_covers_effective_inputs_but_excludes_bookkeeping_paths() -> None:
    base = {
        "top_k_sessions": 3,
        "top_k_pages": 8,
        "max_total_pages": 4,
        "midterm_candidate_pool_multiplier": 4,
        "midterm_rag_threshold": 0.1,
        "retrieval_method": "dense",
        "query_representation": "original",
        "query_artifact_sha256": "query-a",
        "page_representation": "production",
        "derived_artifact_sha256": "derived-a",
        "manifest_sha256": {"/runtime-a/manifest.json": "manifest-a"},
        "source_generation_spec": {
            "source_root": "/cache-a/source",
            "source_identity": {
                "kind": "production-source",
                "prompt_hash": "prompt-a",
                "runtime_dir": "/runtime-a",
            },
        },
        "longterm_top_k": 20,
        "longterm_rag_threshold": 0.1,
        "longterm_hybrid_preset": "balanced",
    }
    identity = build_agentic_parent_retrieval_identity(base)
    path_and_bookkeeping_only = {
        **base,
        "manifest_sha256": {"/runtime-b/manifest.json": "manifest-a"},
        "source_generation_spec": {
            "source_root": "/cache-b/source",
            "source_identity": {
                "kind": "production-source",
                "prompt_hash": "prompt-a",
                "runtime_dir": "/runtime-b",
            },
        },
        "derived_artifact_path": "/cache-b/derived.json",
        "experiment_branch": "report-only",
        "branch_cost_level": "expensive",
        "parent_candidate_hash": "mutable-parent-name",
        "applied_branches": ["report-only"],
    }
    assert build_agentic_parent_retrieval_identity(path_and_bookkeeping_only) == identity
    for field, value in (
        ("top_k_pages", 9),
        ("query_artifact_sha256", "query-b"),
        ("page_representation", "summary"),
        ("longterm_rag_threshold", 0.2),
    ):
        assert build_agentic_parent_retrieval_identity({**base, field: value}) != identity
    changed_source = {
        **base,
        "source_generation_spec": {
            **base["source_generation_spec"],
            "source_identity": {"kind": "production-source", "prompt_hash": "prompt-b"},
        },
    }
    assert build_agentic_parent_retrieval_identity(changed_source) != identity


def test_agentic_semantic_manifest_identity_ignores_runtime_paths(tmp_path: Path) -> None:
    first = _write_semantic_agentic_manifest(
        tmp_path / "runtime-a" / "production_midterm_manifest.json",
        runtime_root=tmp_path / "runtime-a",
    )
    second = _write_semantic_agentic_manifest(
        tmp_path / "runtime-b" / "production_midterm_manifest.json",
        runtime_root=tmp_path / "runtime-b",
    )

    first_identity = build_agentic_parent_retrieval_identity({"manifest_paths": [str(first)]})
    second_identity = build_agentic_parent_retrieval_identity({"manifest_paths": [str(second)]})

    assert sha256_file(first) != sha256_file(second)
    assert first_identity == second_identity
    serialized = json.dumps(first_identity, ensure_ascii=False)
    assert str(tmp_path / "runtime-a") not in serialized
    assert "temporary-runtime-a" not in serialized


def test_agentic_semantic_manifest_identity_changes_with_source_content(tmp_path: Path) -> None:
    base = _write_semantic_agentic_manifest(
        tmp_path / "base" / "production_midterm_manifest.json",
        runtime_root=tmp_path / "base",
    )
    changed_checkpoints = _write_semantic_agentic_manifest(
        tmp_path / "changed" / "production_midterm_manifest.json",
        runtime_root=tmp_path / "changed",
        checkpoints_sha256="checkpoints-b",
    )

    assert build_agentic_parent_retrieval_identity(
        {"manifest_paths": [str(base)]}
    ) != build_agentic_parent_retrieval_identity({"manifest_paths": [str(changed_checkpoints)]})


@pytest.mark.parametrize(
    "change",
    [
        {"embedding_model_id": "embedding-b"},
        {"embedding_model_revision": "revision-b"},
        {"prompt_hash": "prompt-b"},
        {"source_label": "source-b"},
    ],
)
def test_agentic_semantic_manifest_identity_changes_with_embedding_prompt_or_source(
    tmp_path: Path,
    change: dict[str, str],
) -> None:
    base = _write_semantic_agentic_manifest(
        tmp_path / "base" / "production_midterm_manifest.json",
        runtime_root=tmp_path / "base",
    )
    changed = _write_semantic_agentic_manifest(
        tmp_path / "changed" / "production_midterm_manifest.json",
        runtime_root=tmp_path / "changed",
        **change,
    )

    assert build_agentic_parent_retrieval_identity(
        {"manifest_paths": [str(base)]}
    ) != build_agentic_parent_retrieval_identity({"manifest_paths": [str(changed)]})


def test_agentic_trace_rejects_different_retrieval_config_for_same_dataset_and_variant(
    tmp_path: Path,
) -> None:
    dataset = make_dataset(tmp_path, 1)
    parent_config = {
        "top_k_sessions": 3,
        "top_k_pages": 8,
        "max_total_pages": 4,
        "midterm_rag_threshold": 0.1,
        "retrieval_method": "dense",
        "query_representation": "original",
        "manifest_sha256": {"manifest": "manifest-a"},
        **_synthetic_agentic_source(),
    }
    trace_path = tmp_path / "production_agentic_trace_retrieval_a.jsonl"
    _write_production_agentic_trace(trace_path, dataset, [(3, 5)], parent_config=parent_config)
    changed = {**parent_config, "max_total_pages": 2, "midterm_rag_threshold": 0.15}

    with pytest.raises(ValueError, match="parent retrieval identity mismatch"):
        load_production_agentic_trace(
            _agentic_trace_config([trace_path]),
            dataset_sha256=dataset.sha256,
            query_ids_by_session=_agentic_query_ids(dataset),
            expected_parent_retrieval_identity=build_agentic_parent_retrieval_identity(changed),
        )


def test_agentic_trace_rejects_fixed_max_tool_result_chars_mismatch(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 1)
    parent_config = {"max_total_pages": 4, **_synthetic_agentic_source()}
    trace_path = tmp_path / "production_agentic_trace_chars_10000.jsonl"
    _write_production_agentic_trace(
        trace_path,
        dataset,
        [(3, 5)],
        parent_config=parent_config,
        max_tool_result_chars=10000,
    )

    with pytest.raises(ValueError, match="max_tool_result_chars mismatch"):
        load_production_agentic_trace(
            _agentic_trace_config([trace_path], max_tool_result_chars=30000),
            dataset_sha256=dataset.sha256,
            query_ids_by_session=_agentic_query_ids(dataset),
            expected_parent_retrieval_identity=build_agentic_parent_retrieval_identity(parent_config),
        )


def test_agentic_trace_rejects_different_manifest_identity_with_same_retrieval_config(
    tmp_path: Path,
) -> None:
    dataset = make_dataset(tmp_path, 1)
    source_a = _write_semantic_agentic_manifest(
        tmp_path / "source-a" / "production_midterm_manifest.json",
        runtime_root=tmp_path / "source-a",
    )
    source_b = _write_semantic_agentic_manifest(
        tmp_path / "source-b" / "production_midterm_manifest.json",
        runtime_root=tmp_path / "source-b",
        checkpoints_sha256="checkpoints-b",
    )
    parent_config = {
        "top_k_sessions": 3,
        "top_k_pages": 8,
        "max_total_pages": 4,
        "midterm_rag_threshold": 0.1,
        "retrieval_method": "dense",
        "query_representation": "original",
        "manifest_paths": [str(source_a)],
    }
    trace_path = tmp_path / "production_agentic_trace_source_a.jsonl"
    _write_production_agentic_trace(trace_path, dataset, [(3, 5)], parent_config=parent_config)
    changed = {**parent_config, "manifest_paths": [str(source_b)]}

    with pytest.raises(ValueError, match="parent retrieval identity mismatch"):
        load_production_agentic_trace(
            _agentic_trace_config([trace_path]),
            dataset_sha256=dataset.sha256,
            query_ids_by_session=_agentic_query_ids(dataset),
            expected_parent_retrieval_identity=build_agentic_parent_retrieval_identity(changed),
        )


def test_agentic_discovery_aggregates_exact_variants_from_multiple_files(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 1)
    parent_config = {
        "top_k_sessions": 3,
        "top_k_pages": 8,
        "max_total_pages": 4,
        "retrieval_method": "dense",
        "query_representation": "original",
        "manifest_sha256": {"manifest": "manifest-a"},
        **_synthetic_agentic_source(),
    }
    results_root = tmp_path / "results"
    results_root.mkdir()
    first = results_root / "trace_1_5.jsonl"
    second = results_root / "trace_3_1.jsonl"
    _write_production_agentic_trace(first, dataset, [(1, 5)], parent_config=parent_config)
    _write_production_agentic_trace(second, dataset, [(3, 1)], parent_config=parent_config)
    expected = build_agentic_parent_retrieval_identity(parent_config)

    discovered = ArtifactRegistry(tmp_path / "cache", results_root).discover_production_agentic_trace(
        dataset_sha256=dataset.sha256,
        query_ids_by_session=_agentic_query_ids(dataset),
        expected_parent_retrieval_identity=expected,
        expected_max_tool_result_chars=30000,
    )

    assert discovered is not None
    assert discovered["variants"] == [
        {"max_queries": 1, "max_total_results": 5},
        {"max_queries": 3, "max_total_results": 1},
    ]
    assert set(discovered["paths"]) == {str(first.resolve()), str(second.resolve())}
    assert set(discovered["sha256"]) == {str(first.resolve()), str(second.resolve())}
    assert discovered["parent_retrieval_identity"] == expected
    assert discovered["fixed"]["max_tool_result_chars"] == 30000


def test_agentic_discovery_filters_fixed_chars_before_multi_file_aggregation(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 1)
    parent_config = {
        "max_total_pages": 4,
        "retrieval_method": "dense",
        **_synthetic_agentic_source(),
    }
    results_root = tmp_path / "results"
    results_root.mkdir()
    old = results_root / "production_agentic_trace_1_5_10000.jsonl"
    current_first = results_root / "production_agentic_trace_1_5_30000.jsonl"
    current_second = results_root / "production_agentic_trace_3_1_30000.jsonl"
    _write_production_agentic_trace(
        old,
        dataset,
        [(1, 5)],
        parent_config=parent_config,
        max_tool_result_chars=10000,
        supplement_label="old-chars",
    )
    _write_production_agentic_trace(
        current_first,
        dataset,
        [(1, 5)],
        parent_config=parent_config,
        max_tool_result_chars=30000,
    )
    _write_production_agentic_trace(
        current_second,
        dataset,
        [(3, 1)],
        parent_config=parent_config,
        max_tool_result_chars=30000,
    )

    discovered = ArtifactRegistry(tmp_path / "cache", results_root).discover_production_agentic_trace(
        dataset_sha256=dataset.sha256,
        query_ids_by_session=_agentic_query_ids(dataset),
        expected_parent_retrieval_identity=build_agentic_parent_retrieval_identity(parent_config),
        expected_max_tool_result_chars=30000,
    )

    assert discovered is not None
    assert discovered["variants"] == [
        {"max_queries": 1, "max_total_results": 5},
        {"max_queries": 3, "max_total_results": 1},
    ]
    assert set(discovered["paths"]) == {str(current_first.resolve()), str(current_second.resolve())}
    assert str(old.resolve()) not in discovered["sha256"]


def test_agentic_multi_file_trace_rejects_mixed_parent_identities(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 1)
    parent_a = {"max_total_pages": 4, **_synthetic_agentic_source()}
    parent_b = {"max_total_pages": 2, **_synthetic_agentic_source()}
    first = tmp_path / "production_agentic_trace_parent_a.jsonl"
    second = tmp_path / "production_agentic_trace_parent_b.jsonl"
    _write_production_agentic_trace(first, dataset, [(1, 5)], parent_config=parent_a)
    _write_production_agentic_trace(second, dataset, [(3, 1)], parent_config=parent_b)

    with pytest.raises(ValueError, match="parent retrieval identity mismatch"):
        load_production_agentic_trace(
            _agentic_trace_config([first, second]),
            dataset_sha256=dataset.sha256,
            query_ids_by_session=_agentic_query_ids(dataset),
            expected_parent_retrieval_identity=build_agentic_parent_retrieval_identity(parent_a),
        )


def test_agentic_multi_file_trace_rejects_conflicting_duplicate_variant_query(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 1)
    parent_config = {"max_total_pages": 4, **_synthetic_agentic_source()}
    first = tmp_path / "production_agentic_trace_duplicate_a.jsonl"
    second = tmp_path / "production_agentic_trace_duplicate_b.jsonl"
    _write_production_agentic_trace(
        first,
        dataset,
        [(1, 5)],
        parent_config=parent_config,
        supplement_label="first",
    )
    _write_production_agentic_trace(
        second,
        dataset,
        [(1, 5)],
        parent_config=parent_config,
        supplement_label="conflicting",
    )

    with pytest.raises(ValueError, match="duplicate/conflicting"):
        load_production_agentic_trace(
            _agentic_trace_config([first, second]),
            dataset_sha256=dataset.sha256,
            query_ids_by_session=_agentic_query_ids(dataset),
            expected_parent_retrieval_identity=build_agentic_parent_retrieval_identity(parent_config),
        )


def test_agentic_branch_generates_real_parameter_candidates_and_keeps_fixed_limits(tmp_path: Path) -> None:
    import yaml

    dataset = make_dataset(tmp_path, 1)
    baseline, _ = _production_branch_inputs(tmp_path, dataset)
    baseline.config.update({"max_queries": 3, "max_total_results": 5})
    trace_path = tmp_path / "production_agentic_trace.jsonl"
    variants = [(value, 5) for value in (1, 2, 3)] + [(3, value) for value in (1, 2, 3, 4)]
    _write_production_agentic_trace(trace_path, dataset, variants, parent_config=baseline.config)
    baseline.config.update(
        {
            "production_agentic_trace_paths": [str(trace_path)],
            "production_agentic_trace_sha256": {str(trace_path.resolve()): sha256_file(trace_path)},
        }
    )
    space = yaml.safe_load((SCRIPTS.parent / "search_space.yaml").read_text(encoding="utf-8"))
    context = _branch_context(tmp_path, dataset, baseline, search_space=space)
    branch = BranchRegistry().get("AgenticRetrieval")
    outcome = BranchRegistry([branch]).generate(branch, context)

    assert outcome.status == "READY"
    assert {candidate.config["max_queries"] for candidate in outcome.candidates} >= {1, 2}
    assert {candidate.config["max_total_results"] for candidate in outcome.candidates} >= {1, 2, 3, 4}
    assert all(candidate.config["agentic_trace_enabled"] is True for candidate in outcome.candidates)
    assert all(candidate.config["agentic_fixed_max_iterations"] == 2 for candidate in outcome.candidates)
    assert all(candidate.config["agentic_fixed_max_tool_calls"] == 1 for candidate in outcome.candidates)
    assert all(candidate.config["agentic_fixed_max_tool_result_chars"] == 30000 for candidate in outcome.candidates)
    assert all(candidate.config["production_agentic_trace_sha256"] for candidate in outcome.candidates)


def test_agentic_branch_is_unavailable_without_production_trace(tmp_path: Path) -> None:
    import yaml

    dataset = make_dataset(tmp_path, 1)
    baseline, _ = _production_branch_inputs(tmp_path, dataset)
    baseline.config.update({"max_queries": 3, "max_total_results": 5})
    space = yaml.safe_load((SCRIPTS.parent / "search_space.yaml").read_text(encoding="utf-8"))
    context = _branch_context(tmp_path, dataset, baseline, search_space=space)
    branch = AgenticRetrievalBranch()
    outcome = BranchRegistry([branch]).generate(branch, context)

    assert outcome.status == "UNAVAILABLE"
    assert outcome.candidates == []
    assert "production_agentic_trace unavailable" in str(outcome.reason)


def test_agentic_candidates_enter_staged_tune_search(tmp_path: Path) -> None:
    import yaml

    dataset = make_dataset(tmp_path, 2)
    baseline, _ = _production_branch_inputs(tmp_path, dataset)
    baseline.config.update({"max_queries": 3, "max_total_results": 5})
    trace_path = tmp_path / "production_agentic_trace.jsonl"
    variants = [(value, 5) for value in (1, 2, 3)] + [(3, value) for value in (1, 2, 3, 4)]
    _write_production_agentic_trace(trace_path, dataset, variants, parent_config=baseline.config)
    baseline.config.update(
        {
            "production_agentic_trace_paths": [str(trace_path)],
            "production_agentic_trace_sha256": {str(trace_path.resolve()): sha256_file(trace_path)},
        }
    )
    space = yaml.safe_load((SCRIPTS.parent / "search_space.yaml").read_text(encoding="utf-8"))
    space["search"]["branch_coverage"] = {
        "candidate_coverage_bottleneck": {
            "relevant": {
                "RetrievalControl": {"priority": 10, "minimum_attempts": 1, "coverage_class": "required"},
                "AgenticRetrieval": {"priority": 20, "minimum_attempts": 1, "coverage_class": "required"},
            }
        }
    }
    evaluated_scopes: list[tuple[str, list[str]]] = []

    def evaluate(candidates: Any, sessions: Any, scope: str) -> list[CandidateResult]:
        del sessions
        evaluated_scopes.append((scope, [candidate.name for candidate in candidates]))
        values = []
        for candidate in candidates:
            measured = result(candidate.name, 0.5)
            measured.config = candidate.config
            measured.stage = candidate.stage
            measured.complexity = candidate.complexity
            values.append(measured)
        return values

    baseline_result = result("baseline", 0.5)
    baseline_result.config = baseline.config
    search = run_staged_search(
        dataset=dataset,
        baseline=baseline,
        baseline_result=baseline_result,
        tune_sessions=tuple(dataset.sessions),
        registry=BranchRegistry([RetrievalControlBranch(), AgenticRetrievalBranch()]),
        artifact_registry=ArtifactRegistry(tmp_path / "cache", tmp_path / "results"),
        model_discovery=None,
        run_dir=tmp_path / "run",
        search_space=space,
        budget="quick",
        profile={
            "max_stages": 2,
            "max_cost_level": "medium",
            "max_branches_per_stage": 1,
            "max_candidates_per_stage": 20,
            "tune_frontier": 20,
            "max_expensive_candidates": 0,
        },
        k=5,
        ranking_depth=20,
        evaluate=evaluate,
        diagnose=lambda _: {"regime": "candidate_coverage_bottleneck"},
    )

    assert [stage["branches"] for stage in search.stage_history] == [
        ["RetrievalControl"],
        ["AgenticRetrieval"],
    ]
    assert any(
        scope == "stage_2_tune" and any(name.startswith("AgenticRetrieval:") for name in names)
        for scope, names in evaluated_scopes
    )


def test_staged_agentic_rejects_baseline_trace_after_retrieval_anchor_changes(tmp_path: Path) -> None:
    import yaml

    dataset = make_dataset(tmp_path, 2)
    baseline, _ = _production_branch_inputs(tmp_path, dataset)
    baseline.config.update({"max_queries": 3, "max_total_results": 5})
    trace_path = tmp_path / "production_agentic_trace_baseline_anchor.jsonl"
    variants = [(value, 5) for value in (1, 2, 3)] + [(3, value) for value in (1, 2, 3, 4)]
    _write_production_agentic_trace(trace_path, dataset, variants, parent_config=baseline.config)
    baseline.config.update(_agentic_trace_config([trace_path]))
    space = yaml.safe_load((SCRIPTS.parent / "search_space.yaml").read_text(encoding="utf-8"))
    space["search"]["branch_coverage"] = {
        "candidate_coverage_bottleneck": {
            "relevant": {
                "RetrievalControl": {"priority": 10, "minimum_attempts": 1, "coverage_class": "required"},
                "AgenticRetrieval": {"priority": 20, "minimum_attempts": 1, "coverage_class": "required"},
            }
        }
    }

    def evaluate(candidates: Any, sessions: Any, scope: str) -> list[CandidateResult]:
        del sessions, scope
        values = []
        for candidate in candidates:
            recall = 0.8 if candidate.config.get("experiment_branch") == "RetrievalControl" else 0.5
            measured = result(candidate.name, recall)
            measured.config = candidate.config
            measured.stage = candidate.stage
            measured.complexity = candidate.complexity
            values.append(measured)
        return values

    baseline_result = result("baseline", 0.5)
    baseline_result.config = baseline.config
    search = run_staged_search(
        dataset=dataset,
        baseline=baseline,
        baseline_result=baseline_result,
        tune_sessions=tuple(dataset.sessions),
        registry=BranchRegistry([RetrievalControlBranch(), AgenticRetrievalBranch()]),
        artifact_registry=ArtifactRegistry(tmp_path / "cache", tmp_path / "results"),
        model_discovery=None,
        run_dir=tmp_path / "run",
        search_space=space,
        budget="quick",
        profile={
            "max_stages": 2,
            "max_cost_level": "medium",
            "max_branches_per_stage": 1,
            "max_candidates_per_stage": 20,
            "tune_frontier": 5,
            "max_expensive_candidates": 0,
        },
        k=5,
        ranking_depth=20,
        evaluate=evaluate,
        diagnose=lambda _: {"regime": "candidate_coverage_bottleneck"},
    )

    agentic_event = next(event for event in search.branch_events if event["branch"] == "AgenticRetrieval")
    assert agentic_event["stage_index"] == 2
    assert agentic_event["status"] == "UNAVAILABLE"
    assert "parent retrieval identity mismatch" in str(agentic_event["reason"])
    assert not any(result.name.startswith("AgenticRetrieval:") for result in search.tune_results)


def test_agentic_evaluator_uses_only_exact_production_supplement_trace(tmp_path: Path) -> None:
    original = make_dataset(tmp_path, 1)
    session_id = next(iter(original.sessions))
    turns = list(original.sessions[session_id])
    last = turns[-1]
    turns[-1] = Turn(
        session_id=last.session_id,
        session_code=last.session_code,
        turn_index=last.turn_index,
        query_id=last.query_id,
        question=last.question,
        answer=last.answer,
        gold_raw=last.gold_raw,
        requirements=last.requirements,
        required_context="华辰公司授信额度100万元",
    )
    dataset = Dataset(original.path, original.sha256, {session_id: tuple(turns)})
    ranking_path = tmp_path / "ordinary_midterm_rankings.jsonl"
    ranking_path.write_text(
        json.dumps(
            {
                "query_id": last.query_id,
                "source_turn_id": "IRRELEVANT",
                "rank": 1,
                "memory": "ordinary MidTerm result without the required fact",
            }
        )
        + "\n",
        encoding="utf-8",
    )
    trace_path = tmp_path / "production_agentic_trace.jsonl"
    parent_config = {
        "backend": "frozen_ranking",
        "ranking_path": str(ranking_path),
        "max_total_pages": 1,
        "agentic_trace_enabled": True,
        "max_queries": 1,
        "max_total_results": 5,
        "agentic_fixed_max_iterations": 2,
        "agentic_fixed_max_tool_calls": 1,
        "agentic_fixed_max_tool_result_chars": 30000,
        **_synthetic_agentic_source("synthetic-production-source"),
    }
    _write_production_agentic_trace(trace_path, dataset, [(1, 5)], parent_config=parent_config)
    candidate = Candidate(
        "agentic-exact",
        "tune",
        {
            **parent_config,
            "production_agentic_trace_paths": [str(trace_path)],
            "production_agentic_trace_sha256": {str(trace_path.resolve()): sha256_file(trace_path)},
        },
    )
    common = {
        "dataset": dataset,
        "sessions": [session_id],
        "scope": "agentic_exact",
        "k": 5,
        "target": "midterm",
        "shortterm_window": 3,
        "ranking_depth": 20,
        "max_parallel_sessions": 1,
        "registry": ArtifactRegistry(tmp_path / "cache", tmp_path / "results"),
        "run_dir": tmp_path / "run",
    }
    evaluated = evaluate_candidate(candidate=candidate, **common)
    assert evaluated.status == "VALID"
    assert evaluated.metrics["recall_at_k"] == 1.0
    assert evaluated.metrics["agentic_contribution"] == 1.0
    assert evaluated.metrics["midterm_contribution"] == 0.0
    assert evaluated.metrics["candidate_pool_recall"] == 0.0

    missing_trace = Candidate(
        "agentic-missing",
        "tune",
        {key: value for key, value in candidate.config.items() if not key.startswith("production_agentic_trace")},
    )
    blocked = evaluate_candidate(candidate=missing_trace, **common)
    assert blocked.status == "INVALID"
    assert "production_agentic_trace is missing" in blocked.metrics["invalid_reason"]


def test_agentic_evaluator_rejects_parent_retrieval_identity_mismatch(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 1)
    session_id = next(iter(dataset.sessions))
    target = dataset.sessions[session_id][-1]
    ranking_path = tmp_path / "ordinary_midterm_parent_mismatch.jsonl"
    ranking_path.write_text(
        json.dumps({"query_id": target.query_id, "source_turn_id": "IRRELEVANT", "rank": 1}) + "\n",
        encoding="utf-8",
    )
    traced_parent = {
        "backend": "frozen_ranking",
        "max_total_pages": 4,
        "retrieval_method": "dense",
        "query_representation": "original",
        "agentic_fixed_max_tool_result_chars": 30000,
        **_synthetic_agentic_source(),
    }
    trace_path = tmp_path / "production_agentic_trace_evaluator_parent.jsonl"
    _write_production_agentic_trace(trace_path, dataset, [(1, 5)], parent_config=traced_parent)
    candidate = Candidate(
        "agentic-parent-mismatch",
        "tune",
        {
            **traced_parent,
            "ranking_path": str(ranking_path),
            "max_total_pages": 2,
            "agentic_trace_enabled": True,
            "max_queries": 1,
            "max_total_results": 5,
            "agentic_fixed_max_iterations": 2,
            "agentic_fixed_max_tool_calls": 1,
            **_agentic_trace_config([trace_path]),
        },
    )

    evaluated = evaluate_candidate(
        dataset=dataset,
        candidate=candidate,
        sessions=[session_id],
        scope="agentic_parent_mismatch",
        k=5,
        target="midterm",
        shortterm_window=3,
        ranking_depth=20,
        max_parallel_sessions=1,
        registry=ArtifactRegistry(tmp_path / "cache", tmp_path / "results"),
        run_dir=tmp_path / "run",
    )

    assert evaluated.status == "INVALID"
    assert evaluated.metrics["invalid_reason"] == "production Agentic trace parent retrieval identity mismatch"


class _StaticBranch:
    def __init__(self, name: str, *, initial: bool, priority: int, gain: float):
        self.gain = gain
        self.spec = BranchSpec(
            name=name,
            diagnostic_regimes=ALL_REGIMES,
            cost_level="cheap" if initial else "medium",
            required_artifacts=("checkpoint",),
            execution_adapter="test",
            provenance_contract=("dataset_sha256",),
            resource_requirements={},
            priority=priority,
            initial_stage=initial,
        )

    def generate(self, context: Any) -> BranchOutcome:
        candidate = Candidate(
            name=self.spec.name.lower(),
            stage=f"stage_{context.stage_index}",
            config={**context.anchor.config, "gain": self.gain, "branch_cost_level": self.spec.cost_level},
        )
        return BranchOutcome(self.spec.name, "READY", [candidate])

    def validate_provenance(self, candidate: Candidate, context: Any) -> tuple[bool, None]:
        del candidate, context
        return True, None


def test_staged_loop_rediagnoses_and_never_uses_validation(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 3)
    baseline = Candidate(name="baseline", stage="baseline", config={"backend": "production_midterm", "gain": 0.0})

    def measured(candidate: Candidate) -> CandidateResult:
        recall = 0.40 + float(candidate.config.get("gain") or 0.0)
        value = result(candidate.name, recall)
        value.candidate_hash = candidate.name
        value.stage = candidate.stage
        value.config = candidate.config
        value.metrics["recall_at_4k"] = min(1.0, recall + 0.30)
        return value

    baseline_result = measured(baseline)
    scopes: list[str] = []

    def evaluate(
        candidates: list[Candidate] | tuple[Candidate, ...], sessions: Any, scope: str
    ) -> list[CandidateResult]:
        del sessions
        scopes.append(scope)
        return [measured(candidate) for candidate in candidates]

    diagnoses: list[str] = []

    def diagnose(candidate_result: CandidateResult) -> dict[str, Any]:
        diagnoses.append(candidate_result.name)
        return {
            "regime": "ranking_bottleneck",
            "recall_4k_minus_k_pp": 30.0,
            "recall_4k_percent": 80.0,
            "stddev_pp": 0.0,
        }

    registry = BranchRegistry(
        [
            _StaticBranch("Cheap", initial=True, priority=1, gain=0.10),
            _StaticBranch("Secondary", initial=False, priority=2, gain=0.15),
        ]
    )
    search = run_staged_search(
        dataset=dataset,
        baseline=baseline,
        baseline_result=baseline_result,
        tune_sessions=tuple(dataset.sessions),
        registry=registry,
        artifact_registry=None,
        model_discovery=None,
        run_dir=tmp_path,
        search_space={"selection": {"tie_tolerance_pp": 0.25, "min_improvement_pp": 0.25, "patience_stages": 2}},
        budget="quick",
        profile={
            "max_stages": 4,
            "max_cost_level": "medium",
            "max_branches_per_stage": 1,
            "max_candidates_per_stage": 3,
            "validation_frontier": 2,
        },
        k=5,
        ranking_depth=20,
        evaluate=evaluate,
        diagnose=diagnose,
    )
    assert [stage["branches"] for stage in search.stage_history] == [["Cheap"], ["Secondary"]]
    assert len(search.diagnostics) == 3
    assert diagnoses == ["baseline", "cheap", "secondary"]
    assert scopes == ["stage_1_tune", "stage_2_tune"]
    assert search.stop_reason == "converged_after_relevant_branch_coverage"
    assert search.frontier[0].name == "secondary"


def _model_index(*, name: str, task: str, dataset: str, score: float) -> list[dict[str, Any]]:
    return [
        {
            "name": name,
            "results": [
                {
                    "task": {"type": task},
                    "dataset": {"name": dataset},
                    "metrics": [{"type": "ndcg_at_10", "value": score}],
                }
            ],
        }
    ]


@pytest.mark.parametrize(
    ("model_type", "local_id", "online_id", "task", "dataset", "tags"),
    [
        (
            "embedding",
            "local/zh-embedding",
            "dynamic/high-zh-embedding",
            "Retrieval",
            "T2Retrieval",
            ["sentence-similarity", "zh"],
        ),
        (
            "reranker",
            "local/zh-reranker",
            "dynamic/high-zh-reranker",
            "Reranking",
            "T2Reranking",
            ["text-ranking", "zh", "reranker"],
        ),
    ],
)
def test_deep_model_discovery_unifies_cached_and_online_quality_ranking(
    tmp_path: Path,
    model_type: str,
    local_id: str,
    online_id: str,
    task: str,
    dataset: str,
    tags: list[str],
) -> None:
    cache = tmp_path / "hf" / "hub"
    snapshot = cache / f"models--{local_id.replace('/', '--')}" / "snapshots" / "abc123"
    snapshot.mkdir(parents=True)
    (snapshot / "README.md").write_text(
        "---\n"
        + json.dumps(
            {
                "license": "apache-2.0",
                "tags": tags,
                "model-index": _model_index(name="C-MTEB", task=task, dataset=dataset, score=0.3),
            },
            ensure_ascii=False,
        )
        + "\n---\n",
        encoding="utf-8",
    )

    class FakeApi:
        def list_models(self, **kwargs: Any) -> list[Any]:
            return [
                SimpleNamespace(
                    modelId=online_id,
                    tags=tags,
                    sha="online-revision",
                    downloads=100,
                    cardData={
                        "license": "mit",
                        "model-index": _model_index(name="C-MTEB", task=task, dataset=dataset, score=0.8),
                    },
                    safetensors={"total": 20_000_000},
                )
            ]

    discovery = ModelDiscovery(
        output_path=tmp_path / "models.json",
        resources=ResourceEnvelope(0, None, 16.0, 100.0),
        cache_root=cache,
        api=FakeApi(),
    )
    candidates = discovery.discover(
        model_type=model_type,
        allow_network=True,
        general_limit=1,
        finance_limit=0,
    )
    assert [candidate.model_id for candidate in candidates] == [online_id]
    assert candidates[0].cache_status == "NOT_CHECKED"
    assert "cache affects download cost only" in candidates[0].selection_reason
    assert candidates[0].metadata_evidence["benchmark_scores"] == [
        {
            "benchmark": "C-MTEB",
            "task": task,
            "dataset": dataset,
            "metric": "ndcg_at_10",
            "score": 0.8,
        }
    ]


def test_standard_model_discovery_is_cache_only(tmp_path: Path) -> None:
    cache = tmp_path / "hf" / "hub"
    snapshot = cache / "models--local--zh-embedding" / "snapshots" / "abc123"
    snapshot.mkdir(parents=True)

    class OfflineApi:
        def list_models(self, **kwargs: Any) -> list[Any]:
            raise AssertionError(f"standard discovery must not call the network: {kwargs}")

    discovery = ModelDiscovery(
        output_path=tmp_path / "models.json",
        resources=ResourceEnvelope(0, None, 16.0, 100.0),
        cache_root=cache,
        api=OfflineApi(),
    )
    candidates = discovery.discover(
        model_type="embedding",
        allow_network=False,
        general_limit=1,
        finance_limit=0,
    )
    assert [candidate.model_id for candidate in candidates] == ["local/zh-embedding"]
    assert candidates[0].cache_status == "CACHED"


def test_model_discovery_deduplicates_exact_cached_online_revision(tmp_path: Path) -> None:
    cache = tmp_path / "hf" / "hub"
    snapshot = cache / "models--shared--zh-embedding" / "snapshots" / "abc123"
    snapshot.mkdir(parents=True)

    class FakeApi:
        def list_models(self, **_: Any) -> list[Any]:
            return [
                SimpleNamespace(
                    modelId="shared/zh-embedding",
                    tags=["sentence-similarity", "zh"],
                    sha="abc123",
                    downloads=10,
                    cardData={
                        "license": "apache-2.0",
                        "model-index": _model_index(
                            name="C-MTEB",
                            task="Retrieval",
                            dataset="T2Retrieval",
                            score=0.7,
                        ),
                    },
                )
            ]

    discovery = ModelDiscovery(
        output_path=tmp_path / "models.json",
        resources=ResourceEnvelope(0, None, 16.0, 100.0),
        cache_root=cache,
        api=FakeApi(),
    )
    candidates = discovery.discover(
        model_type="embedding",
        allow_network=True,
        general_limit=2,
        finance_limit=0,
    )
    assert len(candidates) == 1
    assert candidates[0].revision == "abc123"
    assert candidates[0].cache_status == "CACHED"
    assert set(candidates[0].source.split("+")) == {"huggingface_search", "local_huggingface_cache"}


def test_model_benchmark_scores_are_structured_and_compared_only_like_for_like() -> None:
    def candidate(model_id: str, benchmark: str, dataset: str, score: float) -> ModelCandidate:
        evidence = _metadata_evidence(
            {
                "tags": ["sentence-similarity", "zh"],
                "cardData": {
                    "model-index": _model_index(
                        name=benchmark,
                        task="Retrieval",
                        dataset=dataset,
                        score=score,
                    )
                },
            }
        )
        return ModelCandidate(model_id, "embedding", "test", revision="revision", metadata_evidence=evidence)

    low = candidate("test/c-mteb-low", "C-MTEB", "T2Retrieval", 0.6)
    high = candidate("test/c-mteb-high", "C-MTEB", "T2Retrieval", 75.0)
    finance = candidate("test/finmteb", "FinMTEB", "FinQA", 99.0)
    _annotate_relative_benchmark_scores([low, high, finance])

    assert low.metadata_evidence["benchmark_scores"][0] == {
        "benchmark": "C-MTEB",
        "task": "Retrieval",
        "dataset": "T2Retrieval",
        "metric": "ndcg_at_10",
        "score": 0.6,
    }
    assert low.metadata_evidence["benchmark_comparisons"][0]["normalized_score"] == pytest.approx(0.6)
    assert high.metadata_evidence["benchmark_comparisons"][0]["normalized_score"] == pytest.approx(0.75)
    assert high.metadata_evidence["benchmark_comparisons"][0]["relative_percentile"] == pytest.approx(1.0)
    assert finance.metadata_evidence["benchmark_comparisons"] == []


def test_model_score_does_not_reward_benchmark_entry_count() -> None:
    common_result = {
        "task": {"type": "Retrieval"},
        "dataset": {"name": "T2Retrieval"},
        "metrics": [{"type": "ndcg_at_10", "value": 0.7}],
    }

    def candidate(model_id: str, *, extra_entries: int) -> ModelCandidate:
        results = [common_result]
        results.extend(
            {
                "task": {"type": "Retrieval"},
                "dataset": {"name": f"PrivateDataset{index}"},
                "metrics": [{"type": "ndcg_at_10", "value": 0.7}],
            }
            for index in range(extra_entries)
        )
        evidence = _metadata_evidence(
            {
                "tags": ["sentence-similarity", "zh"],
                "cardData": {"license": "apache-2.0", "model-index": [{"name": "C-MTEB", "results": results}]},
            }
        )
        return ModelCandidate(
            model_id,
            "embedding",
            "test",
            revision="revision",
            license="apache-2.0",
            tags=["sentence-similarity", "zh"],
            metadata_evidence=evidence,
        )

    concise = candidate("test/concise-card", extra_entries=0)
    verbose = candidate("test/verbose-card", extra_entries=8)
    _annotate_relative_benchmark_scores([concise, verbose])

    assert concise.metadata_evidence["quality_benchmark_score_count"] == 1
    assert verbose.metadata_evidence["quality_benchmark_score_count"] == 9
    assert _candidate_score_value(concise, finance=False) == pytest.approx(
        _candidate_score_value(verbose, finance=False)
    )


def test_exact_cached_revision_reuses_snapshot_without_download(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    snapshot = tmp_path / "hf" / "hub" / "models--cached--embedding" / "snapshots" / "abc123"
    snapshot.mkdir(parents=True)
    discovery = ModelDiscovery(
        output_path=tmp_path / "models.json",
        resources=ResourceEnvelope(0, None, 16.0, 100.0),
        cache_root=tmp_path / "hf" / "hub",
    )

    def unexpected_download(**_: Any) -> str:
        raise AssertionError("an exact cached model_id + revision must not call snapshot_download")

    monkeypatch.setattr("huggingface_hub.snapshot_download", unexpected_download)
    available = discovery.ensure_available(
        ModelCandidate(
            "cached/embedding",
            "embedding",
            "local_huggingface_cache",
            revision="abc123",
            local_path=str(snapshot),
            cache_status="CACHED",
        ),
        allow_download=True,
    )
    assert available.status == "AVAILABLE"
    assert available.resource_usage["cache_reused_without_download"] is True


def test_model_download_failure_is_isolated(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    discovery = ModelDiscovery(
        output_path=tmp_path / "models.json",
        resources=ResourceEnvelope(0, None, 16.0, 100.0),
        cache_root=tmp_path / "hf" / "hub",
    )

    def unavailable(**_: Any) -> str:
        raise PermissionError("gated model")

    monkeypatch.setattr("huggingface_hub.snapshot_download", unavailable)
    failed = discovery.ensure_available(
        ModelCandidate("gated/model", "embedding", "huggingface_search"),
        allow_download=True,
    )
    assert failed.status == "UNAVAILABLE"
    assert "gated model" in str(failed.error)
    payload = json.loads((tmp_path / "models.json").read_text(encoding="utf-8"))
    assert any(item.get("status") == "UNAVAILABLE" for item in payload["models"])


def test_budget_profiles_gate_real_cost_levels() -> None:
    import yaml

    space = yaml.safe_load((SCRIPTS.parent / "search_space.yaml").read_text(encoding="utf-8"))
    profiles = space["budget"]["profiles"]
    assert profiles["quick"]["max_cost_level"] == "medium"
    assert profiles["standard"]["max_cost_level"] == "high"
    assert profiles["deep"]["max_cost_level"] == "expensive"
    assert profiles["quick"]["max_stages"] < profiles["deep"]["max_stages"]


def _production_branch_inputs(
    tmp_path: Path,
    dataset: Dataset,
    *,
    short_term_capacity: int = 6,
) -> tuple[Candidate, Path]:
    tmp_path.mkdir(parents=True, exist_ok=True)
    memory_config = tmp_path / f"memory-{short_term_capacity}.json"
    memory_config.write_text(
        json.dumps(
            {
                "llm": {"provider": "mock", "config": {"model": "test-model"}},
                "midterm": {"enabled": True, "short_term_capacity": short_term_capacity},
            }
        ),
        encoding="utf-8",
    )
    manifests: list[Path] = []
    for session_id, turns in dataset.sessions.items():
        source_dir = tmp_path / "production" / session_id
        source_dir.mkdir(parents=True, exist_ok=True)
        checkpoints = source_dir / "production_midterm_checkpoints.jsonl"
        checkpoint_rows = [
            {
                "query_id": turn.query_id,
                "query": turn.question,
                "query_vector": [1.0, 0.0],
                "filters": {"run_id": session_id},
                "pages": [],
                "sessions": [],
                "source_turn_ids_by_job": {},
            }
            for turn in turns
        ]
        checkpoints.write_text(
            "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in checkpoint_rows),
            encoding="utf-8",
        )
        manifest = source_dir / "production_midterm_manifest.json"
        manifest.write_text(
            json.dumps(
                {
                    "schema": ADAPTER_SCHEMA,
                    "status": "COMPLETE",
                    "dataset_sha256": dataset.sha256,
                    "session_id": session_id,
                    "memory_config_path": str(memory_config),
                    "memory_config_sha256": sha256_file(memory_config),
                    "shortterm_qa_turns": short_term_capacity // 2,
                    "llm_mode": "mock",
                    "checkpoints_path": str(checkpoints),
                    "checkpoints_sha256": sha256_file(checkpoints),
                    "prompt_hashes": {"page_summary": "production"},
                    "production_config": {
                        "short_term_capacity": short_term_capacity,
                        "top_k_sessions": 5,
                        "top_k_pages": 5,
                        "max_total_pages": 5,
                    },
                    "effective_memory_config": {
                        "vector_store": {"config": {"bm25_language": "en"}},
                        "agentic_retrieval": {
                            "max_iterations": 2,
                            "max_tool_calls": 1,
                            "max_queries": 3,
                            "max_total_results": 6,
                            "max_tool_result_chars": 30000,
                        },
                    },
                    "failed_turns": 0,
                    "llm_calls": 0,
                    "embedding_calls": 0,
                }
            ),
            encoding="utf-8",
        )
        manifests.append(manifest)
    candidate = Candidate(
        name="baseline",
        stage="baseline",
        config={
            "backend": "production_midterm",
            "retrieval_contract": "production_midterm_v1",
            "retrieval_method": "dense",
            "query_representation": "original",
            "page_representation": "production",
            "top_k_sessions": 5,
            "top_k_pages": 5,
            "max_total_pages": 5,
            "max_queries": 3,
            "max_total_results": 5,
            "agentic_fixed_max_iterations": 2,
            "agentic_fixed_max_tool_calls": 1,
            "agentic_fixed_max_tool_result_chars": 30000,
            "manifest_paths": [str(manifest) for manifest in manifests],
            "manifest_sha256": {str(manifest.resolve()): sha256_file(manifest) for manifest in manifests},
        },
    )
    return candidate, manifests[0]


def _branch_context(
    tmp_path: Path,
    dataset: Dataset,
    baseline: Candidate,
    *,
    budget: str = "standard",
    generation_round: int = 1,
    model_discovery: Any = None,
    search_space: dict[str, Any] | None = None,
) -> BranchContext:
    return BranchContext(
        dataset=dataset,
        baseline=baseline,
        anchor=baseline,
        anchor_result=result("baseline", 0.4),
        diagnostic={
            "regime": "candidate_coverage_bottleneck",
            "recall_4k_minus_k_pp": 30.0,
        },
        search_space=search_space
        or {
            "search": {
                "stages": {
                    "secondary": {
                        "query_representation": {"max_rounds": 3, "variants_per_round": 3},
                        "reranking": {"min_deep_recall_gap_pp": 15.0},
                    },
                    "expensive": {"require_frontier_candidate": True, "max_prompt_candidates": 4},
                }
            }
        },
        budget=budget,
        k=5,
        ranking_depth=20,
        tune_sessions=(next(iter(dataset.sessions)),),
        registry=ArtifactRegistry(tmp_path / "cache", tmp_path / "results"),
        run_dir=tmp_path / "run",
        model_discovery=model_discovery,
        stage_index=generation_round,
        generation_round=generation_round,
        execution_settings={"max_parallel_sessions": 2, "max_parallel_llm_calls": 2},
    )


def test_fine_grained_longterm_prompt_candidates_change_real_extraction_and_source_identity(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    dataset = make_dataset(tmp_path, 1)
    baseline, _ = _production_branch_inputs(tmp_path, dataset)
    outcome = FineGrainedLongtermExtractionPromptBranch().generate(
        _branch_context(tmp_path, dataset, baseline, budget="deep")
    )
    expected_variants = controlled_fine_grained_longterm_prompt_variants()

    assert outcome.status == "READY"
    assert len(outcome.candidates) == len(expected_variants) == 2
    first, second = outcome.candidates
    prompt_a = first.config["fine_grained_longterm_extraction_prompt"]
    prompt_b = second.config["fine_grained_longterm_extraction_prompt"]
    identity_a = first.config["source_generation_spec"]["source_identity"]
    identity_b = second.config["source_generation_spec"]["source_identity"]
    assert (
        first.config["fine_grained_longterm_extraction_prompt_hash"]
        != second.config["fine_grained_longterm_extraction_prompt_hash"]
    )
    assert identity_a != identity_b
    assert production_prompt_hashes(fine_grained_longterm_extraction_prompt=prompt_a) != production_prompt_hashes(
        fine_grained_longterm_extraction_prompt=prompt_b
    )

    monkeypatch.setattr(memory_main, "capture_event", lambda *args, **kwargs: None)

    class PromptAwareLLM:
        def __init__(self, expected_prompt: str, output: str) -> None:
            self.expected_prompt = expected_prompt
            self.output = output
            self.system_prompts: list[str] = []

        def generate_response(self, *, messages: list[dict[str, Any]], **kwargs: Any) -> str:
            del kwargs
            system_prompt = str(messages[0]["content"])
            self.system_prompts.append(system_prompt)
            assert system_prompt == self.expected_prompt
            return json.dumps({"memory": [{"text": self.output}]}, ensure_ascii=False)

    class Embedding:
        def embed(self, text: str, mode: str) -> list[float]:
            del text, mode
            return [1.0, 0.0]

        def embed_batch(self, texts: list[str], mode: str) -> list[list[float]]:
            del mode
            return [[1.0, 0.0] for _ in texts]

    class VectorStore:
        def __init__(self) -> None:
            self.payloads: list[dict[str, Any]] = []

        def search(self, **kwargs: Any) -> list[Any]:
            del kwargs
            return []

        def get(self, **kwargs: Any) -> None:
            del kwargs
            return None

        def insert(self, *, vectors: list[Any], ids: list[str], payloads: list[dict[str, Any]]) -> None:
            assert len(vectors) == len(ids) == len(payloads)
            self.payloads.extend(payloads)

    def extract(prompt: str, output: str) -> tuple[list[dict[str, Any]], PromptAwareLLM, VectorStore]:
        db = SQLiteManager(":memory:")
        memory = Memory.__new__(Memory)
        delegate = PromptAwareLLM(prompt, output)
        store = VectorStore()
        memory.config = SimpleNamespace(
            midterm=SimpleNamespace(enabled=False, short_term_capacity=20),
            fine_grained_longterm=SimpleNamespace(extraction_prompt=prompt),
        )
        memory.db = db
        memory.llm = delegate
        memory.embedding_model = Embedding()
        memory.vector_store = store
        memory.custom_instructions = None
        memory._bm25_language = "en"
        memory._run_entity_extraction = lambda function, texts: [[] for _ in texts]
        memory.api_version = "v1.1"
        try:
            result = Memory._process_evicted_long_term_memories(
                memory,
                [{"role": "user", "content": "raw-Q"}, {"role": "assistant", "content": "raw-A"}],
                {"user_id": "u1", "run_id": "s1", "source_turn_index": 1},
                {"user_id": "u1", "run_id": "s1"},
                infer=True,
            )
            return result, delegate, store
        finally:
            db.close()

    output_a, delegate_a, store_a = extract(prompt_a, "memory-A")
    output_b, delegate_b, store_b = extract(prompt_b, "memory-B")

    assert output_a != output_b
    assert [item["memory"] for item in output_a] == ["memory-A"]
    assert [item["memory"] for item in output_b] == ["memory-B"]
    assert delegate_a.system_prompts == [prompt_a]
    assert delegate_b.system_prompts == [prompt_b]
    assert [payload["data"] for payload in store_a.payloads] == ["memory-A"]
    assert [payload["data"] for payload in store_b.payloads] == ["memory-B"]
    assert all(payload["data"] not in {"raw-Q", "raw-A"} for payload in [*store_a.payloads, *store_b.payloads])


def test_new_dataset_generates_three_query_variants_and_resumes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import tuner.derived_artifacts as derived_artifacts

    dataset = make_dataset(tmp_path, 2)
    baseline, _ = _production_branch_inputs(tmp_path, dataset)

    class FakeEncoder:
        def __init__(self, manifest_path: Path):
            del manifest_path

        def encode(self, texts: list[str], *, action: str) -> list[list[float]]:
            del action
            return [[float(len(text)), 1.0] for text in texts]

    class FakeLLM:
        def __init__(self) -> None:
            self.calls = 0
            self.active = 0
            self.max_active = 0
            self.lock = threading.Lock()

        def generate_response(self, messages: list[dict[str, str]], **_: Any) -> str:
            payload = json.loads(messages[-1]["content"])
            with self.lock:
                self.calls += 1
                self.active += 1
                self.max_active = max(self.max_active, self.active)
            time.sleep(0.002)
            with self.lock:
                self.active -= 1
            return json.dumps({"resolved_query": f"{payload['current_query']} 已消解"}, ensure_ascii=False)

    llm = FakeLLM()
    monkeypatch.setattr(derived_artifacts, "_ProductionEncoder", FakeEncoder)
    monkeypatch.setattr(QueryPromptArtifactGenerator, "_create_llm", lambda self, config, mode: llm)
    context = _branch_context(tmp_path, dataset, baseline)
    branch = QueryRepresentationBranch()
    first = BranchRegistry([branch]).generate(branch, context)
    assert first.status == "READY"
    assert len(first.candidates) == 3
    assert {item.config["query_prompt_generation_round"] for item in first.candidates} == {1}
    assert first.provenance["analysis_session_ids"] == [context.tune_sessions[0]]
    assert all(Path(item.config["query_artifact_path"]).exists() for item in first.candidates)
    assert llm.max_active <= 2
    calls_after_first = llm.calls

    second = BranchRegistry([branch]).generate(branch, context)
    assert second.status == "READY"
    assert len(second.candidates) == 3
    assert second.llm_calls == 0
    assert llm.calls == calls_after_first
    payload = json.loads(Path(first.candidates[0].config["query_artifact_path"]).read_text(encoding="utf-8"))["payload"]
    assert payload["analysis_session_ids"] == [context.tune_sessions[0]]
    assert set(payload["generated_session_ids"]) == set(dataset.sessions)


@pytest.mark.parametrize(("capacity_messages", "expected_history"), [(6, (2, 3, 4)), (8, (1, 2, 3, 4))])
def test_query_history_uses_exact_production_shortterm_window(
    tmp_path: Path,
    capacity_messages: int,
    expected_history: tuple[int, ...],
) -> None:
    session_id = "S001_历史窗口"
    turns = tuple(make_turn(session_id, index) for index in range(6))
    dataset = Dataset(
        path=str(tmp_path / "history.xlsx"),
        sha256="b" * 64,
        sessions={session_id: turns},
    )
    baseline, _ = _production_branch_inputs(
        tmp_path / f"capacity-{capacity_messages}",
        dataset,
        short_term_capacity=capacity_messages,
    )
    captured: dict[str, dict[str, Any]] = {}

    class CapturingLLM:
        def generate_response(self, messages: list[dict[str, str]], **_: Any) -> str:
            request = json.loads(messages[-1]["content"])
            captured[request["current_query"]] = request
            return json.dumps({"resolved_query": request["current_query"]}, ensure_ascii=False)

    generator = QueryPromptArtifactGenerator(
        ArtifactRegistry(tmp_path / f"cache-{capacity_messages}", tmp_path / "legacy"),
        llm_factory=lambda _config, _mode: CapturingLLM(),
    )
    artifact = generator.generate(
        dataset=dataset,
        anchor=baseline,
        variant=QueryPromptVariant(
            prompt_text="只使用 production 可见历史",
            prompt_hash="query-history-contract",
            parent_prompt_hash="production",
            generation_round=1,
            optimization_direction="explicit_coreference_resolution",
        ),
        tune_sessions=(session_id,),
        max_parallel_llm_calls=1,
    )

    current = turns[4]
    request = captured[current.question]
    assert set(request) == {"current_query", "recent_history"}
    assert [row["content"] for row in request["recent_history"] if row["role"] == "user"] == [
        turns[index - 1].question for index in expected_history
    ]
    assert [row["content"] for row in request["recent_history"] if row["role"] == "assistant"] == [
        turns[index - 1].answer for index in expected_history
    ]
    visible_questions = [row["content"] for row in request["recent_history"] if row["role"] == "user"]
    if capacity_messages == 8:
        assert turns[0].question in visible_questions
    else:
        assert turns[0].question not in visible_questions
    assert turns[5].question not in json.dumps(request, ensure_ascii=False)
    assert artifact.identity["shortterm_capacity_messages"] == capacity_messages
    assert artifact.identity["shortterm_qa_turns"] == capacity_messages // 2
    assert artifact.identity["history_policy"] == PRODUCTION_SHORTTERM_HISTORY_POLICY
    assert artifact.identity["production_config_hash"]
    stored = json.loads(artifact.path.read_text(encoding="utf-8"))["payload"]
    assert stored["shortterm_qa_turns"] == capacity_messages // 2
    assert stored["history_policy"] == PRODUCTION_SHORTTERM_HISTORY_POLICY


@pytest.mark.parametrize(
    ("invalid_mode", "error_match"),
    [
        ("missing_capacity", "short_term_capacity is required"),
        ("odd_capacity", "positive even message count"),
        ("manifest_mismatch", "manifest/config ShortTerm mismatch"),
    ],
)
def test_query_history_rejects_invalid_or_inconsistent_production_window(
    tmp_path: Path,
    invalid_mode: str,
    error_match: str,
) -> None:
    dataset = make_dataset(tmp_path, 1)
    baseline, manifest_path = _production_branch_inputs(tmp_path, dataset)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    config_path = Path(manifest["memory_config_path"])
    memory_config = json.loads(config_path.read_text(encoding="utf-8"))
    if invalid_mode == "missing_capacity":
        memory_config["midterm"].pop("short_term_capacity")
    elif invalid_mode == "odd_capacity":
        memory_config["midterm"]["short_term_capacity"] = 7
    else:
        manifest["shortterm_qa_turns"] = 4
    config_path.write_text(json.dumps(memory_config), encoding="utf-8")
    manifest["memory_config_sha256"] = sha256_file(config_path)
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    baseline.config["manifest_sha256"] = {str(manifest_path.resolve()): sha256_file(manifest_path)}

    with pytest.raises(ValueError, match=error_match):
        QueryPromptArtifactGenerator._production_shortterm_contract(baseline)


def test_query_prompt_variants_are_capped_at_three_rounds_and_three_per_round(tmp_path: Path) -> None:
    dataset = make_dataset(tmp_path, 2)
    baseline = Candidate(name="baseline", stage="baseline", config={"query_representation": "original"})
    baseline_result = result("baseline", 0.4)
    for generation_round in (1, 2, 3):
        variants = controlled_query_prompt_variants(
            dataset=dataset,
            anchor=baseline,
            anchor_result=baseline_result,
            tune_sessions=(next(iter(dataset.sessions)),),
            generation_round=generation_round,
            variants_per_round=99,
        )
        assert len(variants) == 3
        assert {variant.generation_round for variant in variants} == {generation_round}
    assert not controlled_query_prompt_variants(
        dataset=dataset,
        anchor=baseline,
        anchor_result=baseline_result,
        tune_sessions=tuple(dataset.sessions),
        generation_round=4,
        variants_per_round=3,
    )


class _RoundBranch:
    def __init__(self, name: str = "CoarseRefine", *, improve: bool = True):
        self.improve = improve
        self.spec = BranchSpec(
            name=name,
            diagnostic_regimes=ALL_REGIMES,
            cost_level="cheap",
            required_artifacts=("checkpoint",),
            execution_adapter="test",
            provenance_contract=("dataset",),
            resource_requirements={},
            priority=1,
            initial_stage=True,
        )

    def generate(self, context: BranchContext) -> BranchOutcome:
        gain = 0.05 * context.generation_round if self.improve else 0.0
        candidate = Candidate(
            name=f"{self.spec.name}-round-{context.generation_round}",
            stage=f"stage-{context.stage_index}",
            config={
                **context.anchor.config,
                "coarse_refine_round": context.generation_round,
                "gain": gain,
                "branch_cost_level": "cheap",
            },
        )
        return BranchOutcome(self.spec.name, "READY", [candidate])

    def validate_provenance(self, candidate: Candidate, context: BranchContext) -> tuple[bool, None]:
        del candidate, context
        return True, None


def _run_round_search(tmp_path: Path, branch: Any, *, max_rounds: int) -> Any:
    dataset = make_dataset(tmp_path, 3)
    baseline = Candidate(name="baseline", stage="baseline", config={"gain": 0.0})

    def evaluate(candidates: Any, sessions: Any, scope: str) -> list[CandidateResult]:
        del sessions, scope
        return [result(candidate.name, 0.4 + float(candidate.config.get("gain") or 0.0)) for candidate in candidates]

    return run_staged_search(
        dataset=dataset,
        baseline=baseline,
        baseline_result=result("baseline", 0.4),
        tune_sessions=tuple(dataset.sessions),
        registry=BranchRegistry([branch]),
        artifact_registry=None,
        model_discovery=None,
        run_dir=tmp_path,
        search_space={
            "selection": {"min_improvement_pp": 0.25, "patience_stages": 2},
            "search": {"branch_registry": {branch.spec.name: {"enabled": True, "max_rounds": max_rounds}}},
        },
        budget="standard",
        profile={"max_stages": 5, "max_cost_level": "high", "max_candidates_per_stage": 4},
        k=5,
        ranking_depth=20,
        evaluate=evaluate,
        diagnose=lambda _: {"regime": "ranking_bottleneck"},
    )


def test_same_branch_can_reenter_for_coarse_refine(tmp_path: Path) -> None:
    search = _run_round_search(tmp_path, _RoundBranch(), max_rounds=2)
    assert [event["generation_round"] for event in search.branch_events if event["branch"] != "__search_policy__"] == [
        1,
        2,
    ]
    assert search.frontier[0].name == "CoarseRefine-round-2"


def test_query_branch_stops_early_without_improvement(tmp_path: Path) -> None:
    search = _run_round_search(tmp_path, _RoundBranch("QueryRepresentation", improve=False), max_rounds=3)
    assert [event["generation_round"] for event in search.branch_events if event["branch"] != "__search_policy__"] == [
        1
    ]


def test_query_branch_never_exceeds_three_rounds(tmp_path: Path) -> None:
    search = _run_round_search(tmp_path, _RoundBranch("QueryRepresentation", improve=True), max_rounds=3)
    assert [event["generation_round"] for event in search.branch_events if event["branch"] != "__search_policy__"] == [
        1,
        2,
        3,
    ]


class _CoverageBranch:
    def __init__(self, name: str, *, cost_level: str, priority: int, initial: bool = False):
        self.spec = BranchSpec(
            name=name,
            diagnostic_regimes=frozenset({"candidate_coverage_bottleneck"}),
            cost_level=cost_level,
            required_artifacts=("checkpoint",),
            execution_adapter="test",
            provenance_contract=("dataset",),
            resource_requirements={},
            priority=priority,
            initial_stage=initial,
        )

    def generate(self, context: BranchContext) -> BranchOutcome:
        candidate = Candidate(
            name=f"{self.spec.name}-round-{context.generation_round}",
            stage=f"stage-{context.stage_index}",
            config={
                **context.anchor.config,
                f"coverage_{self.spec.name}": context.generation_round,
                "branch_cost_level": self.spec.cost_level,
                "gain": 0.0,
            },
        )
        return BranchOutcome(self.spec.name, "READY", [candidate])

    def validate_provenance(self, candidate: Candidate, context: BranchContext) -> tuple[bool, None]:
        del candidate, context
        return True, None


def _run_coverage_search(
    tmp_path: Path,
    *,
    max_stages: int = 8,
    max_expensive_candidates: int = 20,
) -> Any:
    import yaml

    dataset = make_dataset(tmp_path, 3)
    baseline = Candidate(name="baseline", stage="baseline", config={"gain": 0.0})
    branches = [
        _CoverageBranch("RetrievalControl", cost_level="cheap", priority=10, initial=True),
        _CoverageBranch("HybridRetrieval", cost_level="medium", priority=20),
        _CoverageBranch("Embedding", cost_level="high", priority=30),
        _CoverageBranch("QueryRepresentation", cost_level="high", priority=40),
        _CoverageBranch("PageRepresentation", cost_level="high", priority=50),
        _CoverageBranch("FieldAwareMultiVector", cost_level="high", priority=60),
        _CoverageBranch("MidtermPageSummaryPrompt", cost_level="expensive", priority=70),
    ]
    search_space = yaml.safe_load((SCRIPTS.parent / "search_space.yaml").read_text(encoding="utf-8"))
    search_space["selection"]["patience_stages"] = 1

    def evaluate(candidates: Any, sessions: Any, scope: str) -> list[CandidateResult]:
        del sessions, scope
        return [result(candidate.name, 0.4 + float(candidate.config.get("gain") or 0.0)) for candidate in candidates]

    return run_staged_search(
        dataset=dataset,
        baseline=baseline,
        baseline_result=result("baseline", 0.4),
        tune_sessions=tuple(dataset.sessions),
        registry=BranchRegistry(branches),
        artifact_registry=None,
        model_discovery=None,
        run_dir=tmp_path,
        search_space=search_space,
        budget="deep",
        profile={
            "max_stages": max_stages,
            "max_cost_level": "expensive",
            "max_branches_per_stage": 1,
            "max_candidates_per_stage": 8,
            "screening_sessions": 1,
            "tune_frontier": 8,
            "max_expensive_candidates": max_expensive_candidates,
        },
        k=5,
        ranking_depth=20,
        evaluate=evaluate,
        diagnose=lambda _: {"regime": "candidate_coverage_bottleneck"},
    )


def test_deep_patience_waits_for_relevant_coverage_and_source_prompt(tmp_path: Path) -> None:
    search = _run_coverage_search(tmp_path)
    attempted = [event["branch"] for event in search.branch_events if event["branch"] != "__search_policy__"]
    assert attempted == [
        "RetrievalControl",
        "HybridRetrieval",
        "Embedding",
        "QueryRepresentation",
        "PageRepresentation",
        "FieldAwareMultiVector",
        "MidtermPageSummaryPrompt",
    ]
    assert any(event["status"] == "PATIENCE_SOFT_EXHAUSTED" for event in search.branch_events)
    assert search.stop_reason == "converged_after_relevant_branch_coverage"
    assert search.coverage_audit["remaining_branches"] == []
    assert search.coverage_audit["relevant_coverage_complete"] is True
    assert set(search.coverage_audit["attempted_branches"]) == set(attempted)
    assert all(stage.get("coverage_after") for stage in search.stage_history)
    stop_event = search.branch_events[-1]
    assert stop_event["status"] == "SEARCH_STOP"
    assert stop_event["provenance"]["coverage"]["exhausted_branches"]


def test_stage_and_expensive_resource_budgets_remain_hard_stops(tmp_path: Path) -> None:
    stage_limited = _run_coverage_search(tmp_path / "stage", max_stages=2)
    assert stage_limited.stop_reason == "stage_budget_exhausted"
    assert "QueryRepresentation" in stage_limited.coverage_audit["remaining_branches"]

    resource_limited = _run_coverage_search(
        tmp_path / "resource",
        max_stages=8,
        max_expensive_candidates=0,
    )
    assert resource_limited.stop_reason == "resource_budget_exhausted"
    blocked = {row["branch"]: row["reason"] for row in resource_limited.coverage_audit["blocked_branches"]}
    assert blocked["MidtermPageSummaryPrompt"] == "max_expensive_candidates exhausted"
    assert any(
        row["branch"] == "MidtermPageSummaryPrompt" and row["status"] == "BUDGET_OR_RESOURCE_BLOCKED"
        for row in resource_limited.skipped_branches
    )


def test_candidate_config_hash_deduplicates_branch_bookkeeping() -> None:
    left = {
        "top_k_pages": 10,
        "experiment_branch": "A",
        "branch_cost_level": "cheap",
        "parent_candidate_hash": "one",
        "applied_branches": ["A"],
    }
    right = {
        "top_k_pages": 10,
        "experiment_branch": "B",
        "branch_cost_level": "high",
        "parent_candidate_hash": "two",
        "applied_branches": ["A", "B"],
    }
    assert candidate_config_hash(left) == candidate_config_hash(right)


def test_deep_memory_write_generates_current_dataset_artifacts_without_frozen_rankings(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import tuner.generated_source_artifacts as generated_sources

    dataset = make_dataset(tmp_path, 2)
    baseline, _ = _production_branch_inputs(tmp_path, dataset)
    generated_session_sets: list[tuple[str, ...]] = []

    def fake_generate(**kwargs: Any) -> list[Path]:
        generated_session_sets.append(tuple(kwargs["session_ids"]))
        root = Path(kwargs["source_root"])
        paths = []
        for session_id in kwargs["session_ids"]:
            path = root / session_id / "production_midterm_manifest.json"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(
                json.dumps(
                    {
                        "session_id": session_id,
                        "production_config": {"top_k_sessions": 5, "top_k_pages": 5, "max_total_pages": 5},
                        "effective_memory_config": {"vector_store": {"config": {"bm25_language": "en"}}},
                        "failed_turns": 0,
                        "llm_calls": 2,
                        "embedding_calls": 5,
                    }
                ),
                encoding="utf-8",
            )
            paths.append(path)
        kwargs["generation_stats"].update(
            {"generated_sessions": len(paths), "reused_sessions": 0, "llm_calls": 2, "embedding_calls": 5}
        )
        return paths

    monkeypatch.setattr(generated_sources, "generate_production_sources", fake_generate)
    discovery = SimpleNamespace(resources=SimpleNamespace(gpu_count=0))
    context = _branch_context(tmp_path, dataset, baseline, budget="deep", model_discovery=discovery)
    branch = SourcePromptBranch()
    outcome = BranchRegistry([branch]).generate(branch, context)
    assert outcome.status == "READY"
    assert len(outcome.candidates) == 4
    assert {candidate.config["source_variant"] for candidate in outcome.candidates} == {
        "conservative_add",
        "context_aware_add",
        "evidence_focused_summary",
        "diagnostic_controlled_summary",
    }
    assert outcome.llm_calls == 0
    assert outcome.embedding_calls == 0
    assert generated_session_sets == []

    screening_session = next(iter(dataset.sessions))
    prepared = prepare_generated_source_candidate(
        outcome.candidates[0],
        sessions=[screening_session],
        registry=context.registry,
        run_dir=context.run_dir,
        ranking_depth=context.ranking_depth,
        max_parallel_sessions=2,
        max_parallel_llm_calls=2,
        gpu_count=0,
    )
    assert generated_session_sets == [(screening_session,)]
    assert prepared.config["source_variant"] == "conservative_add"
    assert prepared.provenance["deferred_generation_stats"]["llm_calls"] == 2
    assert prepared.provenance["deferred_generation_stats"]["embedding_calls"] == 5
    selected_manifest = next(
        Path(path)
        for path in prepared.config["manifest_paths"]
        if json.loads(Path(path).read_text(encoding="utf-8"))["session_id"] == screening_session
    )
    assert "production_variants" in str(selected_manifest)


def test_later_branch_combines_with_frontier_anchor(tmp_path: Path) -> None:
    class Retrieval(_StaticBranch):
        def generate(self, context: Any) -> BranchOutcome:
            return BranchOutcome(
                self.spec.name,
                "READY",
                [
                    Candidate(
                        name="retrieval-winner",
                        stage="stage-1",
                        config={**context.anchor.config, "top_k_sessions": 6, "gain": 0.1},
                    )
                ],
            )

    class Embedding(_StaticBranch):
        def generate(self, context: Any) -> BranchOutcome:
            return BranchOutcome(
                self.spec.name,
                "READY",
                [
                    Candidate(
                        name="combined-winner",
                        stage="stage-2",
                        config={**context.anchor.config, "embedding_model_id": "local/model", "gain": 0.15},
                    )
                ],
            )

    dataset = make_dataset(tmp_path, 3)
    baseline = Candidate(name="baseline", stage="baseline", config={"top_k_sessions": 5, "gain": 0.0})
    registry = BranchRegistry(
        [
            Retrieval("Retrieval", initial=True, priority=1, gain=0.1),
            Embedding("Embedding", initial=False, priority=2, gain=0.15),
        ]
    )

    def evaluate(candidates: Any, sessions: Any, scope: str) -> list[CandidateResult]:
        del sessions, scope
        return [result(candidate.name, 0.4 + candidate.config["gain"]) for candidate in candidates]

    search = run_staged_search(
        dataset=dataset,
        baseline=baseline,
        baseline_result=result("baseline", 0.4),
        tune_sessions=tuple(dataset.sessions),
        registry=registry,
        artifact_registry=None,
        model_discovery=None,
        run_dir=tmp_path,
        search_space={"selection": {"min_improvement_pp": 0.25, "patience_stages": 2}},
        budget="standard",
        profile={"max_stages": 3, "max_cost_level": "medium", "max_candidates_per_stage": 3},
        k=5,
        ranking_depth=20,
        evaluate=evaluate,
        diagnose=lambda _: {"regime": "ranking_bottleneck"},
    )
    assert search.candidates["combined-winner"].config["top_k_sessions"] == 6


def test_derived_artifact_rebuild_preserves_frontier_dimensions(tmp_path: Path) -> None:
    query_artifact = tmp_path / "queries.json"
    query_artifact.write_text("{}", encoding="utf-8")
    anchor = Candidate(
        name="combined",
        stage="tune",
        config={
            "query_representation": "bounded_reference_resolution",
            "query_artifact_path": str(query_artifact),
            "query_artifact_variant": "generated:v1",
            "page_representation": "summary_keywords",
            "embedding_model_id": "local/embedding",
            "embedding_model_revision": "immutable-revision",
            "embedding_model_path": "/models/embedding",
            "encoding_contract": {"query_prefix": "query: ", "document_prefix": "passage: "},
            "reranker_method": "multi_vector_maxsim",
        },
    )
    with pytest.raises(ValueError, match="regenerated by Production"):
        _derived_dimensions(anchor, page_representation_name="user_summary")
    dimensions = _derived_dimensions(anchor)
    assert dimensions == {
        "query_representation": "bounded_reference_resolution",
        "query_artifact_path": query_artifact,
        "query_artifact_variant": "generated:v1",
        "embedding_model_id": "local/embedding",
        "embedding_revision": "immutable-revision",
        "embedding_local_path": "/models/embedding",
        "encoding_contract": {"query_prefix": "query: ", "document_prefix": "passage: "},
    }


@pytest.mark.parametrize(("budget", "expected_network"), [("standard", False), ("deep", True)])
def test_reranker_budget_uses_local_first_and_network_only_for_deep(
    tmp_path: Path,
    budget: str,
    expected_network: bool,
) -> None:
    dataset = make_dataset(tmp_path, 1)
    baseline, _ = _production_branch_inputs(tmp_path, dataset)

    class FakeDiscovery:
        resources = SimpleNamespace(gpu_count=0)

        def __init__(self) -> None:
            self.allow_network: list[bool] = []

        def discover(self, *, allow_network: bool, **_: Any) -> list[ModelCandidate]:
            self.allow_network.append(allow_network)
            return [ModelCandidate("local/zh-reranker", "reranker", "local_huggingface_cache")]

        def ensure_available(self, candidate: ModelCandidate, *, allow_download: bool) -> ModelCandidate:
            assert allow_download is expected_network
            candidate.status = "AVAILABLE"
            candidate.local_path = "/tmp/local-reranker"
            return candidate

        def smoke_test(self, candidate: ModelCandidate, *, device: str) -> ModelCandidate:
            del device
            candidate.status = "SMOKE_PASSED"
            return candidate

    discovery = FakeDiscovery()
    context = _branch_context(tmp_path, dataset, baseline, budget=budget, model_discovery=discovery)
    outcome = RerankingBranch().generate(context)
    assert outcome.status == "READY"
    assert discovery.allow_network == [expected_network]
    assert any(candidate.config.get("reranker_method") == "cross_encoder" for candidate in outcome.candidates)


def test_model_specific_encoding_contract_is_applied() -> None:
    class FakeModel:
        def __init__(self) -> None:
            self.calls: list[tuple[list[str], dict[str, Any]]] = []

        def encode(self, texts: list[str], **kwargs: Any) -> list[list[float]]:
            self.calls.append((texts, kwargs))
            return [[1.0, 0.0] for _ in texts]

    model = FakeModel()

    class Delegate:
        def __init__(self, model: FakeModel) -> None:
            self.model = model

        def embed(self, text: str, action: str) -> list[float]:
            return self.model.encode([text], convert_to_numpy=True)[0]

    contract = EncodingContract(
        query_prefix="query: ",
        document_prefix="passage: ",
        query_prompt_name="query",
        normalize_embeddings=True,
        pooling="mean",
        source="test-model-card",
    )
    adapter = EncodingContractEmbedding(Delegate(model), contract.serializable())
    adapter.embed_batch(["现金流"], "search")
    adapter.embed_batch(["经营现金流改善"], "add")
    assert model.calls[0][0] == ["query: 现金流"]
    assert model.calls[0][1]["prompt_name"] == "query"
    assert model.calls[0][1]["normalize_embeddings"] is True
    assert model.calls[1][0] == ["passage: 经营现金流改善"]
    assert "prompt_name" not in model.calls[1][1]


def test_embedding_branch_replays_real_production_sources_before_evaluation(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import tuner.generated_source_artifacts as generated_sources

    dataset = make_dataset(tmp_path, 1)
    baseline, _ = _production_branch_inputs(tmp_path, dataset)
    model_path = tmp_path / "models" / "embedding" / "snapshots" / "immutable-revision"
    model_path.mkdir(parents=True)

    class FakeDiscovery:
        resources = SimpleNamespace(gpu_count=0)

        def discover(self, **_: Any) -> list[ModelCandidate]:
            return [
                ModelCandidate(
                    "local/source-changing-embedding",
                    "embedding",
                    "local_huggingface_cache",
                    revision="immutable-revision",
                    local_path=str(model_path),
                    cache_status="CACHED",
                )
            ]

        def ensure_available(self, candidate: ModelCandidate, **_: Any) -> ModelCandidate:
            candidate.status = "AVAILABLE"
            return candidate

        def smoke_test(self, candidate: ModelCandidate, **_: Any) -> ModelCandidate:
            candidate.status = "SMOKE_PASSED"
            candidate.encoding_contract = {
                "query_prefix": "query: ",
                "document_prefix": "passage: ",
                "normalize_embeddings": True,
                "source": "test-model-contract",
            }
            candidate.resource_usage["embedding_dimension"] = 384
            return candidate

    context = _branch_context(
        tmp_path,
        dataset,
        baseline,
        budget="standard",
        model_discovery=FakeDiscovery(),
    )

    def unexpected_derived_build(*args: Any, **kwargs: Any) -> None:
        del args, kwargs
        raise AssertionError("Embedding source candidates must not rebuild the old Session graph")

    monkeypatch.setattr(DerivedArtifactBuilder, "build", unexpected_derived_build)
    branch = EmbeddingBranch()
    outcome = BranchRegistry([branch]).generate(branch, context)
    assert outcome.status == "READY"
    assert len(outcome.candidates) == 1
    candidate = outcome.candidates[0]
    spec = candidate.config["source_generation_spec"]
    assert parameter_class("embedding_model_id") == "source-changing"
    assert candidate.provenance["requires_source_regeneration"] is True
    assert spec["source_identity"]["real_add_replay"] is True
    assert spec["embedding_model_revision"] == "immutable-revision"
    assert spec["embedding_encoding_contract"] == candidate.config["encoding_contract"]
    assert spec["config_overrides"]["embedder"]["config"]["model"] == str(model_path.resolve())
    assert spec["config_overrides"]["vector_store"]["config"]["embedding_model_dims"] == 384
    assert candidate.config["production_overrides"]["embedder"]["config"]["model"] == str(model_path.resolve())
    assert candidate.config["production_overrides"]["embedder"]["config"]["revision"] == "immutable-revision"
    assert (
        candidate.config["production_overrides"]["embedder"]["config"]["encoding_contract"]
        == candidate.config["encoding_contract"]
    )
    assert candidate.config["production_overrides"]["vector_store"]["config"]["embedding_model_dims"] == 384
    assert candidate.config.get("derived_artifact_path") is None

    generation_calls: list[dict[str, Any]] = []

    def fake_generate(**kwargs: Any) -> list[Path]:
        generation_calls.append(kwargs)
        paths = []
        for session_id in kwargs["session_ids"]:
            path = Path(kwargs["source_root"]) / session_id / "production_midterm_manifest.json"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(
                json.dumps(
                    {
                        "session_id": session_id,
                        "effective_memory_config": {
                            "embedder": {
                                "provider": "huggingface",
                                "config": {"model": str(model_path.resolve()), "embedding_dims": 384},
                            }
                        },
                        "embedding_model_revision": "immutable-revision",
                        "embedding_encoding_contract": kwargs["embedding_encoding_contract"],
                        "failed_turns": 0,
                    }
                ),
                encoding="utf-8",
            )
            paths.append(path)
        kwargs["generation_stats"].update(
            {"generated_sessions": len(paths), "reused_sessions": 0, "llm_calls": 4, "embedding_calls": 12}
        )
        return paths

    monkeypatch.setattr(generated_sources, "generate_production_sources", fake_generate)

    def fake_evaluate_candidate(
        dataset: Dataset,
        candidate: Candidate,
        sessions: Any,
        **_: Any,
    ) -> CandidateResult:
        del dataset, sessions
        assert generation_calls, "source generation must complete before retrieval evaluation"
        generated_manifest = json.loads(Path(candidate.config["manifest_paths"][0]).read_text(encoding="utf-8"))
        assert generated_manifest["embedding_model_revision"] == "immutable-revision"
        measured = result(candidate.name, 0.6)
        measured.config = candidate.config
        measured.stage = candidate.stage
        return measured

    monkeypatch.setattr(orchestrator, "evaluate_candidate", fake_evaluate_candidate)
    evaluated = orchestrator._evaluate_many(
        dataset,
        [candidate],
        list(dataset.sessions),
        scope="embedding_tune",
        config=TunerConfig(dataset=Path(dataset.path)),
        shortterm_window=3,
        ranking_depth=20,
        registry=context.registry,
        run_dir=context.run_dir,
        execution={
            "max_parallel_sessions": 1,
            "max_parallel_candidates": 1,
            "max_parallel_llm_calls": 1,
            "gpu_count": 0,
            "adaptive_reductions": [],
        },
        trace_path=tmp_path / "search_trace.jsonl",
    )

    assert evaluated[0].status == "VALID"
    assert len(generation_calls) == 1
    generated_call = generation_calls[0]
    assert generated_call["config_overrides"]["embedder"]["config"]["model"] == str(model_path.resolve())
    assert generated_call["embedding_model_id"] == "local/source-changing-embedding"
    assert generated_call["embedding_model_revision"] == "immutable-revision"
    assert generated_call["embedding_encoding_contract"] == candidate.config["encoding_contract"]
    assert "derived_artifact_path" not in candidate.config


def test_production_source_runtime_receives_candidate_embedding_and_keeps_isolation(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import asyncio

    dataset_path = tmp_path / "dataset.xlsx"
    dataset_path.write_bytes(b"dataset")
    memory_config_path = tmp_path / "memory_config.json"
    memory_config_path.write_text(
        json.dumps(
            {
                "embedder": {
                    "provider": "huggingface",
                    "config": {"model": "baseline/model", "embedding_dims": 512},
                },
                "vector_store": {
                    "provider": "qdrant",
                    "config": {
                        "path": "/baseline/qdrant",
                        "collection_name": "baseline",
                        "embedding_model_dims": 512,
                    },
                },
                "midterm": {"enabled": True, "short_term_capacity": 6},
            }
        ),
        encoding="utf-8",
    )
    model_path = tmp_path / "models" / "snapshots" / "revision-1"
    model_path.mkdir(parents=True)
    captured_config: dict[str, Any] = {}

    class FakeSentenceTransformer:
        def encode(self, texts: list[str], **_: Any) -> list[list[float]]:
            return [[1.0, 0.0] for _ in texts]

    class FakeEmbedding:
        model = FakeSentenceTransformer()

    class FakeMemory:
        def __init__(self) -> None:
            self.llm = SimpleNamespace()
            self.embedding_model = FakeEmbedding()
            self._midterm_updater = None
            self._midterm_memory = None

        @staticmethod
        def _short_term_capacity() -> int:
            return 6

        async def flush_background_tasks(self, **_: Any) -> None:
            return None

        def close(self) -> None:
            return None

    def fake_create_memory(config: dict[str, Any], **_: Any) -> FakeMemory:
        captured_config.update(config)
        return FakeMemory()

    monkeypatch.setattr(
        production_adapter,
        "load_dataset",
        lambda *args, **kwargs: [SimpleNamespace(session_id="S001_test", turns=[])],
    )
    monkeypatch.setattr(production_adapter, "create_production_memory", fake_create_memory)
    runtime_dir = tmp_path / "isolated-runtime"
    output_dir = tmp_path / "source-output"
    contract = {
        "query_prefix": "query: ",
        "document_prefix": "passage: ",
        "normalize_embeddings": True,
        "source": "test-model-contract",
    }
    manifest = asyncio.run(
        production_adapter.build_production_source(
            {
                "dataset_path": str(dataset_path),
                "session_id": "S001_test",
                "output_dir": str(output_dir),
                "runtime_dir": str(runtime_dir),
                "collection_name": "isolated-candidate",
                "memory_config_path": str(memory_config_path),
                "ranking_depth": 20,
                "llm_mode": "mock",
                "source_variant": "embedding:model@revision-1",
                "source_identity": {"model_id": "model", "model_revision": "revision-1"},
                "embedding_model_id": "model",
                "embedding_model_revision": "revision-1",
                "embedding_encoding_contract": contract,
                "config_overrides": {
                    "embedder": {
                        "provider": "huggingface",
                        "config": {
                            "model": str(model_path.resolve()),
                            "embedding_dims": 2,
                            "model_kwargs": {"local_files_only": True},
                        },
                    },
                    "vector_store": {
                        "config": {
                            "path": "/must-not-escape-isolation",
                            "collection_name": "must-not-escape-isolation",
                            "embedding_model_dims": 2,
                        }
                    },
                },
            }
        )
    )

    assert captured_config["embedder"]["config"]["model"] == str(model_path.resolve())
    assert captured_config["embedder"]["config"]["model_kwargs"]["local_files_only"] is True
    assert captured_config["vector_store"]["config"]["embedding_model_dims"] == 2
    assert captured_config["vector_store"]["config"]["path"] == str((runtime_dir / "qdrant").resolve())
    assert captured_config["vector_store"]["config"]["collection_name"] == "isolated-candidate"
    assert manifest["embedding_model_revision"] == "revision-1"
    assert manifest["embedding_encoding_contract"] == contract
