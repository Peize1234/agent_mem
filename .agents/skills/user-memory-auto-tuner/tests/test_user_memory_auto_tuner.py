from __future__ import annotations

import csv
import json
import sqlite3
import sys
from collections import deque
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_ROOT = SKILL_ROOT / "scripts"
REPO_ROOT = SKILL_ROOT.parents[2]
for path in (REPO_ROOT, SCRIPTS_ROOT):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from auto_tuner_models import (  # noqa: E402
    ConversationTurn,
    ExtractedHistory,
    ExtractedSession,
    RawMessage,
)
from build_benchmark import build_benchmark  # noqa: E402
from config_store import UserTuningConfigStore, production_override_delta  # noqa: E402
from extract_user_history import extract_user_history, parse_session_scope  # noqa: E402
from label_dependencies import TwoStageDependencyAnnotator, annotate_history  # noqa: E402
from run_user_tuning import UserTuningRequest, run_user_tuning  # noqa: E402
from tuner.benchmark_support import load_dataset  # noqa: E402
from tuner_bridge import TUNER_SKILL_ROOT, audit_generated_benchmark, run_tuning  # noqa: E402


class ScriptedLLM:
    def __init__(self, responses: list[dict[str, Any]]):
        self.responses = deque(responses)
        self.calls: list[list[dict[str, str]]] = []

    def generate_response(
        self,
        messages: list[dict[str, str]],
        *,
        response_format: dict[str, Any],
    ) -> str:
        assert response_format == {"type": "json_object"}
        self.calls.append(messages)
        return json.dumps(self.responses.popleft(), ensure_ascii=False)


def _create_history_db(path: Path) -> None:
    connection = sqlite3.connect(path)
    connection.execute(
        """
        CREATE TABLE messages (
            id TEXT PRIMARY KEY,
            session_scope TEXT,
            role TEXT,
            content TEXT,
            name TEXT,
            created_at DATETIME,
            turn_index INTEGER NOT NULL,
            status TEXT NOT NULL DEFAULT 'active',
            migration_job_id TEXT,
            source_operation_key TEXT,
            source_message_index INTEGER
        )
        """
    )
    connection.commit()
    connection.close()


def _insert_message(
    path: Path,
    *,
    message_id: str,
    scope: str,
    role: str,
    content: str,
    turn_index: int,
    created_at: str,
) -> None:
    connection = sqlite3.connect(path)
    connection.execute(
        """
        INSERT INTO messages (
            id, session_scope, role, content, name, created_at, turn_index, status
        ) VALUES (?, ?, ?, ?, NULL, ?, ?, 'active')
        """,
        (message_id, scope, role, content, created_at, turn_index),
    )
    connection.commit()
    connection.close()


def _add_qa(
    path: Path,
    *,
    scope: str,
    prefix: str,
    turn_index: int,
    question: str,
    answer: str,
    day: int,
) -> None:
    _insert_message(
        path,
        message_id=f"{prefix}-u-{turn_index}",
        scope=scope,
        role="user",
        content=question,
        turn_index=turn_index,
        created_at=f"2026-01-{day:02d}T00:00:{turn_index * 2:02d}+08:00",
    )
    _insert_message(
        path,
        message_id=f"{prefix}-a-{turn_index}",
        scope=scope,
        role="assistant",
        content=answer,
        turn_index=turn_index,
        created_at=f"2026-01-{day:02d}T00:00:{turn_index * 2 + 1:02d}+08:00",
    )


def _independent_label() -> dict[str, Any]:
    return {"needs_history": False, "requirements": [], "dependency_type": "", "confidence": 0.95}


def _valid_verifier() -> dict[str, Any]:
    return {
        "valid": True,
        "dependencies_necessary": True,
        "context_supported": True,
        "logic_valid": True,
        "confidence": 0.95,
        "issues": [],
    }


def _dependent_label(dependency_id: str, context: str = "额度100万元") -> dict[str, Any]:
    return {
        "needs_history": True,
        "requirements": [{"dependency_ids": [dependency_id], "required_contexts": [context]}],
        "dependency_type": "事实依赖",
        "confidence": 0.95,
    }


def _history_with_turns(turns: list[tuple[str, str]]) -> ExtractedHistory:
    conversation_turns = []
    for index, (question, answer) in enumerate(turns, start=1):
        raw = (
            RawMessage(
                rowid=index * 2 - 1,
                message_id=f"u-{index}",
                role="user",
                content=question,
                turn_index=index,
                status="active",
            ),
            RawMessage(
                rowid=index * 2,
                message_id=f"a-{index}",
                role="assistant",
                content=answer,
                turn_index=index,
                status="active",
            ),
        )
        conversation_turns.append(
            ConversationTurn(
                source_turn_index=index,
                messages=raw,
                question=question,
                answer=answer,
            )
        )
    return ExtractedHistory(
        user_id="u1",
        history_db_path="/tmp/history.db",
        sessions=(
            ExtractedSession(
                session_scope="run_id=r1&user_id=u1",
                user_id="u1",
                run_id="r1",
                turns=tuple(conversation_turns),
            ),
        ),
    )


def _three_session_db(path: Path, user_id: str = "u1") -> None:
    _create_history_db(path)
    for session_number in range(1, 4):
        scope = f"agent_id=a1&run_id=r{session_number}&user_id={user_id}"
        _add_qa(
            path,
            scope=scope,
            prefix=f"s{session_number}",
            turn_index=1,
            question="本次额度是多少？",
            answer="额度100万元",
            day=session_number,
        )
        _add_qa(
            path,
            scope=scope,
            prefix=f"s{session_number}",
            turn_index=2,
            question="刚才的额度适合吗？",
            answer="适合，因为额度是100万元。",
            day=session_number,
        )


def _three_session_annotator() -> TwoStageDependencyAnnotator:
    responses = []
    for session_number in range(1, 4):
        responses.extend([_dependent_label(f"S{session_number:03d}-Q001"), _valid_verifier()])
    return TwoStageDependencyAnnotator(ScriptedLLM(responses))


def _write_fake_tuner_run(
    config: Any,
    *,
    candidate_config: dict[str, Any],
    candidate_status: str = "VALID",
) -> Path:
    run_dir = Path(config.output_root) / "fake-run"
    run_dir.mkdir(parents=True)
    (run_dir / "dataset_audit.json").write_text(
        json.dumps({"status": "PASS", "hard_errors": []}),
        encoding="utf-8",
    )
    (run_dir / "split_manifest.json").write_text(
        json.dumps(
            {
                "method": "leave_one_session_out",
                "tune_sessions": [],
                "validation_sessions": [],
                "folds": [{"validation_sessions": ["S001"]}],
            }
        ),
        encoding="utf-8",
    )
    metrics = {
        "recall_at_k": 0.8,
        "macro_session_recall_at_k": 0.75,
        "mrr": 0.7,
    }
    (run_dir / "best_config.json").write_text(
        json.dumps(
            {
                "candidate": "candidate-a",
                "config": candidate_config,
                "validation_metrics": metrics,
                "selection": {"primary": "validation_requirement_recall_at_k"},
            }
        ),
        encoding="utf-8",
    )
    (run_dir / "run_metadata.json").write_text(
        json.dumps({"status": "COMPLETE", "failed_turns": 0}),
        encoding="utf-8",
    )
    with (run_dir / "leaderboard.csv").open("w", encoding="utf-8", newline="") as destination:
        writer = csv.DictWriter(
            destination,
            fieldnames=["candidate", "validation_R@5", "validation_macro_R@5", "status"],
        )
        writer.writeheader()
        writer.writerow(
            {
                "candidate": "baseline",
                "validation_R@5": "0.5",
                "validation_macro_R@5": "0.5",
                "status": "VALID",
            }
        )
        writer.writerow(
            {
                "candidate": "candidate-a",
                "validation_R@5": "0.8",
                "validation_macro_R@5": "0.75",
                "status": candidate_status,
            }
        )
    return run_dir


def test_extracts_only_requested_user_groups_sessions_and_orders_turns(tmp_path: Path) -> None:
    database = tmp_path / "history.db"
    _create_history_db(database)
    scope_two = "agent_id=a1&run_id=r2&user_id=u1"
    scope_one = "user_id=u1&run_id=r1"
    _add_qa(
        database,
        scope=scope_two,
        prefix="late",
        turn_index=2,
        question="第二问",
        answer="第二答",
        day=2,
    )
    _add_qa(
        database,
        scope=scope_two,
        prefix="early",
        turn_index=1,
        question="第一问",
        answer="第一答",
        day=2,
    )
    _add_qa(
        database,
        scope=scope_one,
        prefix="one",
        turn_index=1,
        question="另一个 Session",
        answer="回答",
        day=1,
    )
    _add_qa(
        database,
        scope="run_id=other&user_id=u2",
        prefix="other",
        turn_index=1,
        question="不属于 u1",
        answer="不能提取",
        day=1,
    )

    history = extract_user_history(user_id="u1", history_db_path=database, output_dir=tmp_path / "out")

    assert [session.session_scope for session in history.sessions] == [scope_one, scope_two]
    assert [turn.source_turn_index for turn in history.sessions[1].turns] == [1, 2]
    assert [turn.question for turn in history.sessions[1].turns] == ["第一问", "第二问"]
    assert history.qa_turn_count == 3
    serialized = json.loads((tmp_path / "out/extracted_history.json").read_text(encoding="utf-8"))
    assert serialized["sessions"][1]["turns"][0]["messages"][0]["content"] == "第一问"


def test_session_scope_parser_and_sql_filter_use_exact_user_id(tmp_path: Path) -> None:
    database = tmp_path / "history.db"
    _create_history_db(database)
    _add_qa(
        database,
        scope="run_id=r1&user_id=u%",
        prefix="percent",
        turn_index=1,
        question="精确用户",
        answer="是",
        day=1,
    )
    _add_qa(
        database,
        scope="run_id=r2&user_id=ux",
        prefix="wildcard",
        turn_index=1,
        question="不应匹配",
        answer="否",
        day=2,
    )

    history = extract_user_history(user_id="u%", history_db_path=database)

    assert len(history.sessions) == 1
    assert history.sessions[0].turns[0].question == "精确用户"
    assert parse_session_scope("agent_id=a&run_id=r&user_id=u") == {
        "agent_id": "a",
        "run_id": "r",
        "user_id": "u",
    }


def test_future_and_invalid_dependencies_are_filtered_by_deterministic_validation() -> None:
    history = _history_with_turns([("Q1", "额度100万元"), ("Q2", "A2"), ("Q3", "A3")])
    llm = ScriptedLLM(
        [
            _dependent_label("S001-Q003"),
            _valid_verifier(),
            _dependent_label("S999-Q001"),
            _valid_verifier(),
        ]
    )

    annotations = annotate_history(history, TwoStageDependencyAnnotator(llm))

    assert annotations.filtered_samples == 2
    assert annotations.sessions[0].turns[1].status == "INVALID"
    assert "MISSING_OR_FUTURE_DEPENDENCY:S001-Q003" in annotations.sessions[0].turns[1].validation_issues
    assert "MISSING_OR_FUTURE_DEPENDENCY:S999-Q001" in annotations.sessions[0].turns[2].validation_issues
    label_prompt = llm.calls[0][1]["content"]
    assert "S001-Q003" not in label_prompt


def test_required_context_not_present_in_history_is_filtered() -> None:
    history = _history_with_turns([("Q1", "额度100万元"), ("Q2", "回答新增了期限两年")])
    llm = ScriptedLLM([_dependent_label("S001-Q001", "期限两年"), _valid_verifier()])

    annotations = annotate_history(history, TwoStageDependencyAnnotator(llm))

    turn = annotations.sessions[0].turns[1]
    assert turn.status == "INVALID"
    assert "REQUIRED_CONTEXT_NOT_VERBATIM:S001-Q001" in turn.validation_issues


def test_and_or_gold_generation_passes_existing_dataset_audit(tmp_path: Path) -> None:
    history = _history_with_turns(
        [
            ("初始额度？", "额度100万元"),
            ("换种说法？", "金额一百万元"),
            ("期限？", "期限一年"),
            ("综合判断？", "额度和期限均合适"),
        ]
    )
    label = {
        "needs_history": True,
        "requirements": [
            {
                "dependency_ids": ["S001-Q001", "S001-Q002"],
                "required_contexts": ["额度100万元", "金额一百万元"],
            },
            {"dependency_ids": ["S001-Q003"], "required_contexts": ["期限一年"]},
        ],
        "dependency_type": "多事实 AND + 等价来源 OR",
        "confidence": 0.98,
    }
    llm = ScriptedLLM(
        [
            _independent_label(),
            _valid_verifier(),
            _independent_label(),
            _valid_verifier(),
            label,
            _valid_verifier(),
        ]
    )
    annotations = annotate_history(history, TwoStageDependencyAnnotator(llm))
    artifacts = build_benchmark(annotations, output_dir=tmp_path, history_db_path=tmp_path / "history.db")

    workbook = load_workbook(artifacts.benchmark_path, read_only=True, data_only=True)
    try:
        rows = list(workbook["S001"].iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = {name: index for index, name in enumerate(rows[0])}
    last = rows[-1]
    assert last[headers["关联前序对话"]] == "（S001-Q001；S001-Q002）；S001-Q003"
    assert last[headers["所需前文信息"]] == "（额度100万元 OR 金额一百万元）；期限一年"
    loaded = load_dataset(artifacts.benchmark_path)
    assert loaded[0].turns[-1].dependency_turn_ids == ("S001-Q001", "S001-Q002", "S001-Q003")
    audit = audit_generated_benchmark(
        artifacts.benchmark_path,
        output_dir=tmp_path / "audit",
        shortterm_qa_turns=3,
    )
    assert audit["status"] in {"PASS", "DATASET_QUALITY_WARNING"}
    assert audit["future_dependency_count"] == 0
    assert audit["or_requirement_count"] == 1
    assert audit["and_requirement_count"] == 1


def test_not_enough_data_skips_tuner(tmp_path: Path) -> None:
    database = tmp_path / "history.db"
    _create_history_db(database)
    scope = "run_id=r1&user_id=u1"
    _add_qa(
        database,
        scope=scope,
        prefix="s1",
        turn_index=1,
        question="额度？",
        answer="额度100万元",
        day=1,
    )
    _add_qa(
        database,
        scope=scope,
        prefix="s1",
        turn_index=2,
        question="合适吗？",
        answer="合适",
        day=1,
    )
    called = False

    def tuner_runner(*_: Any, **__: Any) -> Path:
        nonlocal called
        called = True
        raise AssertionError("tuner must not run")

    request = UserTuningRequest(user_id="u1", history_db_path=database, output_dir=tmp_path / "run")
    result = run_user_tuning(
        request,
        annotator=TwoStageDependencyAnnotator(ScriptedLLM([_dependent_label("S001-Q001"), _valid_verifier()])),
        tuner_runner=tuner_runner,
    )

    assert result["status"] == "skipped"
    assert result["reason"]["code"] == "NOT_ENOUGH_DATA"
    assert result["valid_dependency_samples"] == 1
    assert not called
    assert Path(result["generated_benchmark_path"]).exists()


def test_pipeline_reuses_existing_tuner_and_saves_only_safe_production_overrides(tmp_path: Path) -> None:
    database = tmp_path / "history.db"
    _three_session_db(database)
    captured: dict[str, Any] = {}

    def tuner_runner(config: Any, *, skill_root: Path) -> Path:
        captured["config"] = config
        captured["skill_root"] = skill_root
        return _write_fake_tuner_run(
            config,
            candidate_config={
                "production_overrides": {
                    "query_rewrite_prompt": "safe query-time prompt",
                    "midterm": {
                        "top_k_pages": 7,
                        "page_representation": "summary",
                        "page_summary_prompt": "source-changing prompt",
                    },
                    "promoted_longterm": {"top_k": 9},
                },
                "experiment_metadata": {"must_not_be_saved": True},
            },
        )

    request = UserTuningRequest(user_id="u1", history_db_path=database, output_dir=tmp_path / "run")
    result = run_user_tuning(
        request,
        annotator=_three_session_annotator(),
        tuner_runner=tuner_runner,
    )

    assert run_tuning.__module__ == "tuner.orchestrator"
    assert captured["config"].__class__.__module__ == "tuner.orchestrator"
    assert captured["skill_root"] == TUNER_SKILL_ROOT
    assert result["status"] == "deployed"
    assert result["production_overrides"] == {
        "query_rewrite_prompt": "safe query-time prompt",
        "midterm": {"top_k_pages": 7},
    }
    assert result["rebuild_required"] == {
        "midterm": {
            "page_representation": "summary",
            "page_summary_prompt": "source-changing prompt",
        }
    }
    assert result["unsupported_overrides"] == {"promoted_longterm": {"top_k": 9}}
    store = UserTuningConfigStore(database)
    active = store.get_active("u1")
    assert active is not None
    assert active.config_overrides == result["production_overrides"]
    assert "candidate" not in active.config_overrides


def test_candidate_delta_does_not_mark_unchanged_source_fields_for_rebuild() -> None:
    delta = production_override_delta(
        {"midterm": {"top_k_pages": 7, "page_representation": "raw"}},
        {"midterm": {"top_k_pages": 5, "page_representation": "raw"}},
    )

    assert delta == {"midterm": {"top_k_pages": 7}}


def test_invalid_memory_config_does_not_overwrite_old_active_config(tmp_path: Path) -> None:
    database = tmp_path / "history.db"
    _three_session_db(database)
    store = UserTuningConfigStore(database)
    original = store.save_active(
        user_id="u1",
        config_overrides={"midterm": {"top_k_pages": 5}},
        source_run_dir=tmp_path / "old-run",
        dataset_hash="old-dataset",
        validation_metrics={"recall_at_k": 0.5},
    )

    def tuner_runner(config: Any, *, skill_root: Path) -> Path:
        assert skill_root == TUNER_SKILL_ROOT
        assert json.loads(Path(config.memory_config).read_text(encoding="utf-8")) == original.config_overrides
        return _write_fake_tuner_run(
            config,
            candidate_config={"production_overrides": {"midterm": {"top_k_pages": 0}}},
        )

    result = run_user_tuning(
        UserTuningRequest(user_id="u1", history_db_path=database, output_dir=tmp_path / "run"),
        annotator=_three_session_annotator(),
        tuner_runner=tuner_runner,
    )

    assert result["status"] == "skipped"
    assert result["reason"]["code"] == "INVALID_PRODUCTION_OVERRIDES"
    unchanged = store.get_active("u1")
    assert unchanged is not None
    assert unchanged.config_version == 1
    assert unchanged.config_overrides == {"midterm": {"top_k_pages": 5}}


def test_overfit_candidate_does_not_overwrite_old_active_config(tmp_path: Path) -> None:
    database = tmp_path / "history.db"
    _three_session_db(database)
    store = UserTuningConfigStore(database)
    store.save_active(
        user_id="u1",
        config_overrides={"midterm": {"top_k_pages": 5}},
        source_run_dir=tmp_path / "old-run",
        dataset_hash="old-dataset",
        validation_metrics={"recall_at_k": 0.5},
    )

    def tuner_runner(config: Any, *, skill_root: Path) -> Path:
        assert skill_root == TUNER_SKILL_ROOT
        return _write_fake_tuner_run(
            config,
            candidate_config={"production_overrides": {"midterm": {"top_k_pages": 7}}},
            candidate_status="OVERFIT",
        )

    result = run_user_tuning(
        UserTuningRequest(user_id="u1", history_db_path=database, output_dir=tmp_path / "run"),
        annotator=_three_session_annotator(),
        tuner_runner=tuner_runner,
    )

    assert result["status"] == "skipped"
    assert result["reason"]["code"] == "BEST_CANDIDATE_OVERFIT"
    unchanged = store.get_active("u1")
    assert unchanged is not None
    assert unchanged.config_version == 1
    assert unchanged.config_overrides == {"midterm": {"top_k_pages": 5}}


def test_config_store_archives_previous_version(tmp_path: Path) -> None:
    database = tmp_path / "history.db"
    _create_history_db(database)
    store = UserTuningConfigStore(database)
    first = store.save_active(
        user_id="u1",
        config_overrides={"midterm": {"top_k_pages": 5}},
        source_run_dir=tmp_path / "run-1",
        dataset_hash="dataset-1",
        validation_metrics={"recall_at_k": 0.5},
    )
    second = store.save_active(
        user_id="u1",
        config_overrides={"midterm": {"top_k_pages": 7}},
        source_run_dir=tmp_path / "run-2",
        dataset_hash="dataset-2",
        validation_metrics={"recall_at_k": 0.8},
    )

    assert (first.config_version, second.config_version) == (1, 2)
    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    try:
        history = connection.execute(
            "SELECT * FROM user_memory_tuning_config_history WHERE user_id = ?",
            ("u1",),
        ).fetchall()
    finally:
        connection.close()
    assert len(history) == 1
    assert history[0]["config_version"] == 1
    assert history[0]["replaced_by_version"] == 2
    assert json.loads(history[0]["config_overrides_json"]) == {"midterm": {"top_k_pages": 5}}
