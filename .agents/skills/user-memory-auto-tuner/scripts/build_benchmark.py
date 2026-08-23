from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook

from auto_tuner_models import AnnotatedTurn, AnnotationResult, DependencyRequirement

BENCHMARK_HEADERS = (
    "编号",
    "当前问题",
    "最终回答",
    "是否需要前文",
    "关联前序对话",
    "所需前文信息",
    "依赖类型",
    "最大回溯轮数",
)


@dataclass(frozen=True)
class BenchmarkArtifacts:
    benchmark_path: Path
    manifest_path: Path
    dataset_hash: str
    valid_dependency_samples: int
    retrieval_required_samples: int
    retrieval_required_sessions: int
    shortterm_only_dependency_samples: int
    filtered_samples: int


def _gold_expression(requirements: tuple[DependencyRequirement, ...]) -> str:
    groups = []
    for requirement in requirements:
        members = list(requirement.dependency_ids)
        groups.append(members[0] if len(members) == 1 else f"（{'；'.join(members)}）")
    return "；".join(groups)


def _required_context_expression(requirements: tuple[DependencyRequirement, ...]) -> str:
    groups = []
    for requirement in requirements:
        contexts = list(requirement.required_contexts)
        groups.append(contexts[0] if len(contexts) == 1 else f"（{' OR '.join(contexts)}）")
    return "；".join(groups)


def _max_lookback(turn: AnnotatedTurn) -> int:
    current = int(turn.benchmark_id.rsplit("Q", 1)[1])
    positions = [
        int(dependency_id.rsplit("Q", 1)[1])
        for requirement in turn.label.requirements
        for dependency_id in requirement.dependency_ids
    ]
    return max((current - position for position in positions), default=0)


def _question_number(benchmark_id: str) -> int:
    return int(benchmark_id.rsplit("Q", 1)[1])


def requirement_is_satisfied_by_shortterm(
    requirement: DependencyRequirement,
    *,
    current_question_number: int,
    shortterm_qa_turns: int,
) -> bool:
    """Return whether any equivalent OR source is visible in ShortTerm."""

    return any(
        current_question_number - _question_number(dependency_id) <= shortterm_qa_turns
        for dependency_id in requirement.dependency_ids
    )


def turn_requires_memory_retrieval(turn: AnnotatedTurn, *, shortterm_qa_turns: int) -> bool:
    """Return whether an accepted Gold query has an unsatisfied AND requirement."""

    if shortterm_qa_turns < 0:
        raise ValueError("shortterm_qa_turns must be non-negative")
    if turn.status != "VALID" or not turn.label.needs_history:
        return False
    current_question_number = _question_number(turn.benchmark_id)
    return any(
        not requirement_is_satisfied_by_shortterm(
            requirement,
            current_question_number=current_question_number,
            shortterm_qa_turns=shortterm_qa_turns,
        )
        for requirement in turn.label.requirements
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_benchmark(
    annotations: AnnotationResult,
    *,
    output_dir: str | Path,
    history_db_path: str | Path,
    shortterm_qa_turns: int,
) -> BenchmarkArtifacts:
    destination = Path(output_dir).expanduser().resolve()
    destination.mkdir(parents=True, exist_ok=True)
    benchmark_path = destination / "generated_benchmark.xlsx"
    manifest_path = destination / "benchmark_manifest.json"

    workbook = Workbook()
    workbook.remove(workbook.active)
    manifest_sessions = []
    retrieval_required_samples = 0
    retrieval_required_sessions = 0
    shortterm_only_dependency_samples = 0
    for session in annotations.sessions:
        sheet = workbook.create_sheet(session.benchmark_session_id)
        sheet.append(BENCHMARK_HEADERS)
        manifest_turns = []
        session_requires_retrieval = False
        for turn in session.turns:
            accepted = turn.status == "VALID" and turn.label.needs_history
            retrieval_required = turn_requires_memory_retrieval(
                turn,
                shortterm_qa_turns=shortterm_qa_turns,
            )
            if retrieval_required:
                retrieval_required_samples += 1
                session_requires_retrieval = True
            elif accepted:
                shortterm_only_dependency_samples += 1
            requirements = turn.label.requirements if accepted else ()
            sheet.append(
                (
                    turn.benchmark_id,
                    turn.question,
                    turn.answer,
                    "是" if accepted else "否",
                    _gold_expression(requirements),
                    _required_context_expression(requirements),
                    turn.label.dependency_type if accepted else "",
                    _max_lookback(turn) if accepted else 0,
                )
            )
            manifest_turns.append(
                {
                    "benchmark_id": turn.benchmark_id,
                    "source_turn_index": turn.source_turn_index,
                    "label_status": turn.status,
                    "validation_issues": list(turn.validation_issues),
                    "included_as_gold": accepted,
                    "retrieval_required": retrieval_required,
                    "shortterm_only_dependency": accepted and not retrieval_required,
                }
            )
        if session_requires_retrieval:
            retrieval_required_sessions += 1
        manifest_sessions.append(
            {
                "benchmark_session_id": session.benchmark_session_id,
                "source_session_scope": session.source_session_scope,
                "user_id": session.user_id,
                "run_id": session.run_id,
                "agent_id": session.agent_id,
                "turns": manifest_turns,
            }
        )
    if not workbook.sheetnames:
        sheet = workbook.create_sheet("README")
        sheet.append(("No complete QA Sessions were extracted",))
    workbook.save(benchmark_path)
    dataset_hash = _sha256(benchmark_path)
    manifest = {
        "schema_version": 1,
        "user_id": annotations.user_id,
        "history_db_path": str(Path(history_db_path).expanduser().resolve()),
        "benchmark_path": str(benchmark_path),
        "dataset_hash": dataset_hash,
        "session_count": len(annotations.sessions),
        "qa_turn_count": annotations.qa_turn_count,
        "valid_dependency_samples": annotations.valid_dependency_samples,
        "shortterm_qa_turns": shortterm_qa_turns,
        "retrieval_required_samples": retrieval_required_samples,
        "retrieval_required_sessions": retrieval_required_sessions,
        "shortterm_only_dependency_samples": shortterm_only_dependency_samples,
        "filtered_samples": annotations.filtered_samples,
        "cross_session_tuning_status": "CROSS_SESSION_TUNING_UNSUPPORTED_NO_GOLD",
        "sessions": manifest_sessions,
    }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return BenchmarkArtifacts(
        benchmark_path=benchmark_path,
        manifest_path=manifest_path,
        dataset_hash=dataset_hash,
        valid_dependency_samples=annotations.valid_dependency_samples,
        retrieval_required_samples=retrieval_required_samples,
        retrieval_required_sessions=retrieval_required_sessions,
        shortterm_only_dependency_samples=shortterm_only_dependency_samples,
        filtered_samples=annotations.filtered_samples,
    )
