"""Offline Query-level completion analysis for the frozen S001-S005 recall run.

This script consumes the existing requirement/query result files and frozen
LongTerm provenance traces. It does not run retrieval, encode text, or invoke
an LLM.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np
from openpyxl import load_workbook

from exp.benchmark.benchmark_common import ensure_repo_root_on_path, load_json, resolve_path

REPO_ROOT = ensure_repo_root_on_path(Path(__file__))

from exp.benchmark.memory_gold_groups import GoldRequirement, parse_gold_requirements  # noqa: E402

DEFAULT_CONFIG = REPO_ROOT / "exp/benchmark/full_memory_recall_s001_s005_v2.json"
DEFAULT_RESULTS_DIR = REPO_ROOT / "exp/results/full_memory_recall_s001_s005_v2"
DEFAULT_MIDTERM_RANKING = (
    REPO_ROOT
    / "exp/results/midterm_bge_m3_multivector_late_interaction/rankings/c3_top60_multivector_rerank.jsonl"
)
MIDTERM_TOP_K = 5
HISTOGRAM_BINS = np.linspace(0.0, 100.0, 11)

COMPLETION_CONFIGS = (
    ("short", "Short", "shortterm_completion"),
    ("mid", "Mid", "midterm_completion"),
    ("long", "Long", "longterm_completion"),
    ("short_mid", "Short + Mid", "short_mid_completion"),
    ("short_long", "Short + Long", "short_long_completion"),
    ("mid_long", "Mid + Long", "mid_long_completion"),
    ("all_memory", "Short + Mid + Long", "all_memory_completion"),
)

EXPECTED_BASELINE = {
    "query_count": 348,
    "gold_requirement_count": 559,
    "or_requirement_count": 5,
    "shortterm_hit_count": 410,
    "midterm_eligible_gold_count": 149,
    "midterm_hit_at_5": 59,
    "overall_hit_count": 484,
}


@dataclass(frozen=True)
class DatasetTurn:
    session_code: str
    session_id: str
    turn_index: int
    query_id: str
    question: str
    answer: str
    gold_raw: str
    gold_groups: tuple[GoldRequirement, ...]

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="S001-S005 Query-level Gold completion offline analysis")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--dataset", type=Path, help="Gold workbook; source QA must match the frozen artifacts")
    parser.add_argument("--artifact-results-dir", type=Path, default=DEFAULT_RESULTS_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_RESULTS_DIR)
    parser.add_argument("--midterm-ranking", type=Path, default=DEFAULT_MIDTERM_RANKING)
    return parser.parse_args()


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_v3_dataset(path: Path, session_codes: Sequence[str]) -> dict[str, list[DatasetTurn]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sessions: dict[str, list[DatasetTurn]] = {}
    required_headers = {"编号", "当前问题", "最终回答", "关联前序对话"}
    try:
        for sheet_name in workbook.sheetnames:
            session_code = sheet_name.split("_", 1)[0]
            if session_code not in session_codes:
                continue
            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows)]
            header_index = {name: index for index, name in enumerate(headers) if name}
            missing_headers = required_headers - set(header_index)
            if missing_headers:
                raise ValueError(f"{sheet_name} missing columns: {sorted(missing_headers)}")
            turns: list[DatasetTurn] = []
            for row in rows:
                query_id = str(row[header_index["编号"]] or "").strip().upper()
                question = str(row[header_index["当前问题"]] or "").strip()
                answer = str(row[header_index["最终回答"]] or "").strip()
                if not query_id and not question and not answer:
                    continue
                gold_raw = str(row[header_index["关联前序对话"]] or "").strip()
                turns.append(
                    DatasetTurn(
                        session_code=session_code,
                        session_id=sheet_name,
                        turn_index=len(turns),
                        query_id=query_id,
                        question=question,
                        answer=answer,
                        gold_raw=gold_raw,
                        gold_groups=parse_gold_requirements(gold_raw),
                    )
                )
            sessions[session_code] = turns
    finally:
        workbook.close()
    if set(sessions) != set(session_codes):
        raise ValueError(f"Session mismatch: expected={list(session_codes)}, actual={sorted(sessions)}")
    validate_dataset_order(sessions)
    return sessions


def validate_dataset_order(sessions: Mapping[str, Sequence[DatasetTurn]]) -> None:
    for session_code, turns in sessions.items():
        positions = {turn.query_id: turn.turn_index for turn in turns}
        if len(positions) != len(turns):
            raise ValueError(f"Duplicate Query ID in {session_code}")
        for turn in turns:
            if not turn.query_id.startswith(f"{session_code}-Q"):
                raise ValueError(f"Query ID does not match Session: {turn.query_id} / {session_code}")
            for group in turn.gold_groups:
                for member in group.members:
                    if member not in positions:
                        raise ValueError(f"{turn.query_id} references missing Gold turn: {member}")
                    if positions[member] >= turn.turn_index:
                        raise ValueError(f"{turn.query_id} references current/future Gold turn: {member}")


def stable_json_hash(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def source_qa_hash(sessions: Mapping[str, Sequence[DatasetTurn]]) -> str:
    rows = [
        {
            "session_code": session_code,
            "query_id": turn.query_id,
            "question": turn.question,
            "answer": turn.answer,
        }
        for session_code, turns in sessions.items()
        for turn in turns
    ]
    return stable_json_hash(rows)


def group_rank(group: GoldRequirement, retrieved_turn_ids: Sequence[str]) -> int | None:
    ranks = {str(turn_id).upper(): rank for rank, turn_id in enumerate(retrieved_turn_ids, start=1)}
    matching_ranks = [ranks[member] for member in group.members if member in ranks]
    return min(matching_ranks) if matching_ranks else None


def static_dataset_stats(
    sessions: Mapping[str, Sequence[DatasetTurn]], shortterm_qa_turns: int, shortterm_top_k: int
) -> dict[str, Any]:
    queries = [turn for turns in sessions.values() for turn in turns]
    groups = [group for turn in queries for group in turn.gold_groups]
    effective_top_k = min(shortterm_qa_turns, shortterm_top_k)
    shortterm_hit_count = 0
    for turns in sessions.values():
        for turn in turns:
            visible_ids = {
                candidate.query_id for candidate in turns[max(0, turn.turn_index - effective_top_k) : turn.turn_index]
            }
            shortterm_hit_count += sum(group.hit_by(visible_ids) for group in turn.gold_groups)
    return {
        "session_count": len(sessions),
        "query_count": len(queries),
        "gold_requirement_count": len(groups),
        "or_requirement_count": sum(group.is_or for group in groups),
        "shortterm_requirement_count": shortterm_hit_count,
        "outside_shortterm_requirement_count": len(groups) - shortterm_hit_count,
        "shortterm_window_qa_turns": shortterm_qa_turns,
        "shortterm_top_k": effective_top_k,
        "session_query_counts": {session_code: len(turns) for session_code, turns in sessions.items()},
    }


def write_jsonl(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(dict(row), ensure_ascii=False) + "\n")


def write_csv(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def load_longterm_retrieved_ids(
    config: Mapping[str, Any], query_ids: set[str], top_k: int = 5
) -> dict[str, tuple[str, ...]]:
    """Load frozen production provenance without applying old eligibility filters."""

    rows_by_file: dict[Path, list[dict[str, Any]]] = {}
    retrieved_by_query: dict[str, tuple[str, ...]] = {}
    for session_code, raw_path in config["longterm"]["trace_files"].items():
        path = resolve_path(REPO_ROOT, raw_path)
        if path not in rows_by_file:
            rows_by_file[path] = load_jsonl(path)
        for row in rows_by_file[path]:
            query_id = str(row.get("turn_id", "")).upper()
            if not query_id.startswith(f"{session_code}-Q"):
                continue
            if row.get("error"):
                raise AssertionError(f"Frozen LongTerm trace error at {query_id}: {row['error']}")
            if "long_retrieved_turn_ids" not in row:
                raise AssertionError(f"Frozen LongTerm trace missing provenance at {query_id}")
            if query_id in retrieved_by_query:
                raise AssertionError(f"Duplicate frozen LongTerm trace: {query_id}")
            retrieved_by_query[query_id] = tuple(
                str(turn_id).upper() for turn_id in row["long_retrieved_turn_ids"][:top_k]
            )
    if set(retrieved_by_query) != query_ids:
        raise AssertionError(
            "Frozen LongTerm Query set mismatch: "
            f"missing={sorted(query_ids - set(retrieved_by_query))[:5]}, "
            f"extra={sorted(set(retrieved_by_query) - query_ids)[:5]}"
        )
    return retrieved_by_query


def load_frozen_midterm_rankings(path: Path) -> dict[str, tuple[str, ...]]:
    """Load the persisted complete C3 candidate order from a later ranking export.

    The export contains reranked rows as well, but ``c3_rank`` and
    ``source_turn_id`` preserve the original C3 ordering used by this analysis.
    """

    rows = load_jsonl(path)
    rows_by_query: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        query_id = str(row["query_id"]).upper()
        rows_by_query[query_id].append(row)

    rankings: dict[str, tuple[str, ...]] = {}
    for query_id, query_rows in rows_by_query.items():
        ordered = sorted(query_rows, key=lambda row: int(row["c3_rank"]))
        ranks = [int(row["c3_rank"]) for row in ordered]
        if ranks != list(range(1, len(ordered) + 1)):
            raise AssertionError(f"Frozen C3 ranking is not complete/contiguous: {query_id}")
        sources = tuple(str(row["source_turn_id"]).upper() for row in ordered)
        if len(set(sources)) != len(sources):
            raise AssertionError(f"Frozen C3 ranking contains duplicate source turns: {query_id}")
        rankings[query_id] = sources
    if not rankings:
        raise AssertionError(f"Frozen C3 ranking is empty: {path}")
    return rankings


def validate_frozen_midterm_ranking(
    rankings: Mapping[str, Sequence[str]],
    existing_requirement_rows: Sequence[Mapping[str, Any]],
    existing_summary: Mapping[str, Any],
) -> dict[str, Any]:
    """Prove that the persisted ranking reproduces the existing formal C3 metric."""

    if int(existing_summary["midterm"]["eligible_gold_count"]) != 149:
        raise AssertionError("Existing MidTerm eligible denominator no longer reproduces 149")
    if int(existing_summary["midterm"]["gold_at_5"]) != 59:
        raise AssertionError("Existing MidTerm C3 @5 no longer reproduces 59 hits")

    mismatches: list[str] = []
    reproduced_hits = 0
    eligible_count = 0
    for row in existing_requirement_rows:
        if not bool(row["midterm_eligible"]):
            continue
        eligible_count += 1
        query_id = str(row["query_id"]).upper()
        members = {str(member).upper() for member in row["gold_members"]}
        rank = next(
            (index for index, source_id in enumerate(rankings.get(query_id, ()), start=1) if source_id in members),
            None,
        )
        stored_rank = int(row["midterm_rank"]) if row.get("midterm_rank") is not None else None
        if rank != stored_rank:
            mismatches.append(str(row["requirement_id"]))
        reproduced_hits += rank is not None and rank <= MIDTERM_TOP_K
    if mismatches:
        raise AssertionError(f"Frozen C3 ranking differs from existing requirement ranks: {mismatches[:5]}")
    if eligible_count != 149 or reproduced_hits != 59:
        raise AssertionError(f"Frozen C3 reproduction failed: {reproduced_hits}/{eligible_count}")
    return {
        "ranking_query_count": len(rankings),
        "existing_eligible_gold_count": eligible_count,
        "existing_gold_at_5": reproduced_hits,
        "status": "PASS",
    }


def build_rechecked_requirement_rows(
    sessions: Mapping[str, Sequence[Any]],
    midterm_rankings: Mapping[str, Sequence[str]],
    longterm_retrieved_ids: Mapping[str, Sequence[str]],
    *,
    shortterm_qa_turns: int,
    shortterm_top_k: int,
    midterm_top_k: int = MIDTERM_TOP_K,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Re-label frozen retrieval results against the selected Gold workbook."""

    requirement_rows: list[dict[str, Any]] = []
    query_rows: list[dict[str, Any]] = []
    effective_shortterm_top_k = min(shortterm_qa_turns, shortterm_top_k)
    for session_code, turns in sessions.items():
        for turn in turns:
            short_ids = tuple(
                candidate.query_id
                for candidate in turns[max(0, turn.turn_index - effective_shortterm_top_k) : turn.turn_index]
            )
            mid_ids = midterm_rankings.get(turn.query_id, ())
            long_ids = longterm_retrieved_ids[turn.query_id]
            query_requirements: list[dict[str, Any]] = []
            for group_index, group in enumerate(turn.gold_groups, start=1):
                short_rank = group_rank(group, short_ids)
                short_hit = short_rank is not None
                midterm_rank = None if short_hit else group_rank(group, mid_ids)
                longterm_rank = group_rank(group, long_ids)
                row = {
                    "requirement_id": f"{turn.query_id}::G{group_index}",
                    "session_id": session_code,
                    "query_id": turn.query_id,
                    "turn_index": turn.turn_index,
                    "group_index": group_index,
                    "gold_members": list(group.members),
                    "is_or": group.is_or,
                    "shortterm_rank": short_rank,
                    "shortterm_hit": short_hit,
                    "midterm_ranking_available": turn.query_id in midterm_rankings,
                    "midterm_rank": midterm_rank,
                    f"midterm_hit_at_{midterm_top_k}": midterm_rank is not None and midterm_rank <= midterm_top_k,
                    "longterm_rank": longterm_rank,
                    "longterm_hit": longterm_rank is not None,
                }
                row["final_hit"] = bool(
                    short_hit or row[f"midterm_hit_at_{midterm_top_k}"] or longterm_rank is not None
                )
                requirement_rows.append(row)
                query_requirements.append(row)
            query_rows.append(
                {
                    "session_id": session_code,
                    "query_id": turn.query_id,
                    "turn_index": turn.turn_index,
                    "parsed_gold_groups": [list(group.members) for group in turn.gold_groups],
                    "shortterm_hit": [row["shortterm_hit"] for row in query_requirements],
                    "midterm_rank": [row["midterm_rank"] for row in query_requirements],
                    "longterm_hit": [row["longterm_hit"] for row in query_requirements],
                    "final_hit": [row["final_hit"] for row in query_requirements],
                }
            )
    return requirement_rows, query_rows


def _completion(hit_count: int, requirement_count: int) -> float:
    if requirement_count <= 0:
        raise ValueError("Query completion requires at least one Gold requirement")
    return hit_count / requirement_count


def build_query_completion_rows(
    requirement_rows: Sequence[Mapping[str, Any]],
    query_rows: Sequence[Mapping[str, Any]],
    longterm_retrieved_ids: Mapping[str, Sequence[str]],
    midterm_top_k: int = MIDTERM_TOP_K,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Aggregate layer hits after requirement-level OR, then compute Query completion."""

    requirements_by_query: dict[str, list[Mapping[str, Any]]] = defaultdict(list)
    seen_requirement_ids: set[str] = set()
    for requirement in requirement_rows:
        requirement_id = str(requirement["requirement_id"])
        if requirement_id in seen_requirement_ids:
            raise AssertionError(f"Duplicate Gold requirement: {requirement_id}")
        seen_requirement_ids.add(requirement_id)
        requirements_by_query[str(requirement["query_id"]).upper()].append(requirement)

    detail_rows: list[dict[str, Any]] = []
    layer_requirement_hits = {key: 0 for key, _, _ in COMPLETION_CONFIGS}
    longterm_scope_difference_ids: list[str] = []
    no_gold_query_count = 0
    seen_query_ids: set[str] = set()
    for query in query_rows:
        query_id = str(query["query_id"]).upper()
        if query_id in seen_query_ids:
            raise AssertionError(f"Duplicate Query result: {query_id}")
        seen_query_ids.add(query_id)
        requirements = requirements_by_query.get(query_id, [])
        if not requirements:
            no_gold_query_count += 1
            continue

        direct_long_ids = {str(turn_id).upper() for turn_id in longterm_retrieved_ids[query_id]}
        hits: list[tuple[bool, bool, bool]] = []
        for requirement in requirements:
            members = tuple(str(member).upper() for member in requirement["gold_members"])
            if not members or len(set(members)) != len(members):
                raise AssertionError(f"Invalid Gold group members: {requirement['requirement_id']}")
            short_hit = bool(requirement["shortterm_hit"])
            midterm_rank = requirement.get("midterm_rank")
            midterm_hit = midterm_rank is not None and int(midterm_rank) <= midterm_top_k
            stored_midterm_hit = bool(requirement.get(f"midterm_hit_at_{midterm_top_k}", midterm_hit))
            if stored_midterm_hit != midterm_hit:
                raise AssertionError(f"MidTerm rank/hit mismatch: {requirement['requirement_id']}")
            longterm_hit = any(member in direct_long_ids for member in members)
            if "longterm_hit" in requirement and bool(requirement["longterm_hit"]) != longterm_hit:
                longterm_scope_difference_ids.append(str(requirement["requirement_id"]))
            hits.append((short_hit, midterm_hit, longterm_hit))

        requirement_count = len(requirements)
        hit_counts = {
            "short": sum(short for short, _, _ in hits),
            "mid": sum(mid for _, mid, _ in hits),
            "long": sum(long for _, _, long in hits),
            "short_mid": sum(short or mid for short, mid, _ in hits),
            "short_long": sum(short or long for short, _, long in hits),
            "mid_long": sum(mid or long for _, mid, long in hits),
            "all_memory": sum(short or mid or long for short, mid, long in hits),
        }
        for key, count in hit_counts.items():
            layer_requirement_hits[key] += count

        detail_rows.append(
            {
                "session_id": str(query["session_id"]),
                "query_id": query_id,
                "gold_requirement_count": requirement_count,
                "shortterm_hit_count": hit_counts["short"],
                "midterm_hit_count": hit_counts["mid"],
                "longterm_hit_count": hit_counts["long"],
                "shortterm_completion": _completion(hit_counts["short"], requirement_count),
                "midterm_completion": _completion(hit_counts["mid"], requirement_count),
                "longterm_completion": _completion(hit_counts["long"], requirement_count),
                "short_mid_completion": _completion(hit_counts["short_mid"], requirement_count),
                "short_long_completion": _completion(hit_counts["short_long"], requirement_count),
                "mid_long_completion": _completion(hit_counts["mid_long"], requirement_count),
                "all_memory_completion": _completion(hit_counts["all_memory"], requirement_count),
            }
        )

    if seen_query_ids != set(longterm_retrieved_ids):
        raise AssertionError("Query results and frozen LongTerm traces do not contain the same Query IDs")
    if set(requirements_by_query) - seen_query_ids:
        raise AssertionError(f"Requirement rows reference unknown Queries: {sorted(set(requirements_by_query)-seen_query_ids)[:5]}")
    return detail_rows, {
        **layer_requirement_hits,
        "no_gold_query_count": no_gold_query_count,
        "longterm_scope_difference_requirement_ids": longterm_scope_difference_ids,
    }


def completion_statistics(
    detail_rows: Sequence[Mapping[str, Any]], requirement_hits: Mapping[str, Any]
) -> dict[str, dict[str, Any]]:
    statistics_by_config: dict[str, dict[str, Any]] = {}
    for key, label, field in COMPLETION_CONFIGS:
        values = np.asarray([float(row[field]) for row in detail_rows], dtype=float)
        if not len(values):
            raise AssertionError("No Gold-bearing Queries found")
        zero_count = int(np.count_nonzero(values == 0.0))
        full_count = int(np.count_nonzero(values == 1.0))
        statistics_by_config[key] = {
            "label": label,
            "query_count": len(values),
            "requirement_hit_count": int(requirement_hits[key]),
            "mean": float(np.mean(values)),
            "median": float(np.median(values)),
            "p25": float(np.percentile(values, 25)),
            "p75": float(np.percentile(values, 75)),
            "zero_complete_query_count": zero_count,
            "zero_complete_query_rate": zero_count / len(values),
            "full_complete_query_count": full_count,
            "full_complete_query_rate": full_count / len(values),
            "completion_ge_50_rate": float(np.mean(values >= 0.5)),
            "completion_ge_80_rate": float(np.mean(values >= 0.8)),
        }
    return statistics_by_config


def validate_baseline_and_completion(
    dataset_stats: Mapping[str, Any],
    baseline_summary: Mapping[str, Any],
    requirement_rows: Sequence[Mapping[str, Any]],
    query_rows: Sequence[Mapping[str, Any]],
    detail_rows: Sequence[Mapping[str, Any]],
    requirement_hits: Mapping[str, Any],
) -> dict[str, Any]:
    actual_baseline = {
        "query_count": int(dataset_stats["query_count"]),
        "gold_requirement_count": int(dataset_stats["gold_requirement_count"]),
        "or_requirement_count": int(dataset_stats["or_requirement_count"]),
        "shortterm_hit_count": int(baseline_summary["shortterm"]["hit_count"]),
        "midterm_eligible_gold_count": int(baseline_summary["midterm"]["eligible_gold_count"]),
        "midterm_hit_at_5": int(baseline_summary["midterm"]["gold_at_5"]),
        "overall_hit_count": int(baseline_summary["overall"]["hit_count"]),
    }
    if actual_baseline != EXPECTED_BASELINE:
        raise AssertionError(f"Existing baseline reproduction failed: {actual_baseline}")
    if len(query_rows) != EXPECTED_BASELINE["query_count"]:
        raise AssertionError(f"Query result row count mismatch: {len(query_rows)}")
    if len(requirement_rows) != EXPECTED_BASELINE["gold_requirement_count"]:
        raise AssertionError(f"Gold requirement row count mismatch: {len(requirement_rows)}")
    if sum(bool(row["is_or"]) for row in requirement_rows) != EXPECTED_BASELINE["or_requirement_count"]:
        raise AssertionError("OR group count mismatch")
    if requirement_hits["short"] != EXPECTED_BASELINE["shortterm_hit_count"]:
        raise AssertionError(f"ShortTerm requirement hits changed: {requirement_hits['short']}")
    if requirement_hits["mid"] != EXPECTED_BASELINE["midterm_hit_at_5"]:
        raise AssertionError(f"MidTerm@5 requirement hits changed: {requirement_hits['mid']}")
    if requirement_hits["all_memory"] != EXPECTED_BASELINE["overall_hit_count"]:
        raise AssertionError(f"Three-layer union must reproduce 484/559, got {requirement_hits['all_memory']}/559")
    if sum(int(row["gold_requirement_count"]) for row in detail_rows) != len(requirement_rows):
        raise AssertionError("Gold requirements were dropped or counted more than once")

    for row in detail_rows:
        short = float(row["shortterm_completion"])
        mid = float(row["midterm_completion"])
        long = float(row["longterm_completion"])
        short_mid = float(row["short_mid_completion"])
        short_long = float(row["short_long_completion"])
        mid_long = float(row["mid_long_completion"])
        all_memory = float(row["all_memory_completion"])
        if short_mid < max(short, mid) or short_long < max(short, long) or mid_long < max(mid, long):
            raise AssertionError(f"Pairwise completion monotonicity failed: {row['query_id']}")
        if all_memory < max(short_mid, short_long, mid_long):
            raise AssertionError(f"Three-layer completion monotonicity failed: {row['query_id']}")

    existing_longterm_hit_count = int(baseline_summary["longterm"]["hit_count"])
    stored_longterm_hit_count = sum(bool(row["longterm_hit"]) for row in requirement_rows)
    if stored_longterm_hit_count != existing_longterm_hit_count:
        raise AssertionError(
            f"Stored LongTerm requirement hits differ from summary: {stored_longterm_hit_count} != "
            f"{existing_longterm_hit_count}"
        )
    return {
        "status": "PASS",
        "baseline": actual_baseline,
        "gold_bearing_query_count": len(detail_rows),
        "no_gold_query_count": int(requirement_hits["no_gold_query_count"]),
        "direct_longterm_requirement_hit_count": int(requirement_hits["long"]),
        "existing_eligible_longterm_hit_count": existing_longterm_hit_count,
        "longterm_count_difference": int(requirement_hits["long"]) - existing_longterm_hit_count,
        "longterm_scope_difference_requirement_ids": list(
            requirement_hits["longterm_scope_difference_requirement_ids"]
        ),
        "all_memory_requirement_hit_count": int(requirement_hits["all_memory"]),
        "or_groups_counted_once": True,
        "query_completion_monotonicity": True,
    }


def validate_rechecked_completion(
    dataset_stats: Mapping[str, Any],
    requirement_rows: Sequence[Mapping[str, Any]],
    query_rows: Sequence[Mapping[str, Any]],
    detail_rows: Sequence[Mapping[str, Any]],
    requirement_hits: Mapping[str, Any],
) -> dict[str, Any]:
    """Validate dynamic Gold counts plus all requirement/Query invariants."""

    expected_query_count = int(dataset_stats["query_count"])
    expected_requirement_count = int(dataset_stats["gold_requirement_count"])
    expected_or_count = int(dataset_stats["or_requirement_count"])
    expected_shortterm_hits = int(dataset_stats["shortterm_requirement_count"])
    midterm_eligible_rows = [row for row in requirement_rows if not bool(row["shortterm_hit"])]
    midterm_evaluated_query_ids = {
        str(row["query_id"]) for row in midterm_eligible_rows if bool(row.get("midterm_ranking_available"))
    }
    midterm_unrouted_query_ids = sorted(
        {str(row["query_id"]) for row in midterm_eligible_rows if not bool(row.get("midterm_ranking_available"))}
    )
    if len(query_rows) != expected_query_count:
        raise AssertionError(f"Query result row count mismatch: {len(query_rows)} != {expected_query_count}")
    if len(requirement_rows) != expected_requirement_count:
        raise AssertionError(
            f"Gold requirement row count mismatch: {len(requirement_rows)} != {expected_requirement_count}"
        )
    if sum(bool(row["is_or"]) for row in requirement_rows) != expected_or_count:
        raise AssertionError("OR group count mismatch")
    if int(requirement_hits["short"]) != expected_shortterm_hits:
        raise AssertionError(
            f"ShortTerm requirement hits mismatch: {requirement_hits['short']} != {expected_shortterm_hits}"
        )
    if sum(int(row["gold_requirement_count"]) for row in detail_rows) != expected_requirement_count:
        raise AssertionError("Gold requirements were dropped or counted more than once")

    for row in detail_rows:
        short = float(row["shortterm_completion"])
        mid = float(row["midterm_completion"])
        long = float(row["longterm_completion"])
        short_mid = float(row["short_mid_completion"])
        short_long = float(row["short_long_completion"])
        mid_long = float(row["mid_long_completion"])
        all_memory = float(row["all_memory_completion"])
        if short_mid < max(short, mid) or short_long < max(short, long) or mid_long < max(mid, long):
            raise AssertionError(f"Pairwise completion monotonicity failed: {row['query_id']}")
        if all_memory < max(short_mid, short_long, mid_long):
            raise AssertionError(f"Three-layer completion monotonicity failed: {row['query_id']}")

    return {
        "status": "PASS",
        "dataset": {
            "query_count": expected_query_count,
            "gold_requirement_count": expected_requirement_count,
            "or_requirement_count": expected_or_count,
            "shortterm_hit_count": expected_shortterm_hits,
        },
        "gold_bearing_query_count": len(detail_rows),
        "no_gold_query_count": int(requirement_hits["no_gold_query_count"]),
        "midterm_hit_at_5": int(requirement_hits["mid"]),
        "midterm_eligible_requirement_count": len(midterm_eligible_rows),
        "midterm_evaluated_query_count": len(midterm_evaluated_query_ids),
        "midterm_no_ranking_query_count": len(midterm_unrouted_query_ids),
        "midterm_no_ranking_query_ids": midterm_unrouted_query_ids,
        "direct_longterm_requirement_hit_count": int(requirement_hits["long"]),
        "all_memory_requirement_hit_count": int(requirement_hits["all_memory"]),
        "longterm_scope_difference_requirement_ids": list(
            requirement_hits["longterm_scope_difference_requirement_ids"]
        ),
        "or_groups_counted_once": True,
        "query_completion_monotonicity": True,
    }


def _histogram_y_limit(detail_rows: Sequence[Mapping[str, Any]]) -> int:
    maximum = 0
    for _, _, field in COMPLETION_CONFIGS:
        counts, _ = np.histogram([float(row[field]) * 100 for row in detail_rows], bins=HISTOGRAM_BINS)
        maximum = max(maximum, int(counts.max()))
    return max(10, int(math.ceil(maximum * 1.08 / 10.0) * 10))


def _plot_histogram(
    axis: Any,
    values: Sequence[float],
    stats: Mapping[str, Any],
    title: str,
    y_limit: int,
    highlight_extremes: bool = False,
) -> None:
    counts, edges = np.histogram(np.asarray(values) * 100, bins=HISTOGRAM_BINS)
    colors = ["#4C78A8"] * len(counts)
    if highlight_extremes:
        colors[0] = "#D55E00"
        colors[-1] = "#009E73"
    axis.bar(edges[:-1], counts, width=np.diff(edges) * 0.9, align="edge", color=colors, edgecolor="white")
    axis.set_xlim(0, 100)
    axis.set_ylim(0, y_limit)
    axis.set_xticks(np.arange(0, 101, 20))
    axis.grid(axis="y", alpha=0.25, linewidth=0.7)
    axis.set_xlabel("Query Completion (%)")
    axis.set_ylabel("Query Count")
    axis.set_title(
        f"{title}\nMean {stats['mean']:.1%} | Median {stats['median']:.1%} | "
        f"100% Complete {stats['full_complete_query_rate']:.1%}",
        fontsize=10,
    )


def plot_completion_matrix(
    output_path: Path,
    detail_rows: Sequence[Mapping[str, Any]],
    statistics_by_config: Mapping[str, Mapping[str, Any]],
    y_limit: int,
) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    matrix = (
        (("short", "shortterm_completion"), ("short_mid", "short_mid_completion"), ("short_long", "short_long_completion")),
        (("short_mid", "short_mid_completion"), ("mid", "midterm_completion"), ("mid_long", "mid_long_completion")),
        (("short_long", "short_long_completion"), ("mid_long", "mid_long_completion"), ("long", "longterm_completion")),
    )
    column_labels = ("ShortTerm", "MidTerm@5", "LongTerm@5")
    row_labels = ("ShortTerm", "MidTerm@5", "LongTerm@5")
    figure, axes = plt.subplots(3, 3, figsize=(18, 14), sharex=True, sharey=True)
    for row_index, row in enumerate(matrix):
        for column_index, (key, field) in enumerate(row):
            _plot_histogram(
                axes[row_index, column_index],
                [float(detail[field]) for detail in detail_rows],
                statistics_by_config[key],
                str(statistics_by_config[key]["label"]),
                y_limit,
            )
            if row_index == 0:
                axes[row_index, column_index].annotate(
                    column_labels[column_index],
                    xy=(0.5, 1.25),
                    xycoords="axes fraction",
                    ha="center",
                    va="bottom",
                    fontsize=13,
                    fontweight="bold",
                )
            if column_index == 0:
                axes[row_index, column_index].annotate(
                    row_labels[row_index],
                    xy=(-0.28, 0.5),
                    xycoords="axes fraction",
                    ha="right",
                    va="center",
                    rotation=90,
                    fontsize=13,
                    fontweight="bold",
                )
    figure.suptitle("Query Completion Matrix (S001-S005, Gold-bearing Queries)", fontsize=16, y=0.995)
    figure.tight_layout(rect=(0.04, 0.02, 1, 0.97), h_pad=2.8, w_pad=1.6)
    figure.savefig(output_path, dpi=220, bbox_inches="tight")
    plt.close(figure)


def plot_all_memory_completion(
    output_path: Path,
    detail_rows: Sequence[Mapping[str, Any]],
    statistics_by_config: Mapping[str, Mapping[str, Any]],
    y_limit: int,
) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    stats = statistics_by_config["all_memory"]
    values = [float(row["all_memory_completion"]) for row in detail_rows]
    partial_count = int(stats["query_count"]) - int(stats["zero_complete_query_count"]) - int(
        stats["full_complete_query_count"]
    )
    figure, axis = plt.subplots(figsize=(11, 6.5))
    _plot_histogram(
        axis,
        values,
        stats,
        "ShortTerm + MidTerm + LongTerm",
        y_limit,
        highlight_extremes=True,
    )
    axis.text(
        0.025,
        0.94,
        f"Exact 0%: {stats['zero_complete_query_count']} ({stats['zero_complete_query_rate']:.1%})\n"
        f"Partial: {partial_count} ({partial_count / stats['query_count']:.1%})\n"
        f"Exact 100%: {stats['full_complete_query_count']} ({stats['full_complete_query_rate']:.1%})",
        transform=axis.transAxes,
        va="top",
        fontsize=11,
        bbox={"boxstyle": "round,pad=0.45", "facecolor": "white", "edgecolor": "#777777", "alpha": 0.92},
    )
    figure.tight_layout()
    figure.savefig(output_path, dpi=220, bbox_inches="tight")
    plt.close(figure)


def render_summary(
    statistics_by_config: Mapping[str, Mapping[str, Any]],
    validation: Mapping[str, Any],
    dataset_path: Path,
    frozen_reproduction: Mapping[str, Any],
) -> str:
    dataset = validation["dataset"]
    lines = [
        "# Query Completion Summary (S001-S005)",
        "",
        f"- Gold workbook: `{dataset_path}`",
        f"- Dataset: {dataset['query_count']} Queries, {dataset['gold_requirement_count']} Gold requirements, "
        f"{dataset['or_requirement_count']} OR groups",
        f"- Gold-bearing Queries: {validation['gold_bearing_query_count']}",
        f"- Queries without Gold requirements (excluded): {validation['no_gold_query_count']}",
        "- Completion values use each Query's complete Gold requirement set as the denominator.",
        "",
        "| Configuration | Queries | Mean | Median | P25 | P75 | 0% Complete | 100% Complete | >=50% | >=80% |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for key, label, _ in COMPLETION_CONFIGS:
        stats = statistics_by_config[key]
        lines.append(
            f"| {label} | {stats['query_count']} | {stats['mean']:.2%} | {stats['median']:.2%} | "
            f"{stats['p25']:.2%} | {stats['p75']:.2%} | "
            f"{stats['zero_complete_query_count']} ({stats['zero_complete_query_rate']:.2%}) | "
            f"{stats['full_complete_query_count']} ({stats['full_complete_query_rate']:.2%}) | "
            f"{stats['completion_ge_50_rate']:.2%} | {stats['completion_ge_80_rate']:.2%} |"
        )
    lines.extend(
        [
            "",
            "## Validation and metric scope",
            "",
            "- Frozen artifact baseline reproduced before relabeling: 348 Queries, 559 Gold requirements, 5 OR groups, "
            "ShortTerm 410/559, MidTerm C3 @5 59/149, Overall Memory Coverage 484/559.",
            f"- Persisted C3 ranking reproduction: {frozen_reproduction['existing_gold_at_5']}/"
            f"{frozen_reproduction['existing_eligible_gold_count']} (PASS).",
            f"- Rechecked ShortTerm: {dataset['shortterm_hit_count']}/{dataset['gold_requirement_count']}.",
            f"- Rechecked routed MidTerm C3 @5 contribution: {validation['midterm_hit_at_5']}/"
            f"{dataset['gold_requirement_count']}.",
            f"- New Gold contains {validation['midterm_eligible_requirement_count']} requirements outside ShortTerm; "
            f"{validation['midterm_no_ranking_query_count']} corresponding Queries have no frozen C3 ranking and are "
            "counted as MidTerm misses under the routing contract.",
            f"- Rechecked direct LongTerm@5 provenance hits: "
            f"{validation['direct_longterm_requirement_hit_count']}/{dataset['gold_requirement_count']}.",
            f"- Rechecked three-layer requirement union: {validation['all_memory_requirement_hit_count']}/"
            f"{dataset['gold_requirement_count']}.",
            "- All pairwise and three-layer per-Query monotonicity checks passed; every OR group is one denominator unit.",
        ]
    )
    return "\n".join(lines) + "\n"


def main() -> int:
    args = parse_args()
    config_path = args.config.resolve()
    artifact_results_dir = args.artifact_results_dir.resolve()
    output_dir = args.output_dir.resolve()
    config = load_json(config_path)
    artifact_dataset_path = resolve_path(REPO_ROOT, config["dataset"])
    dataset_path = (args.dataset or artifact_dataset_path).resolve()
    artifact_dataset_stats = load_json(artifact_results_dir / "dataset_stats.json")
    artifact_summary = load_json(artifact_results_dir / "summary.json")
    artifact_requirement_rows = load_jsonl(artifact_results_dir / "requirement_results.jsonl")
    artifact_query_rows = load_jsonl(artifact_results_dir / "query_results.jsonl")
    query_ids = {str(row["query_id"]).upper() for row in artifact_query_rows}

    existing_baseline = {
        "query_count": int(artifact_dataset_stats["query_count"]),
        "gold_requirement_count": int(artifact_dataset_stats["gold_requirement_count"]),
        "or_requirement_count": int(artifact_dataset_stats["or_requirement_count"]),
        "shortterm_hit_count": int(artifact_summary["shortterm"]["hit_count"]),
        "midterm_eligible_gold_count": int(artifact_summary["midterm"]["eligible_gold_count"]),
        "midterm_hit_at_5": int(artifact_summary["midterm"]["gold_at_5"]),
        "overall_hit_count": int(artifact_summary["overall"]["hit_count"]),
    }
    if existing_baseline != EXPECTED_BASELINE:
        raise AssertionError(f"Existing frozen baseline reproduction failed: {existing_baseline}")

    sessions = load_v3_dataset(dataset_path, tuple(config["sessions"]))
    artifact_sessions = load_v3_dataset(artifact_dataset_path, tuple(config["sessions"]))
    if source_qa_hash(sessions) != source_qa_hash(artifact_sessions):
        raise AssertionError("Selected workbook source QA differs from frozen retrieval artifacts; cannot relabel offline")
    dataset_stats = static_dataset_stats(
        sessions, int(config["shortterm"]["qa_turns"]), int(config["shortterm"]["top_k"])
    )

    midterm_ranking_path = args.midterm_ranking.resolve()
    midterm_rankings = load_frozen_midterm_rankings(midterm_ranking_path)
    frozen_reproduction = validate_frozen_midterm_ranking(
        midterm_rankings, artifact_requirement_rows, artifact_summary
    )
    longterm_ids = load_longterm_retrieved_ids(config, query_ids, top_k=MIDTERM_TOP_K)
    requirement_rows, query_rows = build_rechecked_requirement_rows(
        sessions,
        midterm_rankings,
        longterm_ids,
        shortterm_qa_turns=int(config["shortterm"]["qa_turns"]),
        shortterm_top_k=int(config["shortterm"]["top_k"]),
    )

    detail_rows, requirement_hits = build_query_completion_rows(requirement_rows, query_rows, longterm_ids)
    statistics_by_config = completion_statistics(detail_rows, requirement_hits)
    validation = validate_rechecked_completion(
        dataset_stats,
        requirement_rows,
        query_rows,
        detail_rows,
        requirement_hits,
    )

    output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(output_dir / "requirement_results.csv", requirement_rows)
    write_jsonl(output_dir / "requirement_results.jsonl", requirement_rows)
    write_csv(output_dir / "query_results.csv", query_rows)
    write_jsonl(output_dir / "query_results.jsonl", query_rows)
    write_csv(output_dir / "query_completion_details.csv", detail_rows)
    write_jsonl(output_dir / "query_completion_details.jsonl", detail_rows)
    (output_dir / "dataset_stats.json").write_text(
        json.dumps(dataset_stats, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (output_dir / "query_completion_stats.json").write_text(
        json.dumps(
            {
                "configurations": statistics_by_config,
                "validation": validation,
                "frozen_artifact_reproduction": frozen_reproduction,
                "histogram_bins_percent": HISTOGRAM_BINS.tolist(),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    (output_dir / "run_metadata.json").write_text(
        json.dumps(
            {
                "mode": "offline_gold_relabeling",
                "dataset": str(dataset_path),
                "dataset_sha256": sha256_file(dataset_path),
                "artifact_dataset": str(artifact_dataset_path),
                "artifact_dataset_sha256": sha256_file(artifact_dataset_path),
                "source_qa_sha256": source_qa_hash(sessions),
                "source_qa_matches_frozen_artifacts": True,
                "artifact_results_dir": str(artifact_results_dir),
                "midterm_ranking": str(midterm_ranking_path),
                "midterm_ranking_sha256": sha256_file(midterm_ranking_path),
                "frozen_artifact_reproduction": frozen_reproduction,
                "new_llm_calls": 0,
                "new_embedding_count": 0,
                "midterm_page_regeneration": False,
                "session_rerun": False,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    (output_dir / "query_completion_summary.md").write_text(
        render_summary(statistics_by_config, validation, dataset_path, frozen_reproduction), encoding="utf-8"
    )

    y_limit = _histogram_y_limit(detail_rows)
    plot_completion_matrix(output_dir / "query_completion_matrix.png", detail_rows, statistics_by_config, y_limit)
    plot_all_memory_completion(
        output_dir / "query_completion_all_memory.png", detail_rows, statistics_by_config, y_limit
    )

    print((output_dir / "query_completion_summary.md").read_text(encoding="utf-8"))
    print(f"Validation: {validation['status']}")
    print(f"Results directory: {output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
