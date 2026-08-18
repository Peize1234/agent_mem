from __future__ import annotations

import asyncio
import csv
import json
import math
import os
import re
import shutil
import sqlite3
import statistics
import sys
import time
from collections import defaultdict, deque
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

TURN_ID_PATTERN = re.compile(r"S\d{3}-Q\d{3}", re.IGNORECASE)
MIGRATION_TERMINAL_STATUSES = {"succeeded", "succeeded_degraded", "completed_with_loss"}
PROFILE_TERMINAL_STATUSES = {"succeeded", "discarded"}
BEIJING_TIMEZONE = timezone(timedelta(hours=8))


def find_repo_root(start: Path | None = None) -> Path:
    current = (start or Path(__file__)).resolve()
    if current.is_file():
        current = current.parent
    for candidate in [current, *current.parents]:
        if (candidate / "pyproject.toml").exists() and (candidate / "mem0").exists():
            return candidate
    raise RuntimeError("无法定位仓库根目录：未找到同时包含 pyproject.toml 和 mem0/ 的目录")


def ensure_repo_root_on_path(start: Path | None = None) -> Path:
    root = find_repo_root(start)
    root_text = str(root)
    if root_text not in sys.path:
        sys.path.insert(0, root_text)
    return root


def resolve_path(repo_root: Path, value: str | os.PathLike[str]) -> Path:
    path = Path(os.path.expanduser(os.path.expandvars(str(value))))
    return path if path.is_absolute() else (repo_root / path).resolve()


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as file:
        value = json.load(file)
    if not isinstance(value, dict):
        raise ValueError(f"JSON 顶层必须是对象：{path}")
    return value


def dump_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        json.dump(value, file, ensure_ascii=False, indent=2, default=str)


def expand_env_placeholders(value: Any) -> Any:
    """展开已存在的环境变量；缺失的占位符保留，便于 Mock 模式直接运行。"""
    if isinstance(value, dict):
        return {key: expand_env_placeholders(item) for key, item in value.items()}
    if isinstance(value, list):
        return [expand_env_placeholders(item) for item in value]
    if isinstance(value, str):
        match = re.fullmatch(r"\$\{([A-Za-z_][A-Za-z0-9_]*)\}", value)
        if match:
            variable = match.group(1)
            return os.environ.get(variable, value)
        return os.path.expandvars(value)
    return value


_SENSITIVE_KEYS = {
    "api_key",
    "password",
    "secret",
    "secret_key",
    "private_key",
    "access_key",
    "credentials",
    "credential",
    "token",
    "access_token",
    "refresh_token",
    "auth_token",
    "session_token",
    "client_secret",
}
_SENSITIVE_SUFFIXES = ("_password", "_secret", "_token", "_credential", "_credentials")


def redact_secrets(value: Any) -> Any:
    """在写入结果目录前递归隐藏密钥、令牌和密码。"""
    if isinstance(value, Mapping):
        result: dict[str, Any] = {}
        for key, item in value.items():
            key_text = str(key).lower()
            if key_text in _SENSITIVE_KEYS or any(key_text.endswith(suffix) for suffix in _SENSITIVE_SUFFIXES):
                result[str(key)] = "***REDACTED***" if item not in (None, "") else item
            else:
                result[str(key)] = redact_secrets(item)
        return result
    if isinstance(value, list):
        return [redact_secrets(item) for item in value]
    if isinstance(value, tuple):
        return tuple(redact_secrets(item) for item in value)
    return value


def prepare_runtime_config(
    base_config: Mapping[str, Any],
    *,
    runtime_dir: Path,
    collection_name: str,
    reset_storage: bool,
    agentic_retrieval_enabled: bool,
    profile_enabled: bool,
    profile_update_on_add: bool,
) -> dict[str, Any]:
    runtime_dir = runtime_dir.resolve()
    if reset_storage:
        shutil.rmtree(runtime_dir, ignore_errors=True)
    runtime_dir.mkdir(parents=True, exist_ok=True)

    config = expand_env_placeholders(deepcopy(dict(base_config)))
    config["history_db_path"] = str(runtime_dir / "history.db")
    config.setdefault("vector_store", {}).setdefault("config", {})
    config["vector_store"]["config"]["path"] = str(runtime_dir / "qdrant")
    config["vector_store"]["config"]["collection_name"] = collection_name
    config.setdefault("agentic_retrieval", {})["enabled"] = bool(agentic_retrieval_enabled)
    config.setdefault("profile", {})["enabled"] = bool(profile_enabled)
    config["profile"]["update_on_add"] = bool(profile_update_on_add)
    config.setdefault("background", {})["enabled"] = True
    return config


@dataclass(frozen=True)
class BenchmarkTurn:
    session_id: str
    sheet_name: str
    turn_index: int
    turn_id: str
    question: str
    answer: str
    needs_history: bool
    dependency_turn_ids: tuple[str, ...]
    required_context: str
    dependency_type: str
    max_lookback: int


@dataclass(frozen=True)
class BenchmarkSession:
    session_id: str
    sheet_name: str
    turns: tuple[BenchmarkTurn, ...]

    @property
    def turn_position_by_id(self) -> dict[str, int]:
        return {turn.turn_id: turn.turn_index for turn in self.turns}


def _normalise_header(value: Any) -> str:
    return str(value or "").strip()


def _to_bool(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return text in {"是", "true", "1", "yes", "y"}


def _to_int(value: Any, default: int = 0) -> int:
    if value is None or value == "":
        return default
    try:
        return int(value)
    except (TypeError, ValueError):
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return default


def parse_dependency_turn_ids(value: Any) -> tuple[str, ...]:
    if value is None:
        return ()
    matches = TURN_ID_PATTERN.findall(str(value))
    seen: set[str] = set()
    ordered: list[str] = []
    for match in matches:
        turn_id = match.upper()
        if turn_id not in seen:
            seen.add(turn_id)
            ordered.append(turn_id)
    return tuple(ordered)


def load_dataset(
    path: Path,
    *,
    include_sheets: Sequence[str] | None = None,
    max_sessions: int | None = None,
    max_turns_per_session: int | None = None,
) -> list[BenchmarkSession]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    requested = set(include_sheets or [])
    sessions: list[BenchmarkSession] = []
    required_columns = {"编号", "当前问题", "最终回答", "是否需要前文", "关联前序对话"}

    try:
        for sheet_name in workbook.sheetnames:
            if requested and sheet_name not in requested:
                continue
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            headers = [_normalise_header(value) for value in header_row]
            index_by_name = {name: index for index, name in enumerate(headers) if name}
            missing = sorted(required_columns - set(index_by_name))
            if missing:
                if requested:
                    raise ValueError(f"Sheet {sheet_name} 缺少字段：{', '.join(missing)}")
                # Workbooks may contain audit/readme/helper Sheets. Only Sheets
                # carrying the benchmark schema are Sessions.
                continue

            turns: list[BenchmarkTurn] = []
            for zero_based_index, row in enumerate(rows):
                if max_turns_per_session is not None and len(turns) >= max_turns_per_session:
                    break
                values = list(row)

                def get(column: str, default: Any = None) -> Any:
                    index = index_by_name.get(column)
                    if index is None or index >= len(values):
                        return default
                    return values[index]

                turn_id = str(get("编号") or "").strip().upper()
                question = str(get("当前问题") or "").strip()
                answer = str(get("最终回答") or "").strip()
                if not turn_id and not question and not answer:
                    continue
                if not turn_id:
                    raise ValueError(f"Sheet {sheet_name} 第 {zero_based_index + 2} 行缺少编号")
                if not question:
                    raise ValueError(f"{turn_id} 缺少当前问题")
                if not answer:
                    raise ValueError(f"{turn_id} 缺少最终回答")

                turns.append(
                    BenchmarkTurn(
                        session_id=sheet_name,
                        sheet_name=sheet_name,
                        turn_index=len(turns),
                        turn_id=turn_id,
                        question=question,
                        answer=answer,
                        needs_history=_to_bool(get("是否需要前文")),
                        dependency_turn_ids=parse_dependency_turn_ids(get("关联前序对话")),
                        required_context=str(get("所需前文信息") or "").strip(),
                        dependency_type=str(get("依赖类型") or "").strip(),
                        max_lookback=_to_int(get("最大回溯轮数"), 0),
                    )
                )

            if not turns:
                continue
            if not any(TURN_ID_PATTERN.fullmatch(turn.turn_id) for turn in turns):
                if requested:
                    raise ValueError(f"Sheet {sheet_name} 不包含合法 Session 编号")
                continue
            validate_session_turns(sheet_name, turns)
            sessions.append(BenchmarkSession(session_id=sheet_name, sheet_name=sheet_name, turns=tuple(turns)))
            if max_sessions is not None and len(sessions) >= max_sessions:
                break
    finally:
        workbook.close()

    if requested:
        loaded = {session.sheet_name for session in sessions}
        missing_sheets = sorted(requested - loaded)
        if missing_sheets:
            raise ValueError(f"未找到指定 Sheet：{', '.join(missing_sheets)}")
    if not sessions:
        raise ValueError(f"数据集中没有可用 Session：{path}")
    return sessions


def validate_session_turns(sheet_name: str, turns: Sequence[BenchmarkTurn]) -> None:
    positions: dict[str, int] = {}
    for turn in turns:
        if turn.turn_id in positions:
            raise ValueError(f"Sheet {sheet_name} 出现重复编号：{turn.turn_id}")
        positions[turn.turn_id] = turn.turn_index
    for turn in turns:
        for dependency in turn.dependency_turn_ids:
            if dependency not in positions:
                raise ValueError(f"{turn.turn_id} 依赖不存在的轮次：{dependency}")
            if positions[dependency] >= turn.turn_index:
                raise ValueError(f"{turn.turn_id} 依赖当前或未来轮次：{dependency}")
        if turn.needs_history and not turn.dependency_turn_ids:
            raise ValueError(f"{turn.turn_id} 标记需要前文，但关联前序对话为空")


def dependency_distances(session: BenchmarkSession, turn: BenchmarkTurn) -> list[int]:
    positions = session.turn_position_by_id
    return [turn.turn_index - positions[dependency] for dependency in turn.dependency_turn_ids]


def is_long_range_turn(session: BenchmarkSession, turn: BenchmarkTurn, short_term_qa_capacity: int) -> bool:
    return bool(turn.dependency_turn_ids) and any(
        distance > short_term_qa_capacity for distance in dependency_distances(session, turn)
    )


def parse_turn_id_from_name(name: Any) -> str | None:
    if not isinstance(name, str):
        return None
    match = TURN_ID_PATTERN.search(name)
    return match.group(0).upper() if match else None


def ordered_unique(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


class LineageTracker:
    def __init__(self, short_term_qa_capacity: int):
        if short_term_qa_capacity < 0:
            raise ValueError("short_term_qa_capacity 不能为负数")
        self.short_term_qa_capacity = short_term_qa_capacity
        self._active_turns: dict[str, deque[str]] = defaultdict(deque)
        self.job_to_turn_ids: dict[str, tuple[str, ...]] = {}

    def register_add(self, session_id: str, turn_id: str, migration_job_id: str | None) -> tuple[str, ...]:
        active = self._active_turns[session_id]
        active.append(turn_id)
        overflow: list[str] = []
        while len(active) > self.short_term_qa_capacity:
            overflow.append(active.popleft())
        if overflow:
            if not migration_job_id:
                raise RuntimeError(
                    f"Session {session_id} 发生短期淘汰 {overflow}，但 add 未返回 migration_job_id；"
                    "请确认 background.enabled=True"
                )
            self.job_to_turn_ids[migration_job_id] = tuple(overflow)
        elif migration_job_id:
            # 正常情况下没有淘汰就不会创建迁移任务；保留空映射便于诊断异常。
            self.job_to_turn_ids.setdefault(migration_job_id, ())
        return tuple(overflow)

    def turn_ids_for_job(self, job_id: str | None) -> tuple[str, ...]:
        if not job_id:
            return ()
        return self.job_to_turn_ids.get(job_id, ())


def _source_job_ids(memory: Mapping[str, Any]) -> list[str]:
    result: list[str] = []
    for value in (memory.get("source_job_id"),):
        if value:
            result.append(str(value))
    source_job_ids = memory.get("source_job_ids")
    if isinstance(source_job_ids, (list, tuple, set)):
        result.extend(str(value) for value in source_job_ids if value)
    metadata = memory.get("metadata")
    if isinstance(metadata, Mapping):
        if metadata.get("source_job_id"):
            result.append(str(metadata["source_job_id"]))
        nested = metadata.get("source_job_ids")
        if isinstance(nested, (list, tuple, set)):
            result.extend(str(value) for value in nested if value)
    return ordered_unique(result)


def extract_retrieval_layers(
    context: Mapping[str, Any],
    lineage: LineageTracker,
) -> dict[str, list[str]]:
    short_ids = ordered_unique(
        turn_id
        for message in context.get("short_term_messages") or []
        if isinstance(message, Mapping)
        for turn_id in [parse_turn_id_from_name(message.get("name"))]
        if turn_id
    )

    mid_page_ids: list[str] = []
    mid_session_ids: list[str] = []
    long_ids: list[str] = []
    for memory in context.get("retrieved_memories") or []:
        if not isinstance(memory, Mapping):
            continue
        source = str(memory.get("source") or "")
        if source == "mid_term_page" or source == "midterm":
            target = mid_page_ids
        elif source == "mid_term_session":
            target = mid_session_ids
        else:
            target = long_ids
        for job_id in _source_job_ids(memory):
            target.extend(lineage.turn_ids_for_job(job_id))

    mid_page_ids = ordered_unique(mid_page_ids)
    mid_session_ids = ordered_unique(mid_session_ids)
    mid_ids = ordered_unique([*mid_page_ids, *mid_session_ids])
    long_ids = ordered_unique(long_ids)
    return {
        "short": short_ids,
        "mid_page": mid_page_ids,
        "mid_session": mid_session_ids,
        "mid": mid_ids,
        "long": long_ids,
        "mid_long": ordered_unique([*mid_ids, *long_ids]),
        "all": ordered_unique([*short_ids, *mid_ids, *long_ids]),
    }


def retrieval_metrics(gold_ids: Sequence[str], retrieved_ids: Sequence[str]) -> dict[str, float | int]:
    gold = set(gold_ids)
    retrieved = list(retrieved_ids)
    retrieved_set = set(retrieved)
    if not gold:
        return {
            "recall": 1.0,
            "precision": 1.0 if not retrieved else 0.0,
            "hit": 1,
            "full_recall": 1,
            "mrr": 1.0,
            "gold_count": 0,
            "retrieved_count": len(retrieved),
            "matched_count": 0,
        }
    matched = gold & retrieved_set
    first_rank = next((index + 1 for index, item in enumerate(retrieved) if item in gold), None)
    return {
        "recall": len(matched) / len(gold),
        "precision": len(matched) / len(retrieved_set) if retrieved_set else 0.0,
        "hit": int(bool(matched)),
        "full_recall": int(gold <= retrieved_set),
        "mrr": 1.0 / first_rank if first_rank else 0.0,
        "gold_count": len(gold),
        "retrieved_count": len(retrieved),
        "matched_count": len(matched),
    }


def percentile(values: Sequence[float], percent: float) -> float | None:
    if not values:
        return None
    ordered = sorted(float(value) for value in values)
    if len(ordered) == 1:
        return ordered[0]
    rank = (len(ordered) - 1) * percent
    lower = math.floor(rank)
    upper = math.ceil(rank)
    if lower == upper:
        return ordered[lower]
    fraction = rank - lower
    return ordered[lower] * (1 - fraction) + ordered[upper] * fraction


def latency_summary(values: Sequence[float]) -> dict[str, float | int | None]:
    clean = [float(value) for value in values if value is not None and math.isfinite(float(value))]
    if not clean:
        return {
            "count": 0,
            "min": None,
            "mean": None,
            "p50": None,
            "p90": None,
            "p95": None,
            "p99": None,
            "max": None,
            "stddev": None,
        }
    return {
        "count": len(clean),
        "min": min(clean),
        "mean": statistics.fmean(clean),
        "p50": percentile(clean, 0.50),
        "p90": percentile(clean, 0.90),
        "p95": percentile(clean, 0.95),
        "p99": percentile(clean, 0.99),
        "max": max(clean),
        "stddev": statistics.pstdev(clean) if len(clean) > 1 else 0.0,
    }


def mean_or_none(values: Sequence[float]) -> float | None:
    clean = [float(value) for value in values if value is not None]
    return statistics.fmean(clean) if clean else None


class AsyncJsonlWriter:
    def __init__(self, path: Path):
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._file = self.path.open("w", encoding="utf-8")
        self._lock = asyncio.Lock()

    async def write(self, value: Any) -> None:
        line = json.dumps(value, ensure_ascii=False, default=str)
        async with self._lock:
            self._file.write(line + "\n")
            self._file.flush()

    async def close(self) -> None:
        async with self._lock:
            if not self._file.closed:
                self._file.close()


def write_csv(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fieldnames: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                fieldnames.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _csv_value(row.get(key)) for key in fieldnames})


def _csv_value(value: Any) -> Any:
    if isinstance(value, (list, tuple, set)):
        return "|".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


def parse_iso_timestamp(value: Any) -> float | None:
    if not value:
        return None
    text = str(value).strip().replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=BEIJING_TIMEZONE)
    return parsed.timestamp()


async def wait_for_background_job(
    memory: Any,
    job_id: str | None,
    *,
    job_type: str,
    timeout_seconds: float,
    poll_interval_seconds: float,
) -> dict[str, Any] | None:
    if not job_id:
        return None
    terminal = MIGRATION_TERMINAL_STATUSES if job_type == "migration" else PROFILE_TERMINAL_STATUSES
    deadline = time.monotonic() + timeout_seconds
    while True:
        job = await asyncio.to_thread(memory.db.get_background_job, job_id, job_type)
        if job is None:
            raise RuntimeError(f"找不到后台任务：type={job_type}, id={job_id}")
        if str(job.get("status")) in terminal:
            return job
        if time.monotonic() >= deadline:
            raise TimeoutError(
                f"后台任务超时：type={job_type}, id={job_id}, status={job.get('status')}, "
                f"midterm={job.get('midterm_status')}, longterm={job.get('longterm_status')}"
            )
        await asyncio.sleep(poll_interval_seconds)


async def wait_for_migration_jobs(
    memory: Any,
    job_ids: Sequence[str],
    *,
    timeout_seconds: float,
    poll_interval_seconds: float,
) -> list[dict[str, Any]]:
    unique = ordered_unique(job_ids)
    if not unique:
        return []
    results = await asyncio.gather(
        *(
            wait_for_background_job(
                memory,
                job_id,
                job_type="migration",
                timeout_seconds=timeout_seconds,
                poll_interval_seconds=poll_interval_seconds,
            )
            for job_id in unique
        )
    )
    return [result for result in results if result is not None]


def query_queue_snapshot(db_path: Path) -> dict[str, Any]:
    if not db_path.exists():
        return {}
    connection = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True, timeout=1.0)
    try:
        snapshot: dict[str, Any] = {}
        for prefix, table, columns in (
            ("migration", "memory_migration_jobs", ("status", "midterm_status", "longterm_status")),
            ("profile", "profile_update_jobs", ("status",)),
        ):
            for column in columns:
                rows = connection.execute(
                    f"SELECT {column}, COUNT(*) FROM {table} GROUP BY {column}"
                ).fetchall()
                for status, count in rows:
                    snapshot[f"{prefix}_{column}_{status}"] = int(count)
        message_rows = connection.execute("SELECT status, COUNT(*) FROM messages GROUP BY status").fetchall()
        for status, count in message_rows:
            snapshot[f"messages_{status}"] = int(count)
        return snapshot
    finally:
        connection.close()


def safe_git_commit(repo_root: Path) -> str | None:
    head = repo_root / ".git" / "HEAD"
    if not head.exists():
        return None
    try:
        value = head.read_text(encoding="utf-8").strip()
        if value.startswith("ref: "):
            ref = repo_root / ".git" / value[5:]
            if ref.exists():
                return ref.read_text(encoding="utf-8").strip()
        return value
    except OSError:
        return None
